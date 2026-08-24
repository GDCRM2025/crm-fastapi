from __future__ import annotations

import csv
import hashlib
import io
import re
import unicodedata
from collections import defaultdict
from datetime import date, datetime, timedelta
from zoneinfo import ZoneInfo
from decimal import Decimal, InvalidOperation
from typing import Any

import os
import sys

_FIN_VENDOR = os.path.abspath(
    os.path.join(
        os.path.dirname(__file__),
        "..",
        "..",
        "vendor",
    )
)

if (
    os.path.isdir(_FIN_VENDOR)
    and _FIN_VENDOR not in sys.path
):
    sys.path.insert(0, _FIN_VENDOR)


from fastapi import APIRouter, Body, Depends, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy import text

from backend.routers import finanzas as legacy
from backend.core.database import get_db


router = APIRouter(
    prefix="/finanzas/simple",
    tags=["finanzas-simple"],
)

READ_ROLES = {
    "ADMIN",
    "SUPERADMIN",
    "FINANZAS",
    "COMPRAS",
    "JEFE DE OPERACIONES",
    "CONTROL DE GESTION",
}


# V1 Finanzas:
# - FINANZAS / ADMIN: carga + clasifica + P&L
# - COMPRAS / JEFE DE OPERACIONES: solo clasificacion
FIN_CLASSIFY_ROLES = {
    "ADMIN",
    "SUPERADMIN",
    "FINANZAS",
    "COMPRAS",
    "JEFE DE OPERACIONES",
    "CONTROL DE GESTION",
}

FIN_UPLOAD_ROLES = {
    "ADMIN",
    "SUPERADMIN",
    "FINANZAS",
    "CONTROL DE GESTION",
}

FIN_PNL_ROLES = {
    "ADMIN",
    "SUPERADMIN",
    "FINANZAS",
    "CONTROL DE GESTION",
}


def _role_classify(me: dict) -> str:
    return legacy._ensure_roles(
        me,
        FIN_CLASSIFY_ROLES,
    )


def _role_upload(me: dict) -> str:
    return legacy._ensure_roles(
        me,
        FIN_UPLOAD_ROLES,
    )


def _role_pnl(me: dict) -> str:
    return legacy._ensure_roles(
        me,
        FIN_PNL_ROLES,
    )


SIMPLE_ACCOUNTS = [
    ("5100", "Alimentos / mercadería", "COSTO EVENTOS"),
    ("5120", "Packaging / insumos", "COSTO EVENTOS"),
    ("5130", "Combustible / transporte evento", "COSTO EVENTOS"),
    ("5140", "Operadores / personal evento", "COSTO EVENTOS"),
    ("5150", "Servicios subcontratados", "COSTO EVENTOS"),
    ("5160", "Arriendos de eventos", "COSTO EVENTOS"),

    ("6100", "Transporte / autopistas", "OPERACIÓN"),
    ("6101", "Mantención vehículos", "OPERACIÓN"),
    ("6103", "Fletes / traslados", "OPERACIÓN"),

    ("6210", "Sueldos", "PERSONAL"),
    ("6211", "Imposiciones", "PERSONAL"),

    ("6310", "Arriendo", "GASTOS FIJOS"),
    ("6160", "Servicios básicos", "GASTOS FIJOS"),
    ("6220", "Telefonía / Internet", "GASTOS FIJOS"),
    ("6190", "Seguros", "GASTOS FIJOS"),
    ("6225", "Software / otros administrativos", "GASTOS FIJOS"),

    ("6140", "Marketing / publicidad", "MARKETING"),
    ("6150", "Contabilidad / legal / asesorías", "ADMINISTRACIÓN"),

    ("7010", "Intereses créditos", "FINANCIERO"),
    ("7020", "Gastos bancarios", "FINANCIERO"),
    ("7021", "Comisiones medios de pago", "FINANCIERO"),

    ("8002", "IVA no recuperable", "OTROS"),
]

SIMPLE_CODES = {item[0] for item in SIMPLE_ACCOUNTS}
ACCOUNT_META = {
    code: {"code": code, "label": label, "group": group}
    for code, label, group in SIMPLE_ACCOUNTS
}


def _role(me: dict) -> str:
    return legacy._ensure_roles(me, READ_ROLES)


def _actor(me: dict) -> int | None:
    for key in ("id_usuario", "user_id", "id", "sub"):
        value = me.get(key)
        try:
            if value is not None:
                return int(value)
        except Exception:
            pass
    return None


def _norm(value: Any) -> str:
    raw = str(value or "").strip().lower()
    raw = unicodedata.normalize("NFD", raw)
    raw = "".join(ch for ch in raw if unicodedata.category(ch) != "Mn")
    raw = re.sub(r"[^a-z0-9]+", " ", raw)
    return re.sub(r"\s+", " ", raw).strip()


def _parse_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    raw = str(value or "").strip()
    if not raw:
        return None

    for fmt in (
        "%d/%m/%Y",
        "%d-%m-%Y",
        "%Y-%m-%d",
        "%d/%m/%y",
        "%d-%m-%y",
    ):
        try:
            return datetime.strptime(raw[:10], fmt).date()
        except Exception:
            pass

    return None


def _parse_money(value: Any) -> Decimal:
    if value in (None, ""):
        return Decimal("0")

    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))

    raw = str(value).strip()
    if not raw:
        return Decimal("0")

    negative = raw.startswith("(") and raw.endswith(")")
    raw = raw.replace("(", "").replace(")", "")
    raw = raw.replace("$", "").replace("CLP", "").replace(" ", "")
    raw = re.sub(r"[^0-9,.\-]", "", raw)

    if "," in raw and "." in raw:
        if raw.rfind(",") > raw.rfind("."):
            raw = raw.replace(".", "").replace(",", ".")
        else:
            raw = raw.replace(",", "")
    elif "," in raw:
        parts = raw.split(",")
        if len(parts) == 2 and len(parts[-1]) <= 2:
            raw = raw.replace(",", ".")
        else:
            raw = raw.replace(",", "")
    elif raw.count(".") > 1:
        raw = raw.replace(".", "")
    elif "." in raw:
        left, right = raw.rsplit(".", 1)
        if len(right) == 3 and left.replace("-", "").isdigit():
            raw = raw.replace(".", "")

    try:
        result = Decimal(raw or "0")
    except InvalidOperation:
        return Decimal("0")

    return -result if negative else result


ALIASES = {
    "date": {
        "fecha",
        "fecha movimiento",
        "fecha transaccion",
        "fecha operacion",
        "fecha contable",
    },
    "description": {
        "descripcion",
        "detalle",
        "glosa",
        "movimiento",
        "descripcion movimiento",
        "detalle movimiento",
    },
    "debit": {
        "cargo",
        "cargos",
        "debito",
        "debitos",
        "debe",
        "egreso",
        "egresos",
    },
    "credit": {
        "abono",
        "abonos",
        "credito",
        "creditos",
        "haber",
        "ingreso",
        "ingresos",
    },
    "amount": {
        "monto",
        "importe",
        "monto movimiento",
        "valor",
    },
    "reference": {
        "referencia",
        "numero documento",
        "n documento",
        "nro documento",
        "numero operacion",
        "nro operacion",
    },
    "balance": {
        "saldo",
        "saldo disponible",
        "saldo contable",
    },
}


def _header_map(rows: list[list[Any]]) -> tuple[int, dict[str, int]]:
    best_row = -1
    best: dict[str, int] = {}
    best_score = -1

    for row_index, row in enumerate(rows[:20]):
        current: dict[str, int] = {}

        for col_index, value in enumerate(row):
            normalized = _norm(value)

            for key, aliases in ALIASES.items():
                if normalized in aliases:
                    current.setdefault(key, col_index)

        score = len(current)

        required = (
            "date" in current
            and "description" in current
            and (
                "amount" in current
                or "debit" in current
                or "credit" in current
            )
        )

        if required and score > best_score:
            best_row = row_index
            best = current
            best_score = score

    if best_row < 0:
        raise HTTPException(
            status_code=400,
            detail=(
                "No pude reconocer las columnas de la cartola. "
                "Necesito Fecha + Descripción + Monto, "
                "o Fecha + Descripción + Cargo/Abono."
            ),
        )

    return best_row, best


def _xlsx_rows(
    content: bytes,
) -> list[list[Any]]:
    workbook = load_workbook(
        io.BytesIO(content),
        read_only=True,
        data_only=True,
    )

    sheet = workbook.active

    return [
        list(row)
        for row in sheet.iter_rows(
            values_only=True
        )
    ]


def _xls_rows(
    content: bytes,
) -> list[list[Any]]:
    try:
        import xlrd
    except Exception as exc:
        raise HTTPException(
            status_code=500,
            detail="Lector XLS no disponible.",
        ) from exc

    try:
        workbook = xlrd.open_workbook(
            file_contents=content
        )

        sheet = workbook.sheet_by_index(0)

        rows = []

        for row_index in range(sheet.nrows):
            current = []

            for col_index in range(sheet.ncols):
                cell = sheet.cell(
                    row_index,
                    col_index,
                )

                value = cell.value

                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        value = xlrd.xldate_as_datetime(
                            value,
                            workbook.datemode,
                        )
                    except Exception:
                        pass

                current.append(value)

            rows.append(current)

        return rows

    except HTTPException:
        raise

    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail=(
                "No pude leer el XLS. "
                "Verifica que sea una cartola Excel válida."
            ),
        ) from exc


def _clean_pdf_table(
    table,
) -> list[list[Any]]:
    rows = []

    for row in table or []:
        if not row:
            continue

        cleaned = []

        for value in row:
            if value is None:
                cleaned.append("")
                continue

            cleaned.append(
                str(value)
                .replace("\r", " ")
                .replace("\n", " ")
                .strip()
            )

        if any(
            str(value).strip()
            for value in cleaned
        ):
            rows.append(cleaned)

    return rows


def _pdf_rows(
    content: bytes,
) -> list[list[Any]]:
    try:
        import pdfplumber
    except Exception as exc:
        raise HTTPException(
            status_code=500,
            detail="Lector PDF no disponible.",
        ) from exc

    try:
        with pdfplumber.open(
            io.BytesIO(content)
        ) as pdf:

            # Primero intentamos tablas reales.
            for page in pdf.pages:
                try:
                    tables = (
                        page.extract_tables()
                        or []
                    )
                except Exception:
                    tables = []

                for table in tables:
                    rows = _clean_pdf_table(
                        table
                    )

                    if len(rows) < 2:
                        continue

                    try:
                        _header_map(rows)
                        return rows
                    except HTTPException:
                        pass

            # Segundo intento:
            # texto preservando disposición visual.
            text_lines = []

            for page in pdf.pages:
                try:
                    text_value = (
                        page.extract_text(
                            x_tolerance=2,
                            y_tolerance=3,
                            layout=True,
                        )
                        or ""
                    )
                except TypeError:
                    text_value = (
                        page.extract_text()
                        or ""
                    )

                text_lines.extend(
                    line.rstrip()
                    for line in text_value.splitlines()
                    if line.strip()
                )

    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail=(
                "No pude abrir este PDF. "
                "Puede estar protegido o dañado."
            ),
        ) from exc

    if not text_lines:
        raise HTTPException(
            status_code=400,
            detail=(
                "El PDF no contiene texto bancario "
                "utilizable. Si es un PDF escaneado "
                "como imagen, todavía no se admite."
            ),
        )

    candidate_rows = []

    for line in text_lines:
        columns = [
            value.strip()
            for value in re.split(
                r"\s{2,}",
                line.strip(),
            )
            if value.strip()
        ]

        if columns:
            candidate_rows.append(
                columns
            )

    try:
        _header_map(candidate_rows)
        return candidate_rows

    except HTTPException as exc:
        raise HTTPException(
            status_code=400,
            detail=(
                "Pude leer el PDF, pero no puedo "
                "identificar con seguridad las columnas "
                "Fecha + Descripción + Cargo/Abono/Monto. "
                "No se importó ningún movimiento."
            ),
        ) from exc



def _normalize_banco_chile_rows(
    rows: list[list[Any]],
) -> list[list[Any]]:
    """
    Normaliza el formato nativo de cartola Banco de Chile.

    Formato detectado:
      Fecha
      Descripción
      Canal o Sucursal
      Nro. Docto.
      Cargos (CLP)
      Abonos (CLP)
      Saldo (CLP)

    Banco de Chile intercala columnas vacías y coloca
    información general antes de la tabla de movimientos.
    """

    if not rows:
        return rows

    def find_col(
        values: list[str],
        predicate,
    ) -> int | None:
        for index, value in enumerate(values):
            if predicate(value):
                return index

        return None

    header_index = None
    columns = {}

    # Banco de Chile coloca normalmente la cabecera
    # dentro de las primeras 80 filas.
    for index, row in enumerate(rows[:80]):

        normalized = [
            _norm(value)
            for value in row
        ]

        candidate = {
            "date":
                find_col(
                    normalized,
                    lambda value:
                        value == "fecha",
                ),

            "description":
                find_col(
                    normalized,
                    lambda value:
                        value == "descripcion",
                ),

            "channel":
                find_col(
                    normalized,
                    lambda value:
                        "canal" in value
                        and "sucursal" in value,
                ),

            "document":
                find_col(
                    normalized,
                    lambda value:
                        (
                            value.startswith("nro")
                            or value.startswith("numero")
                        )
                        and "docto" in value,
                ),

            "debit":
                find_col(
                    normalized,
                    lambda value:
                        value.startswith("cargos"),
                ),

            "credit":
                find_col(
                    normalized,
                    lambda value:
                        value.startswith("abonos"),
                ),

            "balance":
                find_col(
                    normalized,
                    lambda value:
                        value.startswith("saldo")
                        and value != "saldos",
                ),
        }

        if (
            candidate["date"] is not None
            and candidate["description"] is not None
            and candidate["debit"] is not None
            and candidate["credit"] is not None
        ):
            header_index = index
            columns = candidate
            break

    # No es Banco de Chile.
    if header_index is None:
        return rows

    expected_debit = None
    expected_credit = None

    # Buscar "Total Cargos" / "Total Abonos"
    # en el resumen anterior a los movimientos.
    for index, row in enumerate(
        rows[:header_index]
    ):

        normalized = [
            _norm(value)
            for value in row
        ]

        debit_label = find_col(
            normalized,
            lambda value:
                value == "total cargos",
        )

        credit_label = find_col(
            normalized,
            lambda value:
                value == "total abonos",
        )

        if (
            debit_label is not None
            and credit_label is not None
            and index + 1 < len(rows)
        ):
            totals_row = rows[index + 1]

            if debit_label < len(totals_row):
                expected_debit = _parse_money(
                    totals_row[debit_label]
                )

            if credit_label < len(totals_row):
                expected_credit = _parse_money(
                    totals_row[credit_label]
                )

            break

    def value(
        row: list[Any],
        key: str,
    ) -> Any:
        position = columns.get(key)

        if (
            position is None
            or position >= len(row)
        ):
            return None

        return row[position]

    normalized_rows: list[list[Any]] = [
        [
            "Fecha",
            "Descripción",
            "Canal o Sucursal",
            "Nro. Docto.",
            "Cargo",
            "Abono",
            "Saldo",
        ]
    ]

    debit_total = Decimal("0")
    credit_total = Decimal("0")

    for row in rows[header_index + 1:]:

        date_value = value(
            row,
            "date",
        )

        description = str(
            value(
                row,
                "description",
            )
            or ""
        ).strip()

        parsed_date = _parse_date(
            date_value
        )

        debit = _parse_money(
            value(
                row,
                "debit",
            )
        )

        credit = _parse_money(
            value(
                row,
                "credit",
            )
        )

        # Ignorar pie de página, textos y filas vacías.
        if parsed_date is None:
            continue

        if not description:
            continue

        if (
            debit == Decimal("0")
            and credit == Decimal("0")
        ):
            continue

        normalized_rows.append(
            [
                date_value,
                description,
                value(
                    row,
                    "channel",
                ),
                value(
                    row,
                    "document",
                ),
                debit,
                credit,
                value(
                    row,
                    "balance",
                ),
            ]
        )

        debit_total += debit
        credit_total += credit

    if len(normalized_rows) <= 1:
        raise HTTPException(
            status_code=400,
            detail=(
                "Reconocí una cartola Banco de Chile, "
                "pero no encontré movimientos válidos."
            ),
        )

    tolerance = Decimal("0.01")

    if (
        expected_debit is not None
        and abs(
            debit_total
            - expected_debit
        ) > tolerance
    ):
        raise HTTPException(
            status_code=400,
            detail=(
                "La cartola Banco de Chile no cuadra: "
                "la suma de cargos no coincide con "
                "el Total Cargos informado por el banco. "
                "No se importó ningún movimiento."
            ),
        )

    if (
        expected_credit is not None
        and abs(
            credit_total
            - expected_credit
        ) > tolerance
    ):
        raise HTTPException(
            status_code=400,
            detail=(
                "La cartola Banco de Chile no cuadra: "
                "la suma de abonos no coincide con "
                "el Total Abonos informado por el banco. "
                "No se importó ningún movimiento."
            ),
        )

    return normalized_rows


def _read_rows(
    content: bytes,
    filename: str,
) -> list[list[Any]]:
    lower = filename.lower().strip()

    if lower.endswith(".xlsx"):
        return _normalize_banco_chile_rows(
            _xlsx_rows(content)
        )

    if lower.endswith(".xls"):
        return _normalize_banco_chile_rows(
            _xls_rows(content)
        )

    if lower.endswith(".pdf"):
        return _pdf_rows(content)

    if (
        lower.endswith(".csv")
        or lower.endswith(".txt")
    ):
        decoded = None

        for encoding in (
            "utf-8-sig",
            "utf-8",
            "latin-1",
        ):
            try:
                decoded = content.decode(
                    encoding
                )
                break

            except UnicodeDecodeError:
                pass

        if decoded is None:
            raise HTTPException(
                status_code=400,
                detail="No pude leer el archivo.",
            )

        sample = decoded[:5000]

        try:
            delimiter = csv.Sniffer().sniff(
                sample,
                delimiters=",;\t|",
            ).delimiter

        except Exception:
            delimiter = ";"

        return list(
            csv.reader(
                io.StringIO(decoded),
                delimiter=delimiter,
            )
        )

    raise HTTPException(
        status_code=400,
        detail=(
            "Formato no soportado. "
            "Usa PDF, XLSX o XLS."
        ),
    )


def _brands(conn) -> list[str]:
    rows = conn.execute(
        text(
            """
            SELECT marca
            FROM public.marcas
            ORDER BY id_marca
            """
        )
    ).scalars().all()

    return [str(item).strip().upper() for item in rows if item]


def _validate_brand(conn, brand: str) -> str:
    normalized = str(brand or "GENERAL").strip().upper()

    if normalized == "GENERAL":
        return normalized

    if normalized not in set(_brands(conn)):
        raise HTTPException(status_code=400, detail="Marca inválida.")

    return normalized


def _validate_simple_account(conn, code: str) -> str:
    code = str(code or "").strip()

    if code not in SIMPLE_CODES:
        raise HTTPException(
            status_code=400,
            detail="Selecciona una cuenta operativa válida.",
        )

    row = conn.execute(
        text(
            """
            SELECT type,is_active
            FROM public.plan_cuentas
            WHERE code=:code
            """
        ),
        {"code": code},
    ).mappings().first()

    if not row:
        raise HTTPException(status_code=400, detail="La cuenta no existe.")

    if str(row["type"] or "").lower() != "expense":
        raise HTTPException(status_code=400, detail="La cuenta no es de gasto.")

    if row["is_active"] is False:
        raise HTTPException(status_code=400, detail="La cuenta está inactiva.")

    return code


@router.get("/meta")
def meta(me: dict = Depends(legacy.get_current_user)):
    _role_classify(me)

    with legacy.get_connection() as conn:
        brands = _brands(conn)

        existing = {
            str(row["code"])
            for row in conn.execute(
                text(
                    """
                    SELECT code
                    FROM public.plan_cuentas
                    WHERE type='expense'
                      AND COALESCE(is_active,TRUE)=TRUE
                    """
                )
            ).mappings().all()
        }

        accounts = [
            ACCOUNT_META[code]
            for code in SIMPLE_CODES
            if code in existing
        ]

        accounts.sort(
            key=lambda item: (
                item["group"],
                item["label"],
            )
        )

        bank_accounts = [
            dict(row)
            for row in conn.execute(
                text(
                    """
                    SELECT
                      id_bank_account,
                      bank_name,
                      label,
                      currency
                    FROM public.fin_bank_accounts
                    WHERE is_active=TRUE
                    ORDER BY bank_name,label
                    """
                )
            ).mappings().all()
        ]

    return {
        "ok": True,
        "brands": ["GENERAL"] + brands,
        "accounts": accounts,
        "bank_accounts": bank_accounts,
    }


@router.post("/bank-accounts")
def create_bank_account(
    body: dict = Body(...),
    me: dict = Depends(legacy.get_current_user),
):
    _role_upload(me)

    bank_name = str(body.get("bank_name") or "").strip()
    label = str(body.get("label") or "").strip()

    if not bank_name or not label:
        raise HTTPException(
            status_code=400,
            detail="Banco y nombre de cuenta son requeridos.",
        )

    with legacy.get_connection() as conn:
        existing = conn.execute(
            text(
                """
                SELECT id_bank_account
                FROM public.fin_bank_accounts
                WHERE upper(bank_name)=upper(:bank)
                  AND upper(label)=upper(:label)
                  AND is_active=TRUE
                LIMIT 1
                """
            ),
            {
                "bank": bank_name,
                "label": label,
            },
        ).scalar()

        if existing is not None:
            return {
                "ok": True,
                "id_bank_account": int(existing),
                "existing": True,
            }

        active_count = conn.execute(
            text(
                """
                SELECT count(*)
                FROM public.fin_bank_accounts
                WHERE is_active=TRUE
                """
            )
        ).scalar_one()

        bank_id = conn.execute(
            text(
                """
                INSERT INTO public.fin_bank_accounts(
                  bank_name,
                  label
                )
                VALUES(
                  :bank,
                  :label
                )
                RETURNING id_bank_account
                """
            ),
            {
                "bank": bank_name,
                "label": label,
            },
        ).scalar_one()

        conn.commit()

    return {
        "ok": True,
        "id_bank_account": int(bank_id),
        "existing": False,
    }


@router.post("/cartolas/upload")
async def upload_statement(
    id_bank_account: int = Form(...),
    file: UploadFile = File(...),
    me: dict = Depends(legacy.get_current_user),
):
    _role_upload(me)

    content = await file.read()

    if not content:
        raise HTTPException(status_code=400, detail="Archivo vacío.")

    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="Archivo demasiado grande.")

    filename = str(file.filename or "cartola")
    rows = _read_rows(content, filename)
    header_row, columns = _header_map(rows)

    sha = hashlib.sha256(content).hexdigest()

    parsed = []
    occurrences: dict[str, int] = defaultdict(int)

    for source_row, row in enumerate(
        rows[header_row + 1:],
        start=header_row + 2,
    ):
        def value(key: str):
            idx = columns.get(key)
            if idx is None or idx >= len(row):
                return None
            return row[idx]

        tx_date = _parse_date(value("date"))
        description = str(value("description") or "").strip()

        if not tx_date or not description:
            continue

        if "amount" in columns:
            amount = _parse_money(value("amount"))
        else:
            debit = abs(_parse_money(value("debit")))
            credit = abs(_parse_money(value("credit")))
            amount = credit - debit

        if amount == 0:
            continue

        reference = str(value("reference") or "").strip() or None
        balance_value = value("balance")
        balance = (
            _parse_money(balance_value)
            if balance_value not in (None, "")
            else None
        )

        base = "|".join(
            [
                str(id_bank_account),
                tx_date.isoformat(),
                _norm(description),
                str(amount.quantize(Decimal("0.01"))),
                _norm(reference),
                (
                    str(balance.quantize(Decimal("0.01")))
                    if balance is not None
                    else ""
                ),
            ]
        )

        occurrences[base] += 1

        dedupe_key = hashlib.sha256(
            f"{base}|{occurrences[base]}".encode("utf-8")
        ).hexdigest()

        parsed.append(
            {
                "source_row": source_row,
                "tx_date": tx_date,
                "description": description[:500],
                "amount": amount,
                "reference": reference[:250] if reference else None,
                "balance": balance,
                "dedupe_key": dedupe_key,
            }
        )

    if not parsed:
        raise HTTPException(
            status_code=400,
            detail="La cartola no contiene movimientos reconocibles.",
        )

    actor = _actor(me)

    with legacy.get_connection() as conn:
        account_exists = conn.execute(
            text(
                """
                SELECT 1
                FROM public.fin_bank_accounts
                WHERE id_bank_account=:id
                  AND is_active=TRUE
                """
            ),
            {"id": id_bank_account},
        ).scalar()

        if not account_exists:
            raise HTTPException(status_code=400, detail="Cuenta bancaria inválida.")

        previous = conn.execute(
            text(
                """
                SELECT id_bank_import,imported_rows
                FROM public.fin_bank_imports
                WHERE id_bank_account=:account
                  AND file_sha256=:sha
                """
            ),
            {
                "account": id_bank_account,
                "sha": sha,
            },
        ).mappings().first()

        if previous:
            return {
                "ok": True,
                "duplicate_file": True,
                "id_bank_import": int(previous["id_bank_import"]),
                "imported_rows": int(previous["imported_rows"]),
            }

        import_id = conn.execute(
            text(
                """
                INSERT INTO public.fin_bank_imports(
                  id_bank_account,
                  original_name,
                  file_sha256,
                  created_by
                )
                VALUES(
                  :account,
                  :name,
                  :sha,
                  :actor
                )
                RETURNING id_bank_import
                """
            ),
            {
                "account": id_bank_account,
                "name": filename[:300],
                "sha": sha,
                "actor": actor,
            },
        ).scalar_one()

        inserted = 0

        for item in parsed:
            result = conn.execute(
                text(
                    """
                    INSERT INTO public.fin_bank_movements(
                      id_bank_import,
                      id_bank_account,
                      source_row,
                      tx_date,
                      description,
                      amount,
                      reference,
                      balance,
                      dedupe_key
                    )
                    VALUES(
                      :import_id,
                      :account,
                      :source_row,
                      :tx_date,
                      :description,
                      :amount,
                      :reference,
                      :balance,
                      :dedupe_key
                    )
                    ON CONFLICT (dedupe_key) DO NOTHING
                    """
                ),
                {
                    **item,
                    "import_id": import_id,
                    "account": id_bank_account,
                },
            )

            inserted += max(int(result.rowcount or 0), 0)

        conn.execute(
            text(
                """
                UPDATE public.fin_bank_imports
                SET imported_rows=:rows
                WHERE id_bank_import=:id
                """
            ),
            {
                "rows": inserted,
                "id": import_id,
            },
        )

        conn.commit()

    return {
        "ok": True,
        "duplicate_file": False,
        "id_bank_import": int(import_id),
        "imported_rows": inserted,
    }


@router.get("/movements")
def movements(
    status: str = Query("PENDING"),
    limit: int = Query(300, ge=1, le=1000),
    me: dict = Depends(legacy.get_current_user),
):
    _role_classify(me)

    status = status.strip().upper()

    if status not in {"PENDING", "CLASSIFIED", "IGNORED"}:
        raise HTTPException(status_code=400, detail="Estado inválido.")

    with legacy.get_connection() as conn:
        rows = conn.execute(
            text(
                """
                SELECT
                  m.id_bank_movement,
                  m.tx_date,
                  m.description,
                  m.amount,
                  m.reference,
                  m.balance,
                  m.status,
                  m.cuenta_code,
                  m.marca,
                  a.bank_name,
                  a.label AS bank_account
                FROM public.fin_bank_movements m
                JOIN public.fin_bank_accounts a
                  ON a.id_bank_account=m.id_bank_account
                WHERE m.status=:status
                ORDER BY m.tx_date DESC,m.id_bank_movement DESC
                LIMIT :limit
                """
            ),
            {
                "status": status,
                "limit": limit,
            },
        ).mappings().all()

    return {
        "ok": True,
        "items": [dict(row) for row in rows],
    }


@router.post("/movements/classify")
def classify_movements(
    body: dict = Body(...),
    me: dict = Depends(legacy.get_current_user),
):
    _role_classify(me)

    ids = body.get("movement_ids") or []
    code = str(body.get("cuenta_code") or "").strip()
    brand = str(body.get("marca") or "GENERAL").strip().upper()

    if not isinstance(ids, list) or not ids:
        raise HTTPException(status_code=400, detail="Selecciona movimientos.")

    actor = _actor(me)

    with legacy.get_connection() as conn:
        code = _validate_simple_account(conn, code)
        brand = _validate_brand(conn, brand)

        classified = 0

        for raw_id in ids:
            movement_id = int(raw_id)

            movement = conn.execute(
                text(
                    """
                    SELECT *
                    FROM public.fin_bank_movements
                    WHERE id_bank_movement=:id
                    FOR UPDATE
                    """
                ),
                {"id": movement_id},
            ).mappings().first()

            if not movement:
                continue

            amount = Decimal(str(movement["amount"]))

            if amount >= 0:
                raise HTTPException(
                    status_code=400,
                    detail=(
                        "Los abonos no se imputan automáticamente "
                        "como gasto. Usa 'No es gasto'."
                    ),
                )

            expense_amount = abs(amount)

            conn.execute(
                text(
                    """
                    INSERT INTO public.fin_gastos(
                      fecha,
                      cuenta_code,
                      monto,
                      descripcion,
                      proveedor,
                      marca,
                      centro_costo,
                      created_at,
                      doc_num,
                      pagado,
                      fecha_pago,
                      tipo_doc,
                      is_active,
                      bank_movement_id
                    )
                    VALUES(
                      :fecha,
                      :cuenta,
                      :monto,
                      :descripcion,
                      :proveedor,
                      :marca,
                      NULL,
                      now(),
                      :doc_num,
                      TRUE,
                      :fecha_pago,
                      'CARTOLA',
                      TRUE,
                      :movement
                    )
                    ON CONFLICT (bank_movement_id)
                    DO UPDATE SET
                      cuenta_code=EXCLUDED.cuenta_code,
                      marca=EXCLUDED.marca,
                      monto=EXCLUDED.monto,
                      descripcion=EXCLUDED.descripcion,
                      proveedor=EXCLUDED.proveedor,
                      fecha=EXCLUDED.fecha,
                      pagado=TRUE,
                      fecha_pago=EXCLUDED.fecha_pago,
                      is_active=TRUE
                    """
                ),
                {
                    "fecha": movement["tx_date"],
                    "cuenta": code,
                    "monto": expense_amount,
                    "descripcion": movement["description"],
                    "proveedor": str(movement["description"])[:200],
                    "marca": brand,
                    "doc_num": movement["reference"],
                    "fecha_pago": movement["tx_date"],
                    "movement": movement_id,
                },
            )

            conn.execute(
                text(
                    """
                    UPDATE public.fin_bank_movements
                    SET
                      status='CLASSIFIED',
                      cuenta_code=:cuenta,
                      marca=:marca,
                      classified_by=:actor,
                      classified_at=now()
                    WHERE id_bank_movement=:id
                    """
                ),
                {
                    "cuenta": code,
                    "marca": brand,
                    "actor": actor,
                    "id": movement_id,
                },
            )

            classified += 1

        conn.commit()

    return {
        "ok": True,
        "classified": classified,
    }


@router.post("/movements/ignore")
def ignore_movements(
    body: dict = Body(...),
    me: dict = Depends(legacy.get_current_user),
):
    _role_classify(me)

    ids = body.get("movement_ids") or []

    if not isinstance(ids, list) or not ids:
        raise HTTPException(status_code=400, detail="Selecciona movimientos.")

    actor = _actor(me)

    with legacy.get_connection() as conn:
        changed = 0

        for raw_id in ids:
            result = conn.execute(
                text(
                    """
                    UPDATE public.fin_bank_movements
                    SET
                      status='IGNORED',
                      classified_by=:actor,
                      classified_at=now()
                    WHERE id_bank_movement=:id
                      AND status='PENDING'
                    """
                ),
                {
                    "actor": actor,
                    "id": int(raw_id),
                },
            )

            changed += max(int(result.rowcount or 0), 0)

        conn.commit()

    return {
        "ok": True,
        "ignored": changed,
    }


def _revenue_for_brand(payload: dict, brand: str) -> float:
    if brand == "CONSOLIDADO":
        return float((payload.get("totals") or {}).get("ingresos") or 0)

    if brand == "GENERAL":
        return 0.0

    # === GD PNL V62 BRAND ALIASES ===
    revenue_aliases = {
        "BRONTOS": "ROLFI",
    }

    needle = _norm(
        revenue_aliases.get(
            brand,
            brand,
        )
    )

    for item in payload.get("items") or []:
        if str(item.get("type") or "").lower() != "revenue":
            continue

        name = _norm(item.get("name"))

        if name.startswith("ingresos") and needle in name:
            return float(item.get("balance") or 0)

    return 0.0


def _classified_by_month(
    year: int,
    brand: str,
) -> dict[tuple[int, str], float]:
    params: dict[str, Any] = {"year": year}
    brand_where = ""

    if brand == "GENERAL":
        brand_where = """
          AND (
            g.marca IS NULL
            OR btrim(g.marca)=''
            OR upper(btrim(g.marca))='GENERAL'
          )
        """
    elif brand != "CONSOLIDADO":
        brand_where = " AND upper(btrim(g.marca))=:brand "
        params["brand"] = brand

    with legacy.get_connection() as conn:
        rows = conn.execute(
            text(
                f"""
                SELECT
                  EXTRACT(MONTH FROM g.fecha)::int AS month,
                  g.cuenta_code,
                  SUM(g.monto)::numeric AS total
                FROM public.fin_gastos g
                WHERE g.is_active=TRUE
                  AND EXTRACT(YEAR FROM g.fecha)::int=:year
                  {brand_where}
                GROUP BY
                  EXTRACT(MONTH FROM g.fecha)::int,
                  g.cuenta_code
                """
            ),
            params,
        ).mappings().all()

    output: dict[tuple[int, str], float] = {}

    for row in rows:
        code = str(row["cuenta_code"] or "")

        if code in SIMPLE_CODES:
            output[(int(row["month"]), code)] = float(row["total"] or 0)

    return output



# === GD PNL V62 CLASSIFIED INCOME HELPER ===

def _pnl_v62_classified_income(
    year: int,
    month: int,
    brand: str,
) -> float:

    brand = str(
        brand or ""
    ).strip().upper()

    params = {
        "year": int(year),
        "month": int(month),
    }

    unit_where = ""

    if brand == "CONSOLIDADO":
        unit_where = ""

    elif brand == "GENERAL":
        unit_where = """
          AND (
            bm.marca IS NULL
            OR btrim(bm.marca)=''
            OR upper(btrim(bm.marca))='GENERAL'
          )
        """

    else:
        unit_where = """
          AND upper(btrim(bm.marca))=:brand
        """
        params["brand"] = brand

    with legacy.get_connection() as conn:

        value = conn.execute(
            text(
                f"""
                SELECT
                  COALESCE(
                    SUM(bm.amount),
                    0
                  )::numeric
                FROM public.fin_bank_movements bm

                JOIN public.plan_cuentas pc
                  ON pc.code=bm.cuenta_code

                WHERE bm.status='CLASSIFIED'
                  AND bm.amount > 0

                  AND pc.type='revenue'
                  AND COALESCE(
                        pc.is_active,
                        TRUE
                      )=TRUE

                  AND EXTRACT(
                        YEAR FROM bm.tx_date
                      )::int=:year

                  AND EXTRACT(
                        MONTH FROM bm.tx_date
                      )::int=:month

                  {unit_where}

                  AND NOT EXISTS (
                    SELECT 1
                    FROM public.fin_bank_event_allocations a
                    WHERE
                      a.id_bank_movement=
                        bm.id_bank_movement
                      AND a.is_active IS TRUE
                  )
                """
            ),
            params,
        ).scalar()

    return float(
        value or 0
    )



# === GD PNL V63 INDIVIDUAL RECURRING ===

PNL_V63_ENTITY_BY_UNIT = {
    "CAMALEON": "CARRITOS_EVENTOS",
    "GOURMET": "GD_ALIMENTOS",
    "EXPRESS": "AFTERMATH",
    "DEL SABOR": "SERVICIOS_GASTRONOMICOS",
    "BRONTOS": "ROLFI",
    "CORPORATIVO": "GREEN_HOLDING",
}


def _pnl_v63_recurring_by_month(
    year: int,
    brand: str,
) -> dict[tuple[int, str], float]:

    brand = str(
        brand or ""
    ).strip().upper()

    output: dict[tuple[int, str], float] = {}

    with legacy.get_connection() as conn:

        # HOLDING:
        # usa asignacion de gestion explicita,
        # no una sociedad bancaria.
        if brand == "HOLDING":

            rows = conn.execute(
                text(
                    """
                    SELECT
                      gs.month::int AS month,
                      t.cuenta_code,
                      SUM(
                        COALESCE(
                          (
                            SELECT rm.amount
                            FROM public.fin_recurring_expense_months rm
                            WHERE
                              rm.id_recurring_template =
                                t.id_recurring_template
                              AND rm.year=:year
                              AND rm.month=gs.month
                              AND rm.is_active IS TRUE
                            ORDER BY
                              rm.id_recurring_month DESC
                            LIMIT 1
                          ),
                          t.monthly_amount
                        )
                      )::numeric AS total

                    FROM public.fin_management_recurring_assignments ma

                    JOIN public.fin_recurring_expense_templates t
                      ON t.recurring_code=ma.recurring_code
                     AND t.is_active IS TRUE

                    CROSS JOIN generate_series(1,12) AS gs(month)

                    WHERE
                      ma.unit_code='HOLDING'
                      AND ma.is_active IS TRUE

                    GROUP BY
                      gs.month,
                      t.cuenta_code
                    """
                ),
                {
                    "year": int(year),
                },
            ).mappings().all()

        else:

            legal_code = PNL_V63_ENTITY_BY_UNIT.get(
                brand
            )

            if not legal_code:
                return output

            # Marcas / Brontos / Corporativo:
            # recurrentes asignados a la sociedad
            # operacional correspondiente.
            rows = conn.execute(
                text(
                    """
                    SELECT
                      gs.month::int AS month,
                      t.cuenta_code,
                      SUM(
                        a.monthly_amount
                      )::numeric AS total

                    FROM public.fin_recurring_expense_allocations a

                    JOIN public.fin_recurring_expense_templates t
                      ON t.id_recurring_template=
                         a.id_recurring_template
                     AND t.is_active IS TRUE

                    JOIN public.fin_legal_entities e
                      ON e.id_legal_entity=
                         a.id_legal_entity
                     AND e.is_active IS TRUE

                    CROSS JOIN generate_series(1,12) AS gs(month)

                    WHERE
                      a.is_active IS TRUE
                      AND e.legal_code=:legal_code

                    GROUP BY
                      gs.month,
                      t.cuenta_code
                    """
                ),
                {
                    "legal_code": legal_code,
                },
            ).mappings().all()

    for row in rows:
        code = str(
            row["cuenta_code"] or ""
        ).strip()

        if not code:
            continue

        output[
            (
                int(row["month"]),
                code,
            )
        ] = float(
            row["total"] or 0
        )

    return output


def _build_matrix(year: int, brand: str, me: dict) -> dict:
    brand = str(brand or "CONSOLIDADO").strip().upper()

    with legacy.get_connection() as conn:
        valid_brands = set(_brands(conn))

    if brand not in valid_brands | {
        "CONSOLIDADO",
        "GENERAL",
        "CORPORATIVO",
        "HOLDING",
    }:
        raise HTTPException(status_code=400, detail="Marca inválida.")

    classified = _classified_by_month(year, brand)
    recurring = _pnl_v63_recurring_by_month(
        year,
        brand,
    )

    rows = []

    ingresos = {
        month: 0.0
        for month in range(1, 13)
    }

    residual = {
        month: 0.0
        for month in range(1, 13)
    }

    total_gastos = {
        month: 0.0
        for month in range(1, 13)
    }

    resultado = {
        month: 0.0
        for month in range(1, 13)
    }

    legacy_totals = {}

    for month in range(1, 13):
        payload = legacy.pnl(
            month=month,
            year=year,
            me=me,
        )

        legacy_totals[month] = payload
        ingresos[month] = (
            _revenue_for_brand(
                payload,
                brand,
            )
            + _pnl_v62_classified_income(
                year,
                month,
                brand,
            )
        )

    account_rows = []

    for code, label, group in SIMPLE_ACCOUNTS:
        values = {}

        for month in range(1, 13):

            classified_value = float(
                classified.get(
                    (month, code),
                    0,
                )
                or 0
            )

            recurring_value = float(
                recurring.get(
                    (month, code),
                    0,
                )
                or 0
            )

            # Mientras no exista reconciliacion
            # accrual/cash por recurrente:
            # si ya existe gasto clasificado real
            # para la cuenta/unidad/mes,
            # ese gasto prevalece.
            #
            # Esto evita sumar dos veces
            # recurrente + pago bancario.
            # V6.4:
            # un pago parcial no debe borrar
            # la base recurrente completa.
            #
            # Mientras no exista matching
            # recurrente<->movimiento uno a uno,
            # mantenemos como P&L el mayor:
            # actual clasificado vs base recurrente.
            values[month] = max(
                classified_value,
                recurring_value,
            )

        account_rows.append(
            {
                "key": code,
                "label": label,
                "group": group,
                "kind": "expense",
                "values": values,
                "total": sum(values.values()),
            }
        )

    for month in range(1, 13):
        simple_total = sum(
            row["values"][month]
            for row in account_rows
        )

        if brand == "CONSOLIDADO":
            old_expenses = float(
                (legacy_totals[month].get("totals") or {}).get("gastos") or 0
            )

            residual[month] = max(old_expenses - simple_total, 0.0)
            total_gastos[month] = max(old_expenses, simple_total)
        else:
            residual[month] = 0.0
            total_gastos[month] = simple_total

        resultado[month] = ingresos[month] - total_gastos[month]

    rows.append(
        {
            "key": "INGRESOS",
            "label": "VENTAS / INGRESOS",
            "group": "INGRESOS",
            "kind": "revenue",
            "values": ingresos,
            "total": sum(ingresos.values()),
        }
    )

    rows.extend(account_rows)

    rows.append(
        {
            "key": "RESIDUAL",
            "label": "Otros gastos CRM / aún no clasificados",
            "group": "OTROS",
            "kind": "expense",
            "values": residual,
            "total": sum(residual.values()),
        }
    )

    rows.append(
        {
            "key": "TOTAL_GASTOS",
            "label": "TOTAL GASTOS",
            "group": "TOTAL",
            "kind": "total",
            "values": total_gastos,
            "total": sum(total_gastos.values()),
        }
    )

    rows.append(
        {
            "key": "RESULTADO",
            "label": "RESULTADO",
            "group": "TOTAL",
            "kind": "result",
            "values": resultado,
            "total": sum(resultado.values()),
        }
    )

    note = (
        "Consolidado conserva el total del P&L actual. "
        "Las cuentas simples muestran los gastos ya clasificados."
        if brand == "CONSOLIDADO"
        else (
            "Por marca, los gastos muestran movimientos que ya fueron "
            "clasificados con esa marca. La cobertura histórica aumentará "
            "a medida que se clasifiquen cartolas."
        )
    )

    return {
        "ok": True,
        "year": year,
        "brand": brand,
        "months": list(range(1, 13)),
        "rows": rows,
        "note": note,
    }



# ============================================================
# FINANCE P&L V2
# Dashboard-compatible commercial source + management P&L.
# ============================================================

FIN_PNL_GROUPS = (
    "COSTO EVENTOS",
    "OPERACIÓN",
    "PERSONAL",
    "GASTOS FIJOS",
    "MARKETING",
    "ADMINISTRACIÓN",
    "FINANCIERO",
    "OTROS",
)


def _fin_month_bounds(
    year: int,
    month: int,
) -> tuple[date, date]:
    start = date(
        int(year),
        int(month),
        1,
    )

    if int(month) == 12:
        end = date(
            int(year) + 1,
            1,
            1,
        )
    else:
        end = date(
            int(year),
            int(month) + 1,
            1,
        )

    return start, end


def _fin_num(value: Any) -> float:
    try:
        return float(
            value or 0
        )
    except Exception:
        return 0.0


def _fin_sales_status(
    actual: float,
    base: float,
    target: float,
) -> str:
    if base <= 0 and target <= 0:
        return "NO_BASE"

    if target > 0 and actual >= target:
        return "GREEN"

    if base > 0 and actual >= base:
        return "BLUE"

    return "RED"


def _fin_expense_status(
    pct_value: float | None,
    ideal_pct: float,
    limit_pct: float,
) -> str:
    if pct_value is None:
        return "NO_DATA"

    if pct_value <= ideal_pct:
        return "GREEN"

    if pct_value <= limit_pct:
        return "BLUE"

    return "RED"


def _fin_confirmed_state_id(db) -> int:
    from backend.routers import tools as dashboard_tools

    return int(
        dashboard_tools._estado_id(
            db,
            "CONFIRM",
        )
        or 0
    )


def _fin_sales_snapshot(
    db,
    *,
    year: int,
    month: int,
) -> dict:
    confirmado_id = _fin_confirmed_state_id(
        db
    )

    if not confirmado_id:
        raise HTTPException(
            status_code=409,
            detail=(
                "No existe estado CONFIRMADO "
                "para calcular ventas."
            ),
        )

    start, end = _fin_month_bounds(
        year,
        month,
    )

    today_cl = datetime.now(
        ZoneInfo(
            "America/Santiago"
        )
    ).date()

    row = db.execute(
        text(
            """
            SELECT
              COALESCE(
                SUM(
                  COALESCE(
                    l.monto_cotizado,
                    0
                  )
                ),
                0
              )::bigint
                AS confirmed_total,

              COALESCE(
                SUM(
                  CASE
                    WHEN l.fecha_evento::date
                         <= :today
                    THEN COALESCE(
                      l.monto_cotizado,
                      0
                    )
                    ELSE 0
                  END
                ),
                0
              )::bigint
                AS realized_total,

              COALESCE(
                SUM(
                  COALESCE(
                    c.traslado,
                    0
                  )
                ),
                0
              )::bigint
                AS traslado,

              COUNT(*)::int
                AS cnt_total

            FROM public.leads l

            LEFT JOIN
              public.cotizaciones c
              ON c.id_cotizacion =
                 l.id_cotizacion_vigente

            WHERE l.id_estado=:conf
              AND l.fecha_evento::date
                  >= :start
              AND l.fecha_evento::date
                  < :end
            """
        ),
        {
            "conf": confirmado_id,
            "start": start,
            "end": end,
            "today": today_cl,
        },
    ).mappings().first() or {}

    by_brand_rows = db.execute(
        text(
            """
            SELECT
              m.id_marca,

              UPPER(
                COALESCE(
                  NULLIF(
                    btrim(
                      COALESCE(
                        m.nombre,
                        m.marca,
                        ''
                      )
                    ),
                    ''
                  ),
                  'SIN MARCA'
                )
              ) AS marca,

              COALESCE(
                SUM(
                  COALESCE(
                    l.monto_cotizado,
                    0
                  )
                ),
                0
              )::bigint
                AS confirmed_total,

              COALESCE(
                SUM(
                  CASE
                    WHEN l.fecha_evento::date
                         <= :today
                    THEN COALESCE(
                      l.monto_cotizado,
                      0
                    )
                    ELSE 0
                  END
                ),
                0
              )::bigint
                AS realized_total,

              COALESCE(
                SUM(
                  COALESCE(
                    c.traslado,
                    0
                  )
                ),
                0
              )::bigint
                AS traslado,

              COUNT(*)::int
                AS cnt_total

            FROM public.leads l

            LEFT JOIN
              public.cotizaciones c
              ON c.id_cotizacion =
                 l.id_cotizacion_vigente

            LEFT JOIN
              public.marcas m
              ON m.id_marca =
                 l.id_marca

            WHERE l.id_estado=:conf
              AND l.fecha_evento::date
                  >= :start
              AND l.fecha_evento::date
                  < :end

            GROUP BY
              m.id_marca,
              UPPER(
                COALESCE(
                  NULLIF(
                    btrim(
                      COALESCE(
                        m.nombre,
                        m.marca,
                        ''
                      )
                    ),
                    ''
                  ),
                  'SIN MARCA'
                )
              )
            """
        ),
        {
            "conf": confirmado_id,
            "start": start,
            "end": end,
            "today": today_cl,
        },
    ).mappings().all()

    target_total = db.execute(
        text(
            """
            SELECT
              COALESCE(
                SUM(
                  COALESCE(
                    ly.venta_base,
                    0
                  )
                ),
                0
              )::bigint
                AS base,

              COALESCE(
                SUM(
                  COALESCE(
                    NULLIF(
                      ly.meta,
                      0
                    ),
                    (
                      COALESCE(
                        ly.venta_base,
                        0
                      )
                      *
                      (
                        1.0
                        +
                        (
                          COALESCE(
                            NULLIF(
                              ly.crecimiento_pct,
                              0
                            ),
                            12
                          )
                          / 100.0
                        )
                      )
                    )
                  )
                ),
                0
              )::bigint
                AS target

            FROM public.metas_marca_mensual ly

            WHERE ly.year=:ly_year
              AND ly.month=:month
            """
        ),
        {
            "ly_year":
                int(year) - 1,
            "month":
                int(month),
        },
    ).mappings().first() or {}

    target_rows = db.execute(
        text(
            """
            SELECT
              ma.id_marca,

              UPPER(
                COALESCE(
                  ma.nombre,
                  ma.marca,
                  ''
                )
              ) AS marca,

              COALESCE(
                ly.venta_base,
                0
              )::float
                AS venta_base,

              COALESCE(
                NULLIF(
                  ly.crecimiento_pct,
                  0
                ),
                12
              )::float
                AS crecimiento_pct,

              COALESCE(
                NULLIF(
                  ly.meta,
                  0
                ),
                (
                  COALESCE(
                    ly.venta_base,
                    0
                  )
                  *
                  (
                    1.0
                    +
                    (
                      COALESCE(
                        NULLIF(
                          ly.crecimiento_pct,
                          0
                        ),
                        12
                      )
                      / 100.0
                    )
                  )
                )
              )::float
                AS target

            FROM public.marcas ma

            LEFT JOIN
              public.metas_marca_mensual ly
              ON ly.id_marca=
                 ma.id_marca
             AND ly.year=:ly_year
             AND ly.month=:month

            ORDER BY ma.id_marca
            """
        ),
        {
            "ly_year":
                int(year) - 1,
            "month":
                int(month),
        },
    ).mappings().all()

    sales_map = {
        int(r["id_marca"]):
            dict(r)
        for r in by_brand_rows
        if r["id_marca"] is not None
    }

    brands = []

    for target_row in target_rows:
        brand_id = int(
            target_row["id_marca"]
        )

        sold = sales_map.get(
            brand_id,
            {},
        )

        confirmed = int(
            sold.get(
                "confirmed_total"
            )
            or 0
        )

        realized = int(
            sold.get(
                "realized_total"
            )
            or 0
        )

        traslado = int(
            sold.get("traslado")
            or 0
        )

        base = _fin_num(
            target_row[
                "venta_base"
            ]
        )

        target = _fin_num(
            target_row[
                "target"
            ]
        )

        status = _fin_sales_status(
            confirmed,
            base,
            target,
        )

        brands.append(
            {
                "id_marca":
                    brand_id,

                "marca":
                    str(
                        target_row[
                            "marca"
                        ]
                        or ""
                    ),

                "venta_base":
                    base,

                "crecimiento_pct":
                    _fin_num(
                        target_row[
                            "crecimiento_pct"
                        ]
                    ),

                "meta":
                    target,

                "realizada":
                    realized,

                "confirmada":
                    confirmed,

                "traslado":
                    traslado,

                "neto_productos":
                    confirmed
                    - traslado,

                "cnt_total":
                    int(
                        sold.get(
                            "cnt_total"
                        )
                        or 0
                    ),

                "gap_meta":
                    target
                    - confirmed,

                "diff_base":
                    confirmed
                    - base,

                "cumplimiento_pct":
                    (
                        confirmed
                        / target
                        * 100.0
                        if target > 0
                        else None
                    ),

                "vs_base_pct":
                    (
                        confirmed
                        / base
                        * 100.0
                        if base > 0
                        else None
                    ),

                "status":
                    status,
            }
        )

    confirmed_total = int(
        row.get(
            "confirmed_total"
        )
        or 0
    )

    realized_total = int(
        row.get(
            "realized_total"
        )
        or 0
    )

    traslado_total = int(
        row.get(
            "traslado"
        )
        or 0
    )

    base_total = int(
        target_total.get("base")
        or 0
    )

    target_value = int(
        target_total.get("target")
        or 0
    )

    return {
        "year":
            int(year),

        "month":
            int(month),

        "confirmed":
            confirmed_total,

        "realized":
            realized_total,

        "traslado":
            traslado_total,

        "neto_productos":
            confirmed_total
            - traslado_total,

        "cnt_total":
            int(
                row.get(
                    "cnt_total"
                )
                or 0
            ),

        "base":
            base_total,

        "target":
            target_value,

        "gap_target":
            target_value
            - confirmed_total,

        "diff_base":
            confirmed_total
            - base_total,

        "attainment_pct":
            (
                confirmed_total
                / target_value
                * 100.0
                if target_value > 0
                else None
            ),

        "vs_base_pct":
            (
                confirmed_total
                / base_total
                * 100.0
                if base_total > 0
                else None
            ),

        "status":
            _fin_sales_status(
                confirmed_total,
                base_total,
                target_value,
            ),

        "by_brand":
            brands,
    }


def _fin_dashboard_current(
    db,
) -> dict:
    from backend.routers import tools as dashboard_tools

    payload = dashboard_tools.dashboard_v2(
        week_offset=0,
        month_offset=0,
        id_marca=None,
        db=db,
        me={
            "role": "SUPERADMIN",
            "rol": "SUPERADMIN",
        },
    )

    return (
        payload.get("sales")
        or {}
    )


def _fin_require_dashboard_parity(
    db,
    snapshot: dict,
) -> dict:
    today_cl = datetime.now(
        ZoneInfo(
            "America/Santiago"
        )
    ).date()

    if (
        int(snapshot["year"])
        != today_cl.year
        or
        int(snapshot["month"])
        != today_cl.month
    ):
        return {
            "checked": False,
            "ok": True,
            "reason":
                "historical_month",
        }

    sale = _fin_dashboard_current(
        db
    )

    dashboard_total = int(
        sale.get(
            "venta_mes_total"
        )
        or 0
    )

    dashboard_neto = int(
        sale.get(
            "venta_mes_neta"
        )
        or 0
    )

    dashboard_traslado = int(
        sale.get(
            "venta_mes_traslado"
        )
        or 0
    )

    dashboard_base = int(
        sale.get(
            "base_mes_total"
        )
        or 0
    )

    dashboard_target = int(
        sale.get(
            "meta_mes_total"
        )
        or 0
    )

    checks = {
        "venta":
            (
                int(
                    snapshot[
                        "confirmed"
                    ]
                )
                == dashboard_total
            ),

        "traslado":
            (
                int(
                    snapshot[
                        "traslado"
                    ]
                )
                == dashboard_traslado
            ),

        "neto":
            (
                int(
                    snapshot[
                        "neto_productos"
                    ]
                )
                == dashboard_neto
            ),

        "base":
            (
                int(
                    snapshot[
                        "base"
                    ]
                )
                == dashboard_base
            ),

        "meta":
            (
                int(
                    snapshot[
                        "target"
                    ]
                )
                == dashboard_target
            ),
    }

    if not all(
        checks.values()
    ):
        raise HTTPException(
            status_code=409,
            detail={
                "message":
                    (
                        "P&L no coincide con "
                        "Dashboard. No se "
                        "mostrarán ventas."
                    ),

                "checks":
                    checks,

                "pnl": {
                    "venta":
                        snapshot[
                            "confirmed"
                        ],
                    "neto":
                        snapshot[
                            "neto_productos"
                        ],
                    "traslado":
                        snapshot[
                            "traslado"
                        ],
                    "base":
                        snapshot[
                            "base"
                        ],
                    "meta":
                        snapshot[
                            "target"
                        ],
                },

                "dashboard": {
                    "venta":
                        dashboard_total,
                    "neto":
                        dashboard_neto,
                    "traslado":
                        dashboard_traslado,
                    "base":
                        dashboard_base,
                    "meta":
                        dashboard_target,
                },
            },
        )

    return {
        "checked": True,
        "ok": True,
        "checks":
            checks,

        "dashboard": {
            "venta":
                dashboard_total,
            "neto":
                dashboard_neto,
            "traslado":
                dashboard_traslado,
            "base":
                dashboard_base,
            "meta":
                dashboard_target,
        },
    }


def _fin_account_config(
    db,
):
    return db.execute(
        text(
            """
            SELECT
              c.cuenta_code
                AS code,

              COALESCE(
                NULLIF(
                  btrim(
                    c.label_override
                  ),
                  ''
                ),
                p.name,
                c.cuenta_code
              ) AS label,

              c.group_name,
              c.group_order,
              c.display_order,
              c.ideal_pct,
              c.limit_pct

            FROM public.fin_pnl_account_config c

            JOIN public.plan_cuentas p
              ON p.code =
                 c.cuenta_code

            WHERE c.is_active=TRUE
              AND p.type='expense'
              AND COALESCE(
                p.is_active,
                TRUE
              )=TRUE

            ORDER BY
              c.group_order,
              c.display_order,
              c.cuenta_code
            """
        )
    ).mappings().all()


@router.get(
    "/classify-meta-v2"
)
def classify_meta_v2(
    me: dict = Depends(
        legacy.get_current_user
    ),
    db=Depends(get_db),
):
    _role_classify(me)

    accounts = list(
        _fin_account_config(
            db
        )
    )

    # FINANCE_UX_CASHFLOW_V3
    # Responsable del gasto NO es cualquier usuario activo.
    # Sólo roles/cargos expresamente autorizados.
    users = db.execute(
        text(
            """
            SELECT
              id_usuario,
              nombre,
              cargo,
              rol

            FROM public.usuarios u

            WHERE u.is_active=TRUE
              AND (
                regexp_replace(
                  upper(
                    COALESCE(
                      u.rol,
                      ''
                    )
                  ),
                  '[^A-Z0-9]+',
                  '',
                  'g'
                ) IN (
                  'COMPRAS',
                  'JEFEDEOPERACIONES',
                  'BODEGUERO',
                  'SUPERADMIN',
                  'ADMIN'
                )

                OR

                regexp_replace(
                  upper(
                    COALESCE(
                      u.cargo,
                      ''
                    )
                  ),
                  '[^A-Z0-9]+',
                  '',
                  'g'
                ) IN (
                  'COMPRAS',
                  'JEFEDEOPERACIONES',
                  'BODEGUERO',
                  'SUPERADMIN',
                  'ADMIN'
                )
              )

            ORDER BY
              CASE
                WHEN regexp_replace(
                       upper(
                         COALESCE(
                           u.rol,
                           u.cargo,
                           ''
                         )
                       ),
                       '[^A-Z0-9]+',
                       '',
                       'g'
                     )='SUPERADMIN'
                  THEN 1

                WHEN regexp_replace(
                       upper(
                         COALESCE(
                           u.rol,
                           u.cargo,
                           ''
                         )
                       ),
                       '[^A-Z0-9]+',
                       '',
                       'g'
                     )='ADMIN'
                  THEN 2

                WHEN regexp_replace(
                       upper(
                         COALESCE(
                           u.rol,
                           u.cargo,
                           ''
                         )
                       ),
                       '[^A-Z0-9]+',
                       '',
                       'g'
                     )='COMPRAS'
                  THEN 3

                WHEN regexp_replace(
                       upper(
                         COALESCE(
                           u.rol,
                           u.cargo,
                           ''
                         )
                       ),
                       '[^A-Z0-9]+',
                       '',
                       'g'
                     )='JEFEDEOPERACIONES'
                  THEN 4

                WHEN regexp_replace(
                       upper(
                         COALESCE(
                           u.rol,
                           u.cargo,
                           ''
                         )
                       ),
                       '[^A-Z0-9]+',
                       '',
                       'g'
                     )='BODEGUERO'
                  THEN 5

                ELSE 99
              END,
              nombre,
              id_usuario
            """
        )
    ).mappings().all()

    return {
        "ok": True,

        "accounts": [
            {
                "code":
                    str(
                        r["code"]
                    ),

                "label":
                    str(
                        r["label"]
                    ),

                "group":
                    str(
                        r[
                            "group_name"
                        ]
                    ),

                "ideal_pct":
                    _fin_num(
                        r[
                            "ideal_pct"
                        ]
                    ),

                "limit_pct":
                    _fin_num(
                        r[
                            "limit_pct"
                        ]
                    ),
            }
            for r in accounts
        ],

        "users": [
            {
                "id_usuario":
                    int(
                        r[
                            "id_usuario"
                        ]
                    ),

                "nombre":
                    str(
                        r[
                            "nombre"
                        ]
                    ),

                "cargo":
                    str(
                        r["cargo"]
                        or ""
                    ),
                "rol":
                    str(
                        r["rol"]
                        or ""
                    ),
            }
            for r in users
        ],
    }



@router.get(
    "/movements/selection-detail"
)
def movement_selection_detail_v3(
    ids: str = Query(
        ...,
        min_length=1,
        max_length=3000,
    ),
    me: dict = Depends(
        legacy.get_current_user
    ),
    db=Depends(get_db),
):
    _role_classify(me)

    raw_ids = []

    for piece in str(ids).split(","):
        piece = piece.strip()

        if not piece:
            continue

        try:
            value = int(piece)
        except Exception as exc:
            raise HTTPException(
                status_code=400,
                detail="ID de movimiento inválido.",
            ) from exc

        if value <= 0:
            raise HTTPException(
                status_code=400,
                detail="ID de movimiento inválido.",
            )

        raw_ids.append(value)

    movement_ids = sorted(
        set(raw_ids)
    )

    if not movement_ids:
        raise HTTPException(
            status_code=400,
            detail="Selecciona movimientos.",
        )

    if len(movement_ids) > 100:
        raise HTTPException(
            status_code=400,
            detail="Máximo 100 movimientos.",
        )

    params = {}

    placeholders = []

    for index, movement_id in enumerate(
        movement_ids
    ):
        key = f"id_{index}"

        params[key] = movement_id

        placeholders.append(
            ":" + key
        )

    sql_ids = ", ".join(
        placeholders
    )

    rows = db.execute(
        text(
            f"""
            SELECT
              bm.id_bank_movement,
              bm.tx_date,
              bm.description,
              bm.amount,
              bm.reference,
              bm.balance,
              bm.status,
              bm.cuenta_code,
              bm.descripcion_interna,

              ba.bank_name,
              ba.label AS bank_account

            FROM public.fin_bank_movements bm

            LEFT JOIN public.fin_bank_accounts ba
              ON ba.id_bank_account=
                 bm.id_bank_account

            WHERE bm.id_bank_movement
                  IN ({sql_ids})

            ORDER BY
              bm.tx_date,
              bm.id_bank_movement
            """
        ),
        params,
    ).mappings().all()

    return {
        "ok": True,
        "items": [
            dict(row)
            for row in rows
        ],
    }


@router.post(
    "/movements/classify-v2"
)
def classify_movements_v2(
    body: dict = Body(...),
    me: dict = Depends(
        legacy.get_current_user
    ),
):
    _role_classify(me)

    ids = body.get(
        "movement_ids"
    ) or []

    code = str(
        body.get(
            "cuenta_code"
        )
        or ""
    ).strip()

    description_internal = str(
        body.get(
            "descripcion_interna"
        )
        or ""
    ).strip()

    try:
        responsible_id = int(
            body.get(
                "responsable_id"
            )
        )
    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail=(
                "Selecciona al responsable "
                "del gasto."
            ),
        ) from exc

    if (
        not isinstance(
            ids,
            list,
        )
        or not ids
    ):
        raise HTTPException(
            status_code=400,
            detail=(
                "Selecciona movimientos."
            ),
        )

    if not code:
        raise HTTPException(
            status_code=400,
            detail=(
                "Selecciona una cuenta."
            ),
        )

    if not description_internal:
        raise HTTPException(
            status_code=400,
            detail=(
                "Describe qué se compró "
                "o a qué corresponde."
            ),
        )

    actor = _actor(me)

    with legacy.get_connection() as conn:

        account_ok = conn.execute(
            text(
                """
                SELECT 1
                FROM public.fin_pnl_account_config c

                JOIN public.plan_cuentas p
                  ON p.code=
                     c.cuenta_code

                WHERE c.cuenta_code=:code
                  AND c.is_active=TRUE
                  AND p.type='expense'
                  AND COALESCE(
                    p.is_active,
                    TRUE
                  )=TRUE

                LIMIT 1
                """
            ),
            {"code": code},
        ).scalar()

        if not account_ok:
            raise HTTPException(
                status_code=400,
                detail=(
                    "Cuenta inválida "
                    "o inactiva."
                ),
            )

        user_ok = conn.execute(
            text(
                """
                SELECT 1

                FROM public.usuarios u

                WHERE u.id_usuario=:id
                  AND u.is_active=TRUE

                  AND (
                    regexp_replace(
                      upper(
                        COALESCE(
                          u.rol,
                          ''
                        )
                      ),
                      '[^A-Z0-9]+',
                      '',
                      'g'
                    ) IN (
                      'COMPRAS',
                      'JEFEDEOPERACIONES',
                      'BODEGUERO',
                      'SUPERADMIN',
                      'ADMIN'
                    )

                    OR

                    regexp_replace(
                      upper(
                        COALESCE(
                          u.cargo,
                          ''
                        )
                      ),
                      '[^A-Z0-9]+',
                      '',
                      'g'
                    ) IN (
                      'COMPRAS',
                      'JEFEDEOPERACIONES',
                      'BODEGUERO',
                      'SUPERADMIN',
                      'ADMIN'
                    )
                  )

                LIMIT 1
                """
            ),
            {
                "id":
                    responsible_id
            },
        ).scalar()

        if not user_ok:
            raise HTTPException(
                status_code=400,
                detail=(
                    "Responsable inválido."
                ),
            )

        classified = 0

        for raw_id in ids:
            movement_id = int(
                raw_id
            )

            movement = conn.execute(
                text(
                    """
                    SELECT *
                    FROM public.fin_bank_movements
                    WHERE id_bank_movement=:id
                    FOR UPDATE
                    """
                ),
                {
                    "id":
                        movement_id
                },
            ).mappings().first()

            if not movement:
                continue


            confirmed_allocation = conn.execute(
                text(
                    """
                    SELECT 1
                    FROM public.fin_payable_bank_allocations
                    WHERE bank_movement_id=:id
                      AND status='CONFIRMED'
                    LIMIT 1
                    """
                ),
                {
                    "id":
                        movement_id
                },
            ).scalar()

            if confirmed_allocation:
                raise HTTPException(
                    status_code=409,
                    detail={
                        "code":
                            "BANK_MOVEMENT_ALREADY_RECONCILED",

                        "message":
                            "Este movimiento ya está conciliado "
                            "con una factura. No se puede "
                            "clasificar nuevamente como gasto.",
                    },
                )

            amount = Decimal(
                str(
                    movement[
                        "amount"
                    ]
                )
            )

            if amount >= 0:
                raise HTTPException(
                    status_code=400,
                    detail=(
                        "Los abonos no se "
                        "imputan automáticamente "
                        "como gasto."
                    ),
                )

            expense_amount = abs(
                amount
            )

            conn.execute(
                text(
                    """
                    INSERT INTO
                    public.fin_gastos(
                      fecha,
                      cuenta_code,
                      monto,
                      descripcion,
                      descripcion_interna,
                      proveedor,
                      marca,
                      responsable_id,
                      centro_costo,
                      created_at,
                      doc_num,
                      pagado,
                      fecha_pago,
                      tipo_doc,
                      is_active,
                      bank_movement_id
                    )
                    VALUES(
                      :fecha,
                      :cuenta,
                      :monto,
                      :descripcion,
                      :descripcion_interna,
                      :proveedor,
                      NULL,
                      :responsable,
                      NULL,
                      now(),
                      :doc_num,
                      TRUE,
                      :fecha_pago,
                      'CARTOLA',
                      TRUE,
                      :movement
                    )
                    ON CONFLICT(
                      bank_movement_id
                    )
                    DO UPDATE SET
                      cuenta_code=
                        EXCLUDED.cuenta_code,
                      monto=
                        EXCLUDED.monto,
                      descripcion=
                        EXCLUDED.descripcion,
                      descripcion_interna=
                        EXCLUDED.descripcion_interna,
                      proveedor=
                        EXCLUDED.proveedor,
                      responsable_id=
                        EXCLUDED.responsable_id,
                      fecha=
                        EXCLUDED.fecha,
                      pagado=TRUE,
                      fecha_pago=
                        EXCLUDED.fecha_pago,
                      is_active=TRUE
                    """
                ),
                {
                    "fecha":
                        movement[
                            "tx_date"
                        ],

                    "cuenta":
                        code,

                    "monto":
                        expense_amount,

                    "descripcion":
                        movement[
                            "description"
                        ],

                    "descripcion_interna":
                        description_internal,

                    "proveedor":
                        str(
                            movement[
                                "description"
                            ]
                        )[:200],

                    "responsable":
                        responsible_id,

                    "doc_num":
                        movement[
                            "reference"
                        ],

                    "fecha_pago":
                        movement[
                            "tx_date"
                        ],

                    "movement":
                        movement_id,
                },
            )

            conn.execute(
                text(
                    """
                    UPDATE
                      public.fin_bank_movements

                    SET
                      status='CLASSIFIED',
                      cuenta_code=:cuenta,
                      marca=NULL,
                      responsable_id=
                        :responsable,
                      descripcion_interna=
                        :descripcion_interna,
                      classified_by=:actor,
                      classified_at=now()

                    WHERE id_bank_movement=:id
                    """
                ),
                {
                    "cuenta":
                        code,

                    "responsable":
                        responsible_id,

                    "descripcion_interna":
                        description_internal,

                    "actor":
                        actor,

                    "id":
                        movement_id,
                },
            )

            classified += 1

        conn.commit()

    return {
        "ok": True,
        "classified":
            classified,
    }


@router.post(
    "/pnl-v2/accounts"
)
def create_pnl_account_v2(
    body: dict = Body(...),
    me: dict = Depends(
        legacy.get_current_user
    ),
):
    _role_upload(me)

    name = str(
        body.get("name")
        or ""
    ).strip()

    group_name = str(
        body.get("group")
        or ""
    ).strip().upper()

    try:
        ideal_pct = float(
            body.get(
                "ideal_pct"
            )
            or 0
        )

        limit_pct = float(
            body.get(
                "limit_pct"
            )
            or 0
        )

    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail=(
                "Porcentajes inválidos."
            ),
        ) from exc

    group_map = {
        _norm(group):
            group
        for group
        in FIN_PNL_GROUPS
    }

    group_name = group_map.get(
        _norm(group_name)
    )

    if not name:
        raise HTTPException(
            status_code=400,
            detail="Nombre requerido.",
        )

    if not group_name:
        raise HTTPException(
            status_code=400,
            detail="Grupo P&L inválido.",
        )

    if (
        ideal_pct < 0
        or
        limit_pct < ideal_pct
        or
        limit_pct > 100
    ):
        raise HTTPException(
            status_code=400,
            detail=(
                "Límite inválido."
            ),
        )

    with legacy.get_connection() as conn:

        used = {
            str(value)
            for value
            in conn.execute(
                text(
                    """
                    SELECT code
                    FROM public.plan_cuentas
                    """
                )
            ).scalars().all()
        }

        code = None

        for candidate in range(
            8900,
            9000,
        ):
            if str(candidate) not in used:
                code = str(
                    candidate
                )
                break

        if not code:
            raise HTTPException(
                status_code=409,
                detail=(
                    "No quedan códigos "
                    "automáticos."
                ),
            )

        order_row = conn.execute(
            text(
                """
                SELECT
                  group_order,
                  COALESCE(
                    MAX(display_order),
                    0
                  ) + 10 AS next_order
                FROM public.fin_pnl_account_config
                WHERE group_name=:group
                GROUP BY group_order
                LIMIT 1
                """
            ),
            {
                "group":
                    group_name
            },
        ).mappings().first()

        if order_row:
            group_order = int(
                order_row[
                    "group_order"
                ]
            )
            display_order = int(
                order_row[
                    "next_order"
                ]
            )
        else:
            group_order = (
                FIN_PNL_GROUPS.index(
                    group_name
                )
                + 1
            )
            display_order = 10

        conn.execute(
            text(
                """
                INSERT INTO
                public.plan_cuentas(
                  code,
                  name,
                  type,
                  classification,
                  description,
                  who_inputs,
                  how_to_impute,
                  parent_code,
                  is_active
                )
                VALUES(
                  :code,
                  :name,
                  'expense',
                  :group_name,
                  'Creada desde P&L V2',
                  'FINANZAS',
                  'Cartola / devengado',
                  NULL,
                  TRUE
                )
                """
            ),
            {
                "code":
                    code,

                "name":
                    name,

                "group_name":
                    group_name,
            },
        )

        conn.execute(
            text(
                """
                INSERT INTO
                public.fin_pnl_account_config(
                  cuenta_code,
                  group_name,
                  group_order,
                  display_order,
                  ideal_pct,
                  limit_pct,
                  is_active
                )
                VALUES(
                  :code,
                  :group_name,
                  :group_order,
                  :display_order,
                  :ideal,
                  :limit,
                  TRUE
                )
                """
            ),
            {
                "code":
                    code,

                "group_name":
                    group_name,

                "group_order":
                    group_order,

                "display_order":
                    display_order,

                "ideal":
                    ideal_pct,

                "limit":
                    limit_pct,
            },
        )

        conn.commit()

    return {
        "ok": True,
        "code": code,
        "name": name,
        "group": group_name,
    }


@router.get(
    "/pnl-v2/management"
)
def pnl_management_v2(
    year: int = Query(
        ...,
        ge=2020,
        le=2100,
    ),
    month: int = Query(
        ...,
        ge=1,
        le=12,
    ),
    me: dict = Depends(
        legacy.get_current_user
    ),
    db=Depends(get_db),
):
    _role_pnl(me)

    sales_months = {}

    for current_month in range(
        1,
        13,
    ):
        sales_months[
            current_month
        ] = _fin_sales_snapshot(
            db,
            year=year,
            month=current_month,
        )

    selected_sales = sales_months[
        int(month)
    ]

    parity = (
        _fin_require_dashboard_parity(
            db,
            selected_sales,
        )
    )

    account_rows = list(
        _fin_account_config(
            db
        )
    )

    group_rows = db.execute(
        text(
            """
            SELECT
              group_name,
              group_order,
              ideal_pct,
              limit_pct
            FROM public.fin_pnl_group_config
            WHERE is_active=TRUE
            ORDER BY group_order
            """
        )
    ).mappings().all()

    # FIN_EXACT_FIXED_BASE_V1
    expense_rows = db.execute(
        text(
            """
            WITH pnl_expenses AS (

              SELECT
                EXTRACT(
                  MONTH FROM g.fecha
                )::int AS month,

                g.cuenta_code,

                COALESCE(
                  g.monto,
                  0
                )::numeric AS amount

              FROM public.fin_gastos g

              JOIN public.fin_pnl_account_config cfg
                ON cfg.cuenta_code=
                   g.cuenta_code

              WHERE g.is_active=TRUE
                AND cfg.is_active=TRUE
                AND EXTRACT(
                      YEAR FROM g.fecha
                    )::int=:year

              UNION ALL

              SELECT
                fm.month::int AS month,

                ft.cuenta_code,

                fm.amount::numeric AS amount

              FROM
                public.fin_recurring_expense_months fm

              JOIN
                public.fin_recurring_expense_templates ft
                ON ft.id_recurring_template=
                   fm.id_recurring_template

              JOIN public.fin_pnl_account_config cfg
                ON cfg.cuenta_code=
                   ft.cuenta_code

              WHERE fm.year=:year
                AND fm.is_active=TRUE
                AND ft.is_active=TRUE
                AND cfg.is_active=TRUE
            )

            SELECT
              month,
              cuenta_code,
              SUM(amount)::numeric AS total

            FROM pnl_expenses

            GROUP BY
              month,
              cuenta_code
            """
        ),
        {
            "year":
                int(year)
        },
    ).mappings().all()

    expense_map = {
        (
            int(r["month"]),
            str(
                r[
                    "cuenta_code"
                ]
            ),
        ):
            _fin_num(
                r["total"]
            )
        for r in expense_rows
    }

    accounts = []

    for config in account_rows:
        code = str(
            config["code"]
        )

        ideal = _fin_num(
            config[
                "ideal_pct"
            ]
        )

        limit_pct = _fin_num(
            config[
                "limit_pct"
            ]
        )

        values = {}
        percentages = {}
        statuses = {}

        for current_month in range(
            1,
            13,
        ):
            amount = _fin_num(
                expense_map.get(
                    (
                        current_month,
                        code,
                    )
                )
            )

            sales_basis = _fin_num(
                sales_months[
                    current_month
                ]["confirmed"]
            )

            pct_value = (
                amount
                / sales_basis
                * 100.0
                if sales_basis > 0
                else None
            )

            values[
                current_month
            ] = amount

            percentages[
                current_month
            ] = pct_value

            statuses[
                current_month
            ] = _fin_expense_status(
                pct_value,
                ideal,
                limit_pct,
            )

        accounts.append(
            {
                "code":
                    code,

                "label":
                    str(
                        config[
                            "label"
                        ]
                    ),

                "group":
                    str(
                        config[
                            "group_name"
                        ]
                    ),

                "ideal_pct":
                    ideal,

                "limit_pct":
                    limit_pct,

                "values":
                    values,

                "percentages":
                    percentages,

                "statuses":
                    statuses,

                "total":
                    sum(
                        values.values()
                    ),
            }
        )

    groups = []

    for group_config in group_rows:
        group_name = str(
            group_config[
                "group_name"
            ]
        )

        children = [
            account
            for account in accounts
            if account["group"]
            == group_name
        ]

        values = {
            current_month:
                sum(
                    child["values"][
                        current_month
                    ]
                    for child in children
                )
            for current_month
            in range(1, 13)
        }

        percentages = {}
        statuses = {}

        ideal = _fin_num(
            group_config[
                "ideal_pct"
            ]
        )

        limit_pct = _fin_num(
            group_config[
                "limit_pct"
            ]
        )

        for current_month in range(
            1,
            13,
        ):
            sales_basis = _fin_num(
                sales_months[
                    current_month
                ]["confirmed"]
            )

            pct_value = (
                values[current_month]
                / sales_basis
                * 100.0
                if sales_basis > 0
                else None
            )

            percentages[
                current_month
            ] = pct_value

            statuses[
                current_month
            ] = _fin_expense_status(
                pct_value,
                ideal,
                limit_pct,
            )

        groups.append(
            {
                "group":
                    group_name,

                "ideal_pct":
                    ideal,

                "limit_pct":
                    limit_pct,

                "values":
                    values,

                "percentages":
                    percentages,

                "statuses":
                    statuses,

                "total":
                    sum(
                        values.values()
                    ),
            }
        )

    expense_totals = {}
    actual_results = {}
    confirmed_results = {}

    for current_month in range(
        1,
        13,
    ):
        expense_total = sum(
            account["values"][
                current_month
            ]
            for account in accounts
        )

        expense_totals[
            current_month
        ] = expense_total

        actual_results[
            current_month
        ] = (
            _fin_num(
                sales_months[
                    current_month
                ]["realized"]
            )
            - expense_total
        )

        confirmed_results[
            current_month
        ] = (
            _fin_num(
                sales_months[
                    current_month
                ]["confirmed"]
            )
            - expense_total
        )

    pending = db.execute(
        text(
            """
            SELECT
              COUNT(*) FILTER(
                WHERE amount < 0
              )::int
                AS count,

              COALESCE(
                SUM(
                  CASE
                    WHEN amount < 0
                    THEN abs(amount)
                    ELSE 0
                  END
                ),
                0
              )::numeric
                AS amount

            FROM
              public.fin_bank_movements

            WHERE status='PENDING'
            """
        )
    ).mappings().first() or {}

    # --------------------------------------------------------
    # FLUJO OPERATIVO AL DIA
    # Venta realizada - base fija - otros gastos clasificados.
    # --------------------------------------------------------

    today_cl = datetime.now(
        ZoneInfo(
            "America/Santiago"
        )
    ).date()

    cf_start, cf_end = _fin_month_bounds(
        year,
        month,
    )

    selected_period = (
        int(year),
        int(month),
    )

    current_period = (
        int(today_cl.year),
        int(today_cl.month),
    )

    if selected_period < current_period:
        cutoff_exclusive = cf_end
        flow_mode = "HISTORICO"

    elif selected_period == current_period:
        cutoff_exclusive = (
            today_cl
            + timedelta(days=1)
        )
        flow_mode = "AL_DIA"

    else:
        cutoff_exclusive = cf_start
        flow_mode = "PROYECCION"

    fixed_base_row = db.execute(
        text(
            """
            SELECT
              COALESCE(
                SUM(
                  fm.amount
                ),
                0
              )::numeric AS total

            FROM
              public.fin_recurring_expense_months fm

            JOIN
              public.fin_recurring_expense_templates ft
              ON ft.id_recurring_template=
                 fm.id_recurring_template

            WHERE fm.year=:year
              AND fm.month=:month
              AND fm.is_active=TRUE
              AND ft.is_active=TRUE
            """
        ),
        {
            "year":
                int(year),

            "month":
                int(month),
        },
    ).scalar()

    fixed_base_total = _fin_num(
        fixed_base_row
    )

    other_expenses_row = db.execute(
        text(
            """
            SELECT
              COALESCE(
                SUM(
                  COALESCE(
                    g.monto,
                    0
                  )
                ),
                0
              )::numeric AS total

            FROM public.fin_gastos g

            WHERE g.is_active=TRUE
              AND g.fecha >= :start
              AND g.fecha < :cutoff

              AND NOT EXISTS (
                SELECT 1

                FROM
                  public.fin_recurring_expense_templates ft

                WHERE ft.is_active=TRUE
                  AND ft.cuenta_code=
                      g.cuenta_code
              )
            """
        ),
        {
            "start":
                cf_start,

            "cutoff":
                cutoff_exclusive,
        },
    ).scalar()

    other_expenses_total = _fin_num(
        other_expenses_row
    )

    fixed_account_movements_row = db.execute(
        text(
            """
            SELECT
              COUNT(*)::int AS count,

              COALESCE(
                SUM(
                  COALESCE(
                    g.monto,
                    0
                  )
                ),
                0
              )::numeric AS total

            FROM public.fin_gastos g

            WHERE g.is_active=TRUE
              AND g.fecha >= :start
              AND g.fecha < :cutoff

              AND EXISTS (
                SELECT 1

                FROM
                  public.fin_recurring_expense_templates ft

                WHERE ft.is_active=TRUE
                  AND ft.cuenta_code=
                      g.cuenta_code
              )
            """
        ),
        {
            "start":
                cf_start,

            "cutoff":
                cutoff_exclusive,
        },
    ).mappings().first() or {}

    pending_period = db.execute(
        text(
            """
            SELECT
              COUNT(*) FILTER(
                WHERE amount < 0
              )::int AS count,

              COALESCE(
                SUM(
                  CASE
                    WHEN amount < 0
                    THEN abs(amount)
                    ELSE 0
                  END
                ),
                0
              )::numeric AS amount

            FROM public.fin_bank_movements

            WHERE status='PENDING'
              AND tx_date >= :start
              AND tx_date < :cutoff
            """
        ),
        {
            "start":
                cf_start,

            "cutoff":
                cutoff_exclusive,
        },
    ).mappings().first() or {}

    sales_to_date = _fin_num(
        selected_sales.get(
            "realized"
        )
    )

    operating_result = (
        sales_to_date
        - fixed_base_total
        - other_expenses_total
    )

    cashflow = {
        "mode":
            flow_mode,

        "venta_al_dia":
            sales_to_date,

        "venta_confirmada_mes":
            _fin_num(
                selected_sales.get(
                    "confirmed"
                )
            ),

        "gastos_fijos":
            fixed_base_total,

        "otros_gastos":
            other_expenses_total,

        "resultado_operativo":
            operating_result,

        "movimientos_cuentas_fijas": {
            "count":
                int(
                    fixed_account_movements_row.get(
                        "count"
                    )
                    or 0
                ),

            "amount":
                _fin_num(
                    fixed_account_movements_row.get(
                        "total"
                    )
                ),
        },

        "pendiente_clasificar": {
            "count":
                int(
                    pending_period.get(
                        "count"
                    )
                    or 0
                ),

            "amount":
                _fin_num(
                    pending_period.get(
                        "amount"
                    )
                ),
        },

        "formula":
            (
                "VENTA_REALIZADA_AL_DIA "
                "- GASTOS_FIJOS "
                "- OTROS_GASTOS_CLASIFICADOS"
            ),

        "is_bank_balance":
            False,
    }


    return {
        "ok": True,

        "year":
            int(year),

        "month":
            int(month),

        "sales_selected":
            selected_sales,

        "cashflow":
            cashflow,

        "sales_by_brand":
            selected_sales[
                "by_brand"
            ],

        "sales_months": {
            str(key): {
                field: value
                for field, value
                in snapshot.items()
                if field != "by_brand"
            }
            for key, snapshot
            in sales_months.items()
        },

        "dashboard_parity":
            parity,

        "accounts":
            accounts,

        "groups":
            groups,

        "expense_totals":
            expense_totals,

        "actual_results":
            actual_results,

        "confirmed_results":
            confirmed_results,

        "selected_expenses":
            expense_totals[
                int(month)
            ],

        "selected_actual_result":
            actual_results[
                int(month)
            ],

        "selected_confirmed_result":
            confirmed_results[
                int(month)
            ],

        "pending": {
            "count":
                int(
                    pending.get(
                        "count"
                    )
                    or 0
                ),

            "amount":
                _fin_num(
                    pending.get(
                        "amount"
                    )
                ),
        },

        "expense_ratio_basis":
            (
                "VENTA_CONFIRMADA_"
                "SIN_IVA_INCLUYE_TRASLADO"
            ),

        "result_warning":
            (
                "Incluye la base fija mensual "
                "ENE-DIC definida por Finanzas. "
                "Los movimientos variables se "
                "suman por separado."
            ),
    }


@router.get(
    "/pnl-v2/detail"
)
def pnl_detail_v2(
    year: int = Query(
        ...,
        ge=2020,
        le=2100,
    ),
    month: int = Query(
        ...,
        ge=1,
        le=12,
    ),
    account: str = Query(
        ...,
        min_length=1,
        max_length=32,
    ),
    me: dict = Depends(
        legacy.get_current_user
    ),
    db=Depends(get_db),
):
    _role_pnl(me)

    start, end = _fin_month_bounds(
        year,
        month,
    )

    rows = db.execute(
        text(
            """
            WITH detail AS (

              SELECT
                'MOVIMIENTO'::text
                  AS source_type,

                g.id_gasto::bigint
                  AS id_item,

                g.fecha::date
                  AS fecha,

                g.fecha_pago::date
                  AS fecha_pago,

                g.monto::numeric
                  AS monto,

                g.cuenta_code,

                pc.name AS cuenta,

                g.descripcion
                  AS descripcion_banco,

                g.descripcion_interna,

                g.proveedor,

                bm.tx_date::date
                  AS fecha_banco,

                bm.description
                  AS movimiento_banco,

                bm.reference
                  AS referencia_banco,

                ba.bank_name,

                ba.label
                  AS bank_account,

                ur.nombre
                  AS responsable,

                uc.nombre
                  AS clasificado_por,

                bm.classified_at

              FROM public.fin_gastos g

              LEFT JOIN public.plan_cuentas pc
                ON pc.code=
                   g.cuenta_code

              LEFT JOIN public.fin_bank_movements bm
                ON bm.id_bank_movement=
                   g.bank_movement_id

              LEFT JOIN public.fin_bank_accounts ba
                ON ba.id_bank_account=
                   bm.id_bank_account

              LEFT JOIN public.usuarios ur
                ON ur.id_usuario=
                   g.responsable_id

              LEFT JOIN public.usuarios uc
                ON uc.id_usuario=
                   bm.classified_by

              WHERE g.is_active=TRUE
                AND g.cuenta_code=:account
                AND g.fecha >= :start
                AND g.fecha < :end

              UNION ALL

              SELECT
                'FIJO_BASE_MENSUAL'::text
                  AS source_type,

                (
                  fm.id_recurring_month
                  * -1
                )::bigint
                  AS id_item,

                fm.economic_date::date
                  AS fecha,

                NULL::date
                  AS fecha_pago,

                fm.amount::numeric
                  AS monto,

                ft.cuenta_code,

                pc.name AS cuenta,

                NULL::text
                  AS descripcion_banco,

                ft.descripcion_interna,

                NULL::text
                  AS proveedor,

                NULL::date
                  AS fecha_banco,

                NULL::text
                  AS movimiento_banco,

                NULL::text
                  AS referencia_banco,

                NULL::text
                  AS bank_name,

                NULL::text
                  AS bank_account,

                NULL::text
                  AS responsable,

                NULL::text
                  AS clasificado_por,

                NULL::timestamptz
                  AS classified_at

              FROM
                public.fin_recurring_expense_months fm

              JOIN
                public.fin_recurring_expense_templates ft
                ON ft.id_recurring_template=
                   fm.id_recurring_template

              LEFT JOIN public.plan_cuentas pc
                ON pc.code=
                   ft.cuenta_code

              WHERE fm.year=:year
                AND fm.month=:month
                AND fm.is_active=TRUE
                AND ft.is_active=TRUE
                AND ft.cuenta_code=:account
            )

            SELECT *
            FROM detail

            ORDER BY
              fecha,
              source_type,
              id_item
            """
        ),
        {
            "account":
                account,

            "start":
                start,

            "end":
                end,

            "year":
                int(year),

            "month":
                int(month),
        },
    ).mappings().all()

    return {
        "ok": True,
        "items": [
            dict(row)
            for row in rows
        ],
    }


@router.get(
    "/pnl-v2/sales-detail"
)
def pnl_sales_detail_v2(
    year: int = Query(
        ...,
        ge=2020,
        le=2100,
    ),
    month: int = Query(
        ...,
        ge=1,
        le=12,
    ),
    brand: str = Query(
        ...,
        min_length=1,
        max_length=100,
    ),
    me: dict = Depends(
        legacy.get_current_user
    ),
    db=Depends(get_db),
):
    _role_pnl(me)

    confirmado_id = _fin_confirmed_state_id(
        db
    )

    start, end = _fin_month_bounds(
        year,
        month,
    )

    rows = db.execute(
        text(
            """
            SELECT
              l.id_lead,
              l.fecha_evento::date
                AS fecha_evento,

              COALESCE(
                NULLIF(
                  btrim(
                    l.cliente
                  ),
                  ''
                ),
                'Sin nombre'
              ) AS cliente,

              UPPER(
                COALESCE(
                  m.nombre,
                  m.marca,
                  ''
                )
              ) AS marca,

              COALESCE(
                l.monto_cotizado,
                0
              )::bigint
                AS venta_confirmada,

              COALESCE(
                c.traslado,
                0
              )::bigint
                AS traslado,

              (
                COALESCE(
                  l.monto_cotizado,
                  0
                )
                -
                COALESCE(
                  c.traslado,
                  0
                )
              )::bigint
                AS neto_productos

            FROM public.leads l

            LEFT JOIN
              public.marcas m
              ON m.id_marca=
                 l.id_marca

            LEFT JOIN
              public.cotizaciones c
              ON c.id_cotizacion=
                 l.id_cotizacion_vigente

            WHERE l.id_estado=:conf
              AND l.fecha_evento::date
                  >= :start
              AND l.fecha_evento::date
                  < :end
              AND UPPER(
                    COALESCE(
                      m.nombre,
                      m.marca,
                      ''
                    )
                  )=:brand

            ORDER BY
              l.fecha_evento,
              l.id_lead
            """
        ),
        {
            "conf":
                confirmado_id,

            "start":
                start,

            "end":
                end,

            "brand":
                str(
                    brand
                ).strip().upper(),
        },
    ).mappings().all()

    return {
        "ok": True,
        "items": [
            dict(row)
            for row in rows
        ],
    }



def _fin_export_width(ws):
    for column in ws.columns:
        letter = column[0].column_letter

        width = 10

        for cell in column:
            if cell.value is None:
                continue

            width = max(
                width,
                min(
                    len(
                        str(
                            cell.value
                        )
                    ) + 2,
                    55,
                ),
            )

        ws.column_dimensions[
            letter
        ].width = width


@router.get(
    "/pnl-v2/export.xlsx"
)
def pnl_export_v2(
    year: int = Query(
        ...,
        ge=2020,
        le=2100,
    ),
    month: int = Query(
        ...,
        ge=1,
        le=12,
    ),
    me: dict = Depends(
        legacy.get_current_user
    ),
    db=Depends(get_db),
):
    _role_pnl(me)

    data = pnl_management_v2(
        year=year,
        month=month,
        me=me,
        db=db,
    )

    wb = Workbook()

    months = [
        "ENE","FEB","MAR","ABR",
        "MAY","JUN","JUL","AGO",
        "SEP","OCT","NOV","DIC",
    ]

    ws = wb.active
    ws.title = "P&L"

    ws.append(
        [
            "Cuenta",
            *months,
            "TOTAL",
            "Ideal %",
            "Límite %",
        ]
    )

    for label, key in (
        ("VENTA CONFIRMADA","confirmed"),
        ("VENTA REALIZADA","realized"),
        ("AÑO PASADO","base"),
        ("META","target"),
    ):
        values = [
            _fin_num(
                data[
                    "sales_months"
                ][str(m)].get(key)
            )
            for m in range(1,13)
        ]

        ws.append(
            [
                label,
                *values,
                sum(values),
                None,
                None,
            ]
        )

    ws.append([])

    for group in (
        data.get("groups")
        or []
    ):
        values = [
            _fin_num(
                group[
                    "values"
                ][m]
            )
            for m in range(1,13)
        ]

        ws.append(
            [
                group["group"],
                *values,
                sum(values),
                _fin_num(
                    group[
                        "ideal_pct"
                    ]
                ) / 100,
                _fin_num(
                    group[
                        "limit_pct"
                    ]
                ) / 100,
            ]
        )

        for account in (
            data.get("accounts")
            or []
        ):
            if (
                account["group"]
                != group["group"]
            ):
                continue

            account_values = [
                _fin_num(
                    account[
                        "values"
                    ][m]
                )
                for m in range(1,13)
            ]

            ws.append(
                [
                    "   "
                    + str(
                        account["label"]
                    ),
                    *account_values,
                    sum(account_values),
                    _fin_num(
                        account[
                            "ideal_pct"
                        ]
                    ) / 100,
                    _fin_num(
                        account[
                            "limit_pct"
                        ]
                    ) / 100,
                ]
            )

    ws.freeze_panes = "B2"

    for row in ws.iter_rows(
        min_row=2
    ):
        for index in range(
            1,
            14,
        ):
            row[index].number_format = (
                '#,##0;[Red]-#,##0'
            )

        row[14].number_format = "0.0%"
        row[15].number_format = "0.0%"

    _fin_export_width(ws)

    ws_sales = wb.create_sheet(
        "Ventas por marca"
    )

    ws_sales.append(
        [
            "Marca",
            "Año pasado",
            "Meta",
            "Realizada",
            "Confirmada",
            "Neto productos",
            "Traslado",
            "Falta meta",
            "Dif. año pasado",
            "Cumplimiento %",
            "Estado",
            "Eventos",
        ]
    )

    for item in (
        data.get(
            "sales_by_brand"
        )
        or []
    ):
        ws_sales.append(
            [
                item.get("marca"),
                item.get(
                    "venta_base"
                ),
                item.get("meta"),
                item.get(
                    "realizada"
                ),
                item.get(
                    "confirmada"
                ),
                item.get(
                    "neto_productos"
                ),
                item.get(
                    "traslado"
                ),
                item.get(
                    "gap_meta"
                ),
                item.get(
                    "diff_base"
                ),
                (
                    _fin_num(
                        item.get(
                            "cumplimiento_pct"
                        )
                    ) / 100
                    if item.get(
                        "cumplimiento_pct"
                    )
                    is not None
                    else None
                ),
                item.get("status"),
                item.get(
                    "cnt_total"
                ),
            ]
        )

    ws_sales.freeze_panes = "A2"

    for row in ws_sales.iter_rows(
        min_row=2
    ):
        for index in range(
            1,
            9,
        ):
            row[index].number_format = (
                '#,##0;[Red]-#,##0'
            )

        row[9].number_format = "0.0%"

    _fin_export_width(
        ws_sales
    )

    detail = db.execute(
        text(
            """
            WITH d AS (

              SELECT
                g.fecha::date
                  AS fecha,

                EXTRACT(
                  MONTH FROM g.fecha
                )::int
                  AS mes,

                'MOVIMIENTO'
                  AS origen,

                g.cuenta_code,

                pc.name
                  AS cuenta,

                cfg.group_name,

                g.monto,

                g.descripcion
                  AS descripcion_banco,

                g.descripcion_interna

              FROM public.fin_gastos g

              LEFT JOIN public.plan_cuentas pc
                ON pc.code=
                   g.cuenta_code

              LEFT JOIN
                public.fin_pnl_account_config cfg
                ON cfg.cuenta_code=
                   g.cuenta_code

              WHERE g.is_active=TRUE
                AND EXTRACT(
                      YEAR FROM g.fecha
                    )::int=:year

              UNION ALL

              SELECT
                fm.economic_date,

                fm.month,

                'FIJO_BASE_MENSUAL',

                ft.cuenta_code,

                pc.name,

                cfg.group_name,

                fm.amount,

                NULL,

                ft.descripcion_interna

              FROM
                public.fin_recurring_expense_months fm

              JOIN
                public.fin_recurring_expense_templates ft
                ON ft.id_recurring_template=
                   fm.id_recurring_template

              LEFT JOIN public.plan_cuentas pc
                ON pc.code=
                   ft.cuenta_code

              LEFT JOIN
                public.fin_pnl_account_config cfg
                ON cfg.cuenta_code=
                   ft.cuenta_code

              WHERE fm.year=:year
                AND fm.is_active=TRUE
                AND ft.is_active=TRUE
            )

            SELECT *
            FROM d
            ORDER BY
              fecha,
              group_name,
              cuenta_code,
              origen
            """
        ),
        {
            "year":
                int(year)
        },
    ).mappings().all()

    ws_detail = wb.create_sheet(
        "Detalle gastos"
    )

    ws_detail.append(
        [
            "Fecha",
            "Mes",
            "Origen",
            "Código",
            "Cuenta",
            "Grupo P&L",
            "Monto",
            "Descripción banco",
            "Descripción interna",
        ]
    )

    for item in detail:
        ws_detail.append(
            list(
                item.values()
            )
        )

    ws_detail.freeze_panes = "A2"

    for row in ws_detail.iter_rows(
        min_row=2
    ):
        row[6].number_format = (
            '#,##0;[Red]-#,##0'
        )

    _fin_export_width(
        ws_detail
    )

    fixed = db.execute(
        text(
            """
            SELECT
              fm.year,
              fm.month,
              fm.economic_date,
              ft.recurring_code,
              ft.name,
              ft.cuenta_code,
              pc.name AS cuenta,
              cfg.group_name,
              fm.amount,
              ft.descripcion_interna

            FROM
              public.fin_recurring_expense_months fm

            JOIN
              public.fin_recurring_expense_templates ft
              ON ft.id_recurring_template=
                 fm.id_recurring_template

            LEFT JOIN public.plan_cuentas pc
              ON pc.code=
                 ft.cuenta_code

            LEFT JOIN
              public.fin_pnl_account_config cfg
              ON cfg.cuenta_code=
                 ft.cuenta_code

            WHERE fm.year=:year
              AND fm.is_active=TRUE
              AND ft.is_active=TRUE

            ORDER BY
              fm.month,
              cfg.group_order,
              ft.cuenta_code
            """
        ),
        {
            "year":
                int(year)
        },
    ).mappings().all()

    ws_fixed = wb.create_sheet(
        "Fijos 12 meses"
    )

    ws_fixed.append(
        [
            "Año",
            "Mes",
            "Fecha económica",
            "Código fijo",
            "Concepto",
            "Código cuenta",
            "Cuenta P&L",
            "Grupo P&L",
            "Monto exacto",
            "Descripción",
        ]
    )

    for item in fixed:
        ws_fixed.append(
            list(
                item.values()
            )
        )

    ws_fixed.freeze_panes = "A2"

    for row in ws_fixed.iter_rows(
        min_row=2
    ):
        row[8].number_format = (
            '#,##0;[Red]-#,##0'
        )

    _fin_export_width(
        ws_fixed
    )

    pending = db.execute(
        text(
            """
            SELECT
              bm.tx_date::date,
              ba.bank_name,
              ba.label,
              bm.description,
              bm.reference,
              abs(bm.amount)

            FROM public.fin_bank_movements bm

            LEFT JOIN
              public.fin_bank_accounts ba
              ON ba.id_bank_account=
                 bm.id_bank_account

            WHERE bm.status='PENDING'
              AND bm.amount < 0

            ORDER BY
              bm.tx_date,
              bm.id_bank_movement
            """
        )
    ).all()

    ws_pending = wb.create_sheet(
        "Pendientes cartola"
    )

    ws_pending.append(
        [
            "Fecha",
            "Banco",
            "Cuenta bancaria",
            "Descripción",
            "Referencia",
            "Monto",
        ]
    )

    for item in pending:
        ws_pending.append(
            list(item)
        )

    ws_pending.freeze_panes = "A2"

    for row in ws_pending.iter_rows(
        min_row=2
    ):
        row[5].number_format = (
            '#,##0;[Red]-#,##0'
        )

    _fin_export_width(
        ws_pending
    )

    ws_cash = wb.create_sheet(
        "Flujo operativo"
    )

    cf = (
        data.get(
            "cashflow"
        )
        or {}
    )

    ws_cash.append(
        [
            "Concepto",
            "Monto",
        ]
    )

    ws_cash.append(
        [
            "Venta realizada al día",
            cf.get(
                "venta_al_dia"
            ),
        ]
    )

    ws_cash.append(
        [
            "Venta confirmada mes",
            cf.get(
                "venta_confirmada_mes"
            ),
        ]
    )

    ws_cash.append(
        [
            "(-) Gastos fijos",
            cf.get(
                "gastos_fijos"
            ),
        ]
    )

    ws_cash.append(
        [
            "(-) Otros gastos clasificados",
            cf.get(
                "otros_gastos"
            ),
        ]
    )

    ws_cash.append(
        [
            "FLUJO OPERATIVO AL DÍA",
            cf.get(
                "resultado_operativo"
            ),
        ]
    )

    ws_cash.append([])
    ws_cash.append(
        [
            "Pendiente clasificar",
            (
                cf.get(
                    "pendiente_clasificar"
                )
                or {}
            ).get(
                "amount"
            ),
        ]
    )

    ws_cash.append(
        [
            "Movimientos en cuentas fijas por revisar",
            (
                cf.get(
                    "movimientos_cuentas_fijas"
                )
                or {}
            ).get(
                "amount"
            ),
        ]
    )

    for row in ws_cash.iter_rows(
        min_row=2
    ):
        row[1].number_format = (
            '#,##0;[Red]-#,##0'
        )

    _fin_export_width(
        ws_cash
    )

    output = io.BytesIO()
    wb.save(output)

    filename = (
        "P&L_Green_Diamond_"
        + str(year)
        + "_"
        + str(month).zfill(2)
        + ".xlsx"
    )

    return StreamingResponse(
        io.BytesIO(
            output.getvalue()
        ),
        media_type=(
            "application/vnd."
            "openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition":
                (
                    'attachment; filename="'
                    + filename
                    + '"'
                ),
        },
    )



@router.get(
    "/pnl-v2/income-reconciliation"
)
def pnl_income_reconciliation_v3(
    year: int = Query(
        ...,
        ge=2020,
        le=2100,
    ),
    month: int = Query(
        ...,
        ge=1,
        le=12,
    ),
    brand: str = Query(
        "",
        max_length=100,
    ),
    me: dict = Depends(
        legacy.get_current_user
    ),
    db=Depends(get_db),
):
    _role_pnl(me)

    confirmado_id = _fin_confirmed_state_id(
        db
    )

    start, end = _fin_month_bounds(
        year,
        month,
    )

    brand_norm = str(
        brand
        or ""
    ).strip().upper()

    event_params = {
        "conf":
            confirmado_id,

        "start":
            start,

        "end":
            end,
    }

    brand_sql = ""

    if brand_norm:
        brand_sql = """
          AND UPPER(
                COALESCE(
                  m.nombre,
                  m.marca,
                  ''
                )
              )=:brand
        """

        event_params[
            "brand"
        ] = brand_norm

    events = db.execute(
        text(
            f"""
            SELECT
              l.id_lead,

              l.fecha_evento::date
                AS fecha_evento,

              COALESCE(
                NULLIF(
                  btrim(
                    l.cliente
                  ),
                  ''
                ),
                'Sin nombre'
              ) AS cliente,

              UPPER(
                COALESCE(
                  m.nombre,
                  m.marca,
                  ''
                )
              ) AS marca,

              COALESCE(
                l.monto_cotizado,
                0
              )::numeric AS monto

            FROM public.leads l

            LEFT JOIN public.marcas m
              ON m.id_marca=
                 l.id_marca

            WHERE l.id_estado=:conf

              AND l.fecha_evento::date
                  >= (
                    CAST(
                      :start AS date
                    )
                    - INTERVAL '60 days'
                  )

              AND l.fecha_evento::date
                  < (
                    CAST(
                      :end AS date
                    )
                    + INTERVAL '90 days'
                  )

              {brand_sql}

            ORDER BY
              l.fecha_evento,
              l.cliente,
              l.id_lead
            """
        ),
        event_params,
    ).mappings().all()

    income_params = {
        "start":
            start,

        "end":
            end,
    }

    associated_brand_filter = ""

    if brand_norm:
        associated_brand_filter = """
          AND (
            bm.associated_lead_id IS NULL

            OR UPPER(
                 COALESCE(
                   am.nombre,
                   am.marca,
                   ''
                 )
               )=:brand
          )
        """

        income_params[
            "brand"
        ] = brand_norm

    incomes = db.execute(
        text(
            f"""
            SELECT
              bm.id_bank_movement,
              bm.tx_date,
              bm.description,
              bm.amount,
              bm.reference,

              ba.bank_name,
              ba.label AS bank_account,

              bm.associated_lead_id,
              bm.association_note,
              bm.associated_at,
              bm.associated_by,

              al.fecha_evento::date
                AS associated_event_date,

              al.cliente
                AS associated_client,

              UPPER(
                COALESCE(
                  am.nombre,
                  am.marca,
                  ''
                )
              ) AS associated_brand

            FROM public.fin_bank_movements bm

            LEFT JOIN public.fin_bank_accounts ba
              ON ba.id_bank_account=
                 bm.id_bank_account

            LEFT JOIN public.leads al
              ON al.id_lead=
                 bm.associated_lead_id

            LEFT JOIN public.marcas am
              ON am.id_marca=
                 al.id_marca

            WHERE bm.amount > 0
              AND bm.tx_date >= :start
              AND bm.tx_date < :end

              {associated_brand_filter}

            ORDER BY
              bm.tx_date DESC,
              bm.id_bank_movement DESC
            """
        ),
        income_params,
    ).mappings().all()

    return {
        "ok": True,

        "brand":
            brand_norm,

        "events": [
            dict(row)
            for row in events
        ],

        "incomes": [
            dict(row)
            for row in incomes
        ],
    }


@router.post(
    "/pnl-v2/income-associate"
)
def pnl_income_associate_v3(
    body: dict = Body(...),
    me: dict = Depends(
        legacy.get_current_user
    ),
    db=Depends(get_db),
):
    _role_pnl(me)

    try:
        movement_id = int(
            body.get(
                "bank_movement_id"
            )
        )
    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail="Movimiento bancario inválido.",
        ) from exc

    clear = bool(
        body.get(
            "clear"
        )
    )

    note = str(
        body.get(
            "note"
        )
        or ""
    ).strip()[:500]

    movement = db.execute(
        text(
            """
            SELECT
              id_bank_movement,
              amount

            FROM public.fin_bank_movements

            WHERE id_bank_movement=:id

            FOR UPDATE
            """
        ),
        {
            "id":
                movement_id
        },
    ).mappings().first()

    if not movement:
        raise HTTPException(
            status_code=404,
            detail="Movimiento bancario no existe.",
        )

    if _fin_num(
        movement.get(
            "amount"
        )
    ) <= 0:
        raise HTTPException(
            status_code=400,
            detail=(
                "Sólo los abonos/ingresos "
                "pueden asociarse a eventos."
            ),
        )

    if clear:
        db.execute(
            text(
                """
                UPDATE public.fin_bank_movements

                SET
                  associated_lead_id=NULL,
                  association_note=NULL,
                  associated_by=NULL,
                  associated_at=NULL

                WHERE id_bank_movement=:id
                """
            ),
            {
                "id":
                    movement_id
            },
        )

        db.commit()

        return {
            "ok": True,
            "cleared": True,
        }

    try:
        lead_id = int(
            body.get(
                "id_lead"
            )
        )
    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail="Selecciona un evento.",
        ) from exc

    confirmado_id = _fin_confirmed_state_id(
        db
    )

    lead = db.execute(
        text(
            """
            SELECT
              l.id_lead,
              l.cliente,
              l.fecha_evento,
              l.id_estado

            FROM public.leads l

            WHERE l.id_lead=:id

            LIMIT 1
            """
        ),
        {
            "id":
                lead_id
        },
    ).mappings().first()

    if not lead:
        raise HTTPException(
            status_code=404,
            detail="Evento/lead no existe.",
        )

    if int(
        lead.get(
            "id_estado"
        )
        or 0
    ) != int(
        confirmado_id
    ):
        raise HTTPException(
            status_code=400,
            detail=(
                "El ingreso sólo puede asociarse "
                "a un evento confirmado."
            ),
        )

    actor = _actor(me)

    db.execute(
        text(
            """
            UPDATE public.fin_bank_movements

            SET
              associated_lead_id=:lead,
              association_note=:note,
              associated_by=:actor,
              associated_at=now()

            WHERE id_bank_movement=:movement
            """
        ),
        {
            "lead":
                lead_id,

            "note":
                note or None,

            "actor":
                actor,

            "movement":
                movement_id,
        },
    )

    db.commit()

    return {
        "ok": True,

        "bank_movement_id":
            movement_id,

        "id_lead":
            lead_id,
    }



# ============================================================
# FINANCE_ENTITIES_RECONCILE_V4
# Sociedad/RUT + P&L por entidad + cobranza parcial.
# ============================================================

def _fin_v4_scope_allowed(
    bank_entity,
    event_entity,
):
    if not event_entity:
        return False

    if not bank_entity:
        return True

    if bool(
        bank_entity.get(
            "is_rolfi"
        )
    ):
        return (
            int(
                bank_entity[
                    "id_legal_entity"
                ]
            )
            ==
            int(
                event_entity[
                    "id_legal_entity"
                ]
            )
        )

    if bool(
        bank_entity.get(
            "is_holding"
        )
    ):
        return (
            bool(
                event_entity.get(
                    "is_group_member"
                )
            )
            and not bool(
                event_entity.get(
                    "is_rolfi"
                )
            )
        )

    return (
        int(
            bank_entity[
                "id_legal_entity"
            ]
        )
        ==
        int(
            event_entity[
                "id_legal_entity"
            ]
        )
    )


@router.get(
    "/entities-v4"
)
def finance_entities_v4(
    me: dict = Depends(
        legacy.get_current_user
    ),
    db=Depends(get_db),
):
    _role_pnl(me)

    rows = db.execute(
        text(
            """
            SELECT
              e.id_legal_entity,
              e.legal_code,
              e.legal_name,
              e.rut,
              e.is_group_member,
              e.is_holding,
              e.is_rolfi,

              COALESCE(
                STRING_AGG(
                  m.nombre,
                  ', '
                  ORDER BY m.nombre
                ) FILTER(
                  WHERE m.id_marca
                        IS NOT NULL
                ),
                ''
              ) AS brands,

              COUNT(
                DISTINCT ba.id_bank_account
              )::int AS bank_accounts

            FROM public.fin_legal_entities e

            LEFT JOIN public.marcas m
              ON m.id_legal_entity=
                 e.id_legal_entity
             AND m.finance_enabled=TRUE

            LEFT JOIN public.fin_bank_accounts ba
              ON ba.id_legal_entity=
                 e.id_legal_entity
             AND ba.is_active=TRUE

            WHERE e.is_active=TRUE

            GROUP BY
              e.id_legal_entity,
              e.legal_code,
              e.legal_name,
              e.rut,
              e.is_group_member,
              e.is_holding,
              e.is_rolfi

            ORDER BY
              e.is_rolfi,
              e.is_holding DESC,
              e.legal_name
            """
        )
    ).mappings().all()

    return {
        "ok": True,
        "entities": [
            dict(row)
            for row in rows
        ],
    }


@router.post(
    "/entity-bank-accounts-v4"
)
def finance_entity_bank_account_create_v4(
    body: dict = Body(...),
    me: dict = Depends(
        legacy.get_current_user
    ),
    db=Depends(get_db),
):
    _role_pnl(me)

    try:
        entity_id = int(
            body.get(
                "id_legal_entity"
            )
        )
    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail="Sociedad invalida.",
        ) from exc

    bank_name = str(
        body.get(
            "bank_name"
        )
        or ""
    ).strip()[:100]

    label = str(
        body.get(
            "label"
        )
        or ""
    ).strip()[:100]

    currency = str(
        body.get(
            "currency"
        )
        or "CLP"
    ).strip().upper()[:3]

    if not bank_name:
        raise HTTPException(
            status_code=400,
            detail="Indica el banco.",
        )

    if not label:
        raise HTTPException(
            status_code=400,
            detail="Indica una etiqueta de cuenta.",
        )

    entity = db.execute(
        text(
            """
            SELECT
              id_legal_entity
            FROM public.fin_legal_entities
            WHERE id_legal_entity=:id
              AND is_active=TRUE
            """
        ),
        {
            "id":
                entity_id
        },
    ).scalar()

    if not entity:
        raise HTTPException(
            status_code=404,
            detail="Sociedad no existe.",
        )

    row = db.execute(
        text(
            """
            INSERT INTO public.fin_bank_accounts(
              bank_name,
              label,
              currency,
              is_active,
              id_legal_entity
            )
            VALUES(
              :bank_name,
              :label,
              :currency,
              TRUE,
              :entity
            )
            RETURNING
              id_bank_account
            """
        ),
        {
            "bank_name":
                bank_name,
            "label":
                label,
            "currency":
                currency,
            "entity":
                entity_id,
        },
    ).scalar()

    db.commit()

    return {
        "ok": True,
        "id_bank_account":
            int(row),
    }


@router.get(
    "/movements/selection-detail-v4"
)
def movement_selection_detail_v4(
    ids: str = Query(
        ...,
        min_length=1,
        max_length=3000,
    ),
    me: dict = Depends(
        legacy.get_current_user
    ),
    db=Depends(get_db),
):
    _role_classify(me)

    movement_ids = []

    for piece in str(ids).split(","):
        piece = piece.strip()

        if not piece:
            continue

        try:
            value = int(piece)
        except Exception as exc:
            raise HTTPException(
                status_code=400,
                detail="ID de movimiento invalido.",
            ) from exc

        if value > 0:
            movement_ids.append(value)

    movement_ids = sorted(
        set(movement_ids)
    )

    if not movement_ids:
        raise HTTPException(
            status_code=400,
            detail="Selecciona movimientos.",
        )

    if len(movement_ids) > 100:
        raise HTTPException(
            status_code=400,
            detail="Maximo 100 movimientos.",
        )

    params = {}

    placeholders = []

    for index, movement_id in enumerate(
        movement_ids
    ):
        key = f"id_{index}"

        params[key] = movement_id

        placeholders.append(
            ":" + key
        )

    rows = db.execute(
        text(
            f"""
            SELECT
              bm.id_bank_movement,
              bm.tx_date,
              bm.description,
              bm.amount,
              bm.reference,
              bm.balance,
              bm.status,

              ba.bank_name,
              ba.label AS bank_account,
              ba.id_legal_entity,

              e.legal_name
                AS bank_legal_name,

              e.rut
                AS bank_rut

            FROM public.fin_bank_movements bm

            LEFT JOIN public.fin_bank_accounts ba
              ON ba.id_bank_account=
                 bm.id_bank_account

            LEFT JOIN public.fin_legal_entities e
              ON e.id_legal_entity=
                 ba.id_legal_entity

            WHERE bm.id_bank_movement IN (
              {", ".join(placeholders)}
            )

            ORDER BY
              bm.tx_date,
              bm.id_bank_movement
            """
        ),
        params,
    ).mappings().all()

    return {
        "ok": True,
        "items": [
            dict(row)
            for row in rows
        ],
    }


@router.post(
    "/movements/classify-v4"
)
def classify_movements_v4(
    body: dict = Body(...),
    me: dict = Depends(
        legacy.get_current_user
    ),
    db=Depends(get_db),
):
    _role_classify(me)

    try:
        entity_id = int(
            body.get(
                "id_legal_entity"
            )
        )
    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail=(
                "Selecciona la sociedad/RUT "
                "que soporta el gasto."
            ),
        ) from exc

    entity = db.execute(
        text(
            """
            SELECT
              id_legal_entity
            FROM public.fin_legal_entities
            WHERE id_legal_entity=:id
              AND is_active=TRUE
            """
        ),
        {
            "id":
                entity_id
        },
    ).scalar()

    if not entity:
        raise HTTPException(
            status_code=400,
            detail="Sociedad/RUT invalida.",
        )

    raw_ids = (
        body.get(
            "movement_ids"
        )
        or []
    )

    movement_ids = sorted(
        {
            int(value)
            for value in raw_ids
            if int(value) > 0
        }
    )

    if not movement_ids:
        raise HTTPException(
            status_code=400,
            detail="Selecciona movimientos.",
        )

    result = classify_movements_v2(
        body,
        me,
    )

    params = {
        "entity":
            entity_id,
    }

    placeholders = []

    for index, movement_id in enumerate(
        movement_ids
    ):
        key = f"id_{index}"

        params[key] = movement_id

        placeholders.append(
            ":" + key
        )

    db.execute(
        text(
            f"""
            UPDATE public.fin_gastos

            SET id_legal_entity=:entity

            WHERE is_active=TRUE
              AND bank_movement_id IN (
                {", ".join(placeholders)}
              )
            """
        ),
        params,
    )

    db.commit()

    if isinstance(
        result,
        dict
    ):
        result[
            "id_legal_entity"
        ] = entity_id

    return result


@router.get(
    "/entity-dashboard-v4"
)
def finance_entity_dashboard_v4(
    year: int = Query(
        ...,
        ge=2020,
        le=2100,
    ),
    month: int = Query(
        ...,
        ge=1,
        le=12,
    ),
    me: dict = Depends(
        legacy.get_current_user
    ),
    db=Depends(get_db),
):
    _role_pnl(me)

    start, end = _fin_month_bounds(
        year,
        month,
    )

    today_cl = datetime.now(
        ZoneInfo(
            "America/Santiago"
        )
    ).date()

    selected_period = (
        int(year),
        int(month),
    )

    current_period = (
        today_cl.year,
        today_cl.month,
    )

    if selected_period < current_period:
        cutoff = end

    elif selected_period == current_period:
        cutoff = (
            today_cl
            + timedelta(days=1)
        )

    else:
        cutoff = start

    confirmed_id = _fin_confirmed_state_id(
        db
    )

    entities = db.execute(
        text(
            """
            SELECT
              e.id_legal_entity,
              e.legal_code,
              e.legal_name,
              e.rut,
              e.is_group_member,
              e.is_holding,
              e.is_rolfi,

              COALESCE(
                STRING_AGG(
                  m.nombre,
                  ', '
                  ORDER BY m.nombre
                ) FILTER(
                  WHERE m.id_marca
                        IS NOT NULL
                ),
                ''
              ) AS brands

            FROM public.fin_legal_entities e

            LEFT JOIN public.marcas m
              ON m.id_legal_entity=
                 e.id_legal_entity
             AND m.finance_enabled=TRUE

            WHERE e.is_active=TRUE

            GROUP BY
              e.id_legal_entity,
              e.legal_code,
              e.legal_name,
              e.rut,
              e.is_group_member,
              e.is_holding,
              e.is_rolfi

            ORDER BY
              e.is_rolfi,
              e.is_holding DESC,
              e.legal_name
            """
        )
    ).mappings().all()

    sales_rows = db.execute(
        text(
            """
            SELECT
              m.id_legal_entity,

              COALESCE(
                SUM(
                  COALESCE(
                    l.monto_cotizado,
                    0
                  )
                ),
                0
              )::numeric
                AS confirmed,

              COALESCE(
                SUM(
                  CASE
                    WHEN l.fecha_evento::date
                         < :cutoff
                    THEN COALESCE(
                           l.monto_cotizado,
                           0
                         )
                    ELSE 0
                  END
                ),
                0
              )::numeric
                AS realized

            FROM public.leads l

            JOIN public.marcas m
              ON m.id_marca=
                 l.id_marca
             AND m.finance_enabled=TRUE
             AND m.id_legal_entity
                 IS NOT NULL

            WHERE l.id_estado=:confirmed
              AND COALESCE(
                    l.is_deleted,
                    FALSE
                  )=FALSE
              AND l.fecha_evento::date
                  >= :start
              AND l.fecha_evento::date
                  < :end

            GROUP BY
              m.id_legal_entity
            """
        ),
        {
            "confirmed":
                confirmed_id,
            "start":
                start,
            "end":
                end,
            "cutoff":
                cutoff,
        },
    ).mappings().all()

    sales_map = {
        int(row["id_legal_entity"]):
            {
                "confirmed":
                    _fin_num(
                        row["confirmed"]
                    ),
                "realized":
                    _fin_num(
                        row["realized"]
                    ),
            }
        for row in sales_rows
    }

    collection_rows = db.execute(
        text(
            """
            SELECT
              m.id_legal_entity,

              COALESCE(
                SUM(
                  COALESCE(
                    NULLIF(
                      c.total,
                      0
                    ),
                    NULLIF(
                      fe.monto_bruto,
                      0
                    ),
                    l.monto_cotizado,
                    0
                  )
                ),
                0
              )::numeric
                AS collection_target,

              COALESCE(
                SUM(
                  COALESCE(
                    paid.allocated,
                    0
                  )
                ),
                0
              )::numeric
                AS collected

            FROM public.leads l

            JOIN public.marcas m
              ON m.id_marca=
                 l.id_marca
             AND m.finance_enabled=TRUE
             AND m.id_legal_entity
                 IS NOT NULL

            LEFT JOIN public.cotizaciones c
              ON c.id_cotizacion=
                 l.id_cotizacion_vigente

            LEFT JOIN LATERAL (
              SELECT
                f.monto_bruto
              FROM public.fin_eventos f
              WHERE f.id_lead=
                    l.id_lead
              ORDER BY
                f.id_evento DESC
              LIMIT 1
            ) fe ON TRUE

            LEFT JOIN LATERAL (
              SELECT
                SUM(a.amount)::numeric
                  AS allocated
              FROM
                public.fin_bank_event_allocations a
              WHERE a.id_lead=
                    l.id_lead
                AND a.is_active=TRUE
            ) paid ON TRUE

            WHERE l.id_estado=:confirmed
              AND COALESCE(
                    l.is_deleted,
                    FALSE
                  )=FALSE
              AND l.fecha_evento::date
                  >= :start
              AND l.fecha_evento::date
                  < :end

            GROUP BY
              m.id_legal_entity
            """
        ),
        {
            "confirmed":
                confirmed_id,
            "start":
                start,
            "end":
                end,
        },
    ).mappings().all()

    collection_map = {
        int(row["id_legal_entity"]):
            {
                "target":
                    _fin_num(
                        row[
                            "collection_target"
                        ]
                    ),
                "collected":
                    _fin_num(
                        row[
                            "collected"
                        ]
                    ),
            }
        for row in collection_rows
    }

    fixed_rows = db.execute(
        text(
            """
            WITH base AS (
              SELECT
                fm.id_recurring_template,
                ft.recurring_code,
                ft.name,
                ft.cuenta_code,
                fm.amount
              FROM
                public.fin_recurring_expense_months fm
              JOIN
                public.fin_recurring_expense_templates ft
                ON ft.id_recurring_template=
                   fm.id_recurring_template
              WHERE fm.year=:year
                AND fm.month=:month
                AND fm.is_active=TRUE
                AND ft.is_active=TRUE
            ),

            allocation_totals AS (
              SELECT
                a.id_recurring_template,
                SUM(
                  a.monthly_amount
                )::numeric AS allocated
              FROM
                public.fin_recurring_expense_allocations a
              WHERE a.is_active=TRUE
              GROUP BY
                a.id_recurring_template
            ),

            rows AS (
              SELECT
                a.id_legal_entity,
                b.recurring_code,
                b.name,
                b.cuenta_code,
                a.monthly_amount::numeric
                  AS amount
              FROM base b
              JOIN
                public.fin_recurring_expense_allocations a
                ON a.id_recurring_template=
                   b.id_recurring_template
               AND a.is_active=TRUE

              UNION ALL

              SELECT
                NULL::bigint,
                b.recurring_code,
                b.name,
                b.cuenta_code,
                (
                  b.amount
                  -
                  COALESCE(
                    at.allocated,
                    0
                  )
                )::numeric
                  AS amount
              FROM base b
              LEFT JOIN allocation_totals at
                ON at.id_recurring_template=
                   b.id_recurring_template
              WHERE (
                b.amount
                -
                COALESCE(
                  at.allocated,
                  0
                )
              ) > 0
            )

            SELECT *
            FROM rows
            ORDER BY
              id_legal_entity NULLS LAST,
              cuenta_code,
              recurring_code
            """
        ),
        {
            "year":
                int(year),
            "month":
                int(month),
        },
    ).mappings().all()

    variable_rows = db.execute(
        text(
            """
            SELECT
              g.id_legal_entity,
              g.cuenta_code,
              COALESCE(
                pc.name,
                g.cuenta_code
              ) AS cuenta,
              SUM(
                COALESCE(
                  g.monto,
                  0
                )
              )::numeric AS amount

            FROM public.fin_gastos g

            LEFT JOIN public.plan_cuentas pc
              ON pc.code=
                 g.cuenta_code

            WHERE g.is_active=TRUE
              AND g.fecha >= :start
              AND g.fecha < :end

            GROUP BY
              g.id_legal_entity,
              g.cuenta_code,
              COALESCE(
                pc.name,
                g.cuenta_code
              )

            ORDER BY
              g.id_legal_entity NULLS LAST,
              g.cuenta_code
            """
        ),
        {
            "start":
                start,
            "end":
                end,
        },
    ).mappings().all()

    bank_rows = db.execute(
        text(
            """
            WITH last_balance AS (
              SELECT DISTINCT ON (
                bm.id_bank_account
              )
                bm.id_bank_account,
                bm.balance
              FROM public.fin_bank_movements bm
              WHERE bm.tx_date < :cutoff
              ORDER BY
                bm.id_bank_account,
                bm.tx_date DESC,
                bm.id_bank_movement DESC
            )

            SELECT
              ba.id_legal_entity,

              COUNT(
                ba.id_bank_account
              )::int AS accounts,

              COALESCE(
                SUM(
                  COALESCE(
                    lb.balance,
                    0
                  )
                ),
                0
              )::numeric AS bank_balance

            FROM public.fin_bank_accounts ba

            LEFT JOIN last_balance lb
              ON lb.id_bank_account=
                 ba.id_bank_account

            WHERE ba.is_active=TRUE

            GROUP BY
              ba.id_legal_entity
            """
        ),
        {
            "cutoff":
                cutoff
        },
    ).mappings().all()

    bank_map = {
        (
            int(row["id_legal_entity"])
            if row["id_legal_entity"]
            is not None
            else None
        ):
            {
                "accounts":
                    int(
                        row["accounts"]
                        or 0
                    ),
                "balance":
                    _fin_num(
                        row["bank_balance"]
                    ),
            }
        for row in bank_rows
    }

    fixed_by_entity = {}

    fixed_unassigned = []

    for row in fixed_rows:
        item = dict(row)

        entity_id = item.pop(
            "id_legal_entity"
        )

        item[
            "amount"
        ] = _fin_num(
            item[
                "amount"
            ]
        )

        if entity_id is None:
            fixed_unassigned.append(
                item
            )
            continue

        fixed_by_entity.setdefault(
            int(entity_id),
            [],
        ).append(item)

    variable_by_entity = {}

    variable_unassigned = []

    for row in variable_rows:
        item = dict(row)

        entity_id = item.pop(
            "id_legal_entity"
        )

        item[
            "amount"
        ] = _fin_num(
            item[
                "amount"
            ]
        )

        if entity_id is None:
            variable_unassigned.append(
                item
            )
            continue

        variable_by_entity.setdefault(
            int(entity_id),
            [],
        ).append(item)

    entity_items = []

    for row in entities:
        item = dict(row)

        entity_id = int(
            item[
                "id_legal_entity"
            ]
        )

        sales = sales_map.get(
            entity_id,
            {}
        )

        collection = collection_map.get(
            entity_id,
            {}
        )

        fixed_items = fixed_by_entity.get(
            entity_id,
            [],
        )

        variable_items = (
            variable_by_entity.get(
                entity_id,
                [],
            )
        )

        fixed_total = sum(
            _fin_num(
                value[
                    "amount"
                ]
            )
            for value in fixed_items
        )

        variable_total = sum(
            _fin_num(
                value[
                    "amount"
                ]
            )
            for value in variable_items
        )

        realized = _fin_num(
            sales.get(
                "realized"
            )
        )

        target = _fin_num(
            collection.get(
                "target"
            )
        )

        collected = _fin_num(
            collection.get(
                "collected"
            )
        )

        bank = bank_map.get(
            entity_id,
            {}
        )

        item.update(
            {
                "venta_confirmada":
                    _fin_num(
                        sales.get(
                            "confirmed"
                        )
                    ),

                "venta_realizada":
                    realized,

                "cobranza_objetivo":
                    target,

                "cobrado":
                    collected,

                "por_cobrar":
                    max(
                        target
                        - collected,
                        0,
                    ),

                "gastos_fijos":
                    fixed_total,

                "otros_gastos":
                    variable_total,

                "resultado_operativo":
                    (
                        realized
                        - fixed_total
                        - variable_total
                    ),

                "bank_accounts":
                    int(
                        bank.get(
                            "accounts"
                        )
                        or 0
                    ),

                "bank_balance":
                    _fin_num(
                        bank.get(
                            "balance"
                        )
                    ),

                "fixed_items":
                    fixed_items,

                "expense_items":
                    variable_items,
            }
        )

        entity_items.append(item)

    group_entities = [
        item
        for item in entity_items
        if bool(
            item[
                "is_group_member"
            ]
        )
        and not bool(
            item[
                "is_rolfi"
            ]
        )
    ]

    rolfi_items = [
        item
        for item in entity_items
        if bool(
            item[
                "is_rolfi"
            ]
        )
    ]

    unassigned_fixed_total = sum(
        _fin_num(
            item[
                "amount"
            ]
        )
        for item in fixed_unassigned
    )

    unassigned_variable_total = sum(
        _fin_num(
            item[
                "amount"
            ]
        )
        for item in variable_unassigned
    )

    group = {
        "legal_code":
            "GROUP_CARRITOS",

        "legal_name":
            "GRUPO CARRITOS",

        "rut":
            "",

        "venta_confirmada":
            sum(
                _fin_num(
                    item[
                        "venta_confirmada"
                    ]
                )
                for item in group_entities
            ),

        "venta_realizada":
            sum(
                _fin_num(
                    item[
                        "venta_realizada"
                    ]
                )
                for item in group_entities
            ),

        "cobranza_objetivo":
            sum(
                _fin_num(
                    item[
                        "cobranza_objetivo"
                    ]
                )
                for item in group_entities
            ),

        "cobrado":
            sum(
                _fin_num(
                    item[
                        "cobrado"
                    ]
                )
                for item in group_entities
            ),

        "por_cobrar":
            sum(
                _fin_num(
                    item[
                        "por_cobrar"
                    ]
                )
                for item in group_entities
            ),

        "gastos_fijos":
            (
                sum(
                    _fin_num(
                        item[
                            "gastos_fijos"
                        ]
                    )
                    for item in group_entities
                )
                + unassigned_fixed_total
            ),

        "otros_gastos":
            (
                sum(
                    _fin_num(
                        item[
                            "otros_gastos"
                        ]
                    )
                    for item in group_entities
                )
                + unassigned_variable_total
            ),

        "bank_accounts":
            sum(
                int(
                    item[
                        "bank_accounts"
                    ]
                    or 0
                )
                for item in group_entities
            ),

        "bank_balance":
            sum(
                _fin_num(
                    item[
                        "bank_balance"
                    ]
                )
                for item in group_entities
            ),

        "fixed_items":
            [
                value
                for item in group_entities
                for value in item[
                    "fixed_items"
                ]
            ]
            + fixed_unassigned,

        "expense_items":
            [
                value
                for item in group_entities
                for value in item[
                    "expense_items"
                ]
            ]
            + variable_unassigned,

        "unassigned_fixed":
            unassigned_fixed_total,

        "unassigned_expenses":
            unassigned_variable_total,
    }

    group[
        "resultado_operativo"
    ] = (
        _fin_num(
            group[
                "venta_realizada"
            ]
        )
        -
        _fin_num(
            group[
                "gastos_fijos"
            ]
        )
        -
        _fin_num(
            group[
                "otros_gastos"
            ]
        )
    )

    rolfi = (
        rolfi_items[0]
        if rolfi_items
        else None
    )

    return {
        "ok": True,
        "year":
            int(year),
        "month":
            int(month),
        "group":
            group,
        "rolfi":
            rolfi,
        "entities":
            entity_items,
        "unassigned": {
            "fixed":
                unassigned_fixed_total,
            "expenses":
                unassigned_variable_total,
        },
    }


@router.get(
    "/pnl-v4/reconciliation"
)
def pnl_reconciliation_v4(
    year: int = Query(
        ...,
        ge=2020,
        le=2100,
    ),
    month: int = Query(
        ...,
        ge=1,
        le=12,
    ),
    brand: str = Query(
        "",
        max_length=100,
    ),
    me: dict = Depends(
        legacy.get_current_user
    ),
    db=Depends(get_db),
):
    from difflib import SequenceMatcher
    import unicodedata

    _role_pnl(me)

    start, end = _fin_month_bounds(
        year,
        month,
    )

    confirmed_id = _fin_confirmed_state_id(
        db
    )

    brand_norm = str(
        brand
        or ""
    ).strip().upper()

    incomes = db.execute(
        text(
            """
            SELECT
              bm.id_bank_movement,
              bm.tx_date,
              bm.description,
              bm.reference,
              bm.amount,

              ba.id_bank_account,
              ba.bank_name,
              ba.label AS bank_account,

              e.id_legal_entity,
              e.legal_name,
              e.rut,
              e.is_group_member,
              e.is_holding,
              e.is_rolfi,

              COALESCE(
                paid.allocated,
                0
              )::numeric AS allocated

            FROM public.fin_bank_movements bm

            JOIN public.fin_bank_accounts ba
              ON ba.id_bank_account=
                 bm.id_bank_account

            LEFT JOIN public.fin_legal_entities e
              ON e.id_legal_entity=
                 ba.id_legal_entity

            LEFT JOIN LATERAL (
              SELECT
                SUM(a.amount)::numeric
                  AS allocated
              FROM
                public.fin_bank_event_allocations a
              WHERE a.id_bank_movement=
                    bm.id_bank_movement
                AND a.is_active=TRUE
            ) paid ON TRUE

            WHERE bm.amount > 0
              AND bm.tx_date >= :start
              AND bm.tx_date < :end

            ORDER BY
              bm.tx_date DESC,
              bm.id_bank_movement DESC
            """
        ),
        {
            "start":
                start,
            "end":
                end,
        },
    ).mappings().all()

    event_params = {
        "confirmed":
            confirmed_id,
        "start":
            start,
        "end":
            end,
    }

    brand_sql = ""

    if brand_norm:
        brand_sql = """
          AND UPPER(
                COALESCE(
                  m.nombre,
                  m.marca,
                  ''
                )
              )=:brand
        """

        event_params[
            "brand"
        ] = brand_norm

    events = db.execute(
        text(
            f"""
            SELECT
              l.id_lead,
              l.cliente,
              l.fecha_evento::date
                AS fecha_evento,

              UPPER(
                COALESCE(
                  m.nombre,
                  m.marca,
                  ''
                )
              ) AS marca,

              e.id_legal_entity,
              e.legal_name,
              e.rut,
              e.is_group_member,
              e.is_holding,
              e.is_rolfi,

              COALESCE(
                NULLIF(
                  c.total,
                  0
                ),
                NULLIF(
                  fe.monto_bruto,
                  0
                ),
                l.monto_cotizado,
                0
              )::numeric
                AS total_evento,

              COALESCE(
                paid.allocated,
                0
              )::numeric
                AS cobrado

            FROM public.leads l

            JOIN public.marcas m
              ON m.id_marca=
                 l.id_marca
             AND m.finance_enabled=TRUE
             AND m.id_legal_entity
                 IS NOT NULL

            JOIN public.fin_legal_entities e
              ON e.id_legal_entity=
                 m.id_legal_entity

            LEFT JOIN public.cotizaciones c
              ON c.id_cotizacion=
                 l.id_cotizacion_vigente

            LEFT JOIN LATERAL (
              SELECT
                f.monto_bruto
              FROM public.fin_eventos f
              WHERE f.id_lead=
                    l.id_lead
              ORDER BY
                f.id_evento DESC
              LIMIT 1
            ) fe ON TRUE

            LEFT JOIN LATERAL (
              SELECT
                SUM(a.amount)::numeric
                  AS allocated
              FROM
                public.fin_bank_event_allocations a
              WHERE a.id_lead=
                    l.id_lead
                AND a.is_active=TRUE
            ) paid ON TRUE

            WHERE l.id_estado=:confirmed
              AND COALESCE(
                    l.is_deleted,
                    FALSE
                  )=FALSE
              AND l.fecha_evento::date
                  >= :start
              AND l.fecha_evento::date
                  < :end

              {brand_sql}

            ORDER BY
              l.fecha_evento,
              l.cliente,
              l.id_lead
            """
        ),
        event_params,
    ).mappings().all()

    allocation_rows = db.execute(
        text(
            """
            SELECT
              a.id_bank_event_allocation,
              a.id_bank_movement,
              a.id_lead,
              a.amount,
              a.note,
              a.allocated_by,
              a.allocated_at,

              l.cliente,
              l.fecha_evento,

              COALESCE(
                m.nombre,
                m.marca,
                ''
              ) AS marca

            FROM
              public.fin_bank_event_allocations a

            JOIN public.leads l
              ON l.id_lead=
                 a.id_lead

            LEFT JOIN public.marcas m
              ON m.id_marca=
                 l.id_marca

            JOIN public.fin_bank_movements bm
              ON bm.id_bank_movement=
                 a.id_bank_movement

            WHERE a.is_active=TRUE
              AND bm.tx_date >= :start
              AND bm.tx_date < :end

            ORDER BY
              a.allocated_at,
              a.id_bank_event_allocation
            """
        ),
        {
            "start":
                start,
            "end":
                end,
        },
    ).mappings().all()

    allocations_by_movement = {}

    for row in allocation_rows:
        allocations_by_movement.setdefault(
            int(
                row[
                    "id_bank_movement"
                ]
            ),
            [],
        ).append(
            dict(row)
        )

    def normalize(value):
        text_value = unicodedata.normalize(
            "NFKD",
            str(
                value
                or ""
            ).upper(),
        )

        text_value = "".join(
            char
            for char in text_value
            if not unicodedata.combining(
                char
            )
        )

        return "".join(
            char
            if char.isalnum()
            else " "
            for char in text_value
        )

    event_list = []

    for row in events:
        item = dict(row)

        item[
            "total_evento"
        ] = _fin_num(
            item[
                "total_evento"
            ]
        )

        item[
            "cobrado"
        ] = _fin_num(
            item[
                "cobrado"
            ]
        )

        item[
            "pendiente"
        ] = max(
            item[
                "total_evento"
            ]
            -
            item[
                "cobrado"
            ],
            0,
        )

        if item[
            "total_evento"
        ] > 0:
            item[
                "pct_cobrado"
            ] = min(
                100.0,
                (
                    item[
                        "cobrado"
                    ]
                    /
                    item[
                        "total_evento"
                    ]
                )
                * 100.0,
            )
        else:
            item[
                "pct_cobrado"
            ] = 0.0

        event_list.append(
            item
        )

    income_list = []

    for row in incomes:
        item = dict(row)

        amount = _fin_num(
            item[
                "amount"
            ]
        )

        allocated = _fin_num(
            item[
                "allocated"
            ]
        )

        available = max(
            amount
            - allocated,
            0,
        )

        item[
            "amount"
        ] = amount

        item[
            "allocated"
        ] = allocated

        item[
            "available"
        ] = available

        item[
            "allocations"
        ] = (
            allocations_by_movement.get(
                int(
                    item[
                        "id_bank_movement"
                    ]
                ),
                [],
            )
        )

        bank_entity = {
            "id_legal_entity":
                item.get(
                    "id_legal_entity"
                ),
            "is_group_member":
                item.get(
                    "is_group_member"
                ),
            "is_holding":
                item.get(
                    "is_holding"
                ),
            "is_rolfi":
                item.get(
                    "is_rolfi"
                ),
        }

        candidates = []

        for event in event_list:
            if event[
                "pendiente"
            ] <= 0:
                continue

            event_entity = {
                "id_legal_entity":
                    event.get(
                        "id_legal_entity"
                    ),
                "is_group_member":
                    event.get(
                        "is_group_member"
                    ),
                "is_holding":
                    event.get(
                        "is_holding"
                    ),
                "is_rolfi":
                    event.get(
                        "is_rolfi"
                    ),
            }

            if not _fin_v4_scope_allowed(
                bank_entity,
                event_entity,
            ):
                continue

            score = 0.0

            if (
                item.get(
                    "id_legal_entity"
                )
                ==
                event.get(
                    "id_legal_entity"
                )
            ):
                score += 45.0

            elif bool(
                item.get(
                    "is_holding"
                )
            ):
                score += 35.0

            if (
                available > 0
                and event[
                    "pendiente"
                ] > 0
            ):
                ratio = (
                    min(
                        available,
                        event[
                            "pendiente"
                        ],
                    )
                    /
                    max(
                        available,
                        event[
                            "pendiente"
                        ],
                    )
                )

                score += (
                    ratio
                    * 25.0
                )

            try:
                days = abs(
                    (
                        event[
                            "fecha_evento"
                        ]
                        -
                        item[
                            "tx_date"
                        ]
                    ).days
                )
            except Exception:
                days = 999

            score += max(
                0.0,
                20.0
                -
                (
                    min(
                        days,
                        80
                    )
                    / 4.0
                ),
            )

            bank_text = normalize(
                item.get(
                    "description"
                )
            )

            client_text = normalize(
                event.get(
                    "cliente"
                )
            )

            if (
                bank_text
                and client_text
            ):
                client_words = [
                    word
                    for word in
                    client_text.split()
                    if len(word) >= 4
                ]

                if any(
                    word in bank_text
                    for word in client_words
                ):
                    score += 7.0

                score += (
                    SequenceMatcher(
                        None,
                        bank_text,
                        client_text,
                    ).ratio()
                    * 3.0
                )

            candidate = dict(
                event
            )

            candidate[
                "score"
            ] = round(
                score,
                1,
            )

            candidates.append(
                candidate
            )

        candidates.sort(
            key=lambda value: (
                -float(
                    value[
                        "score"
                    ]
                ),
                value[
                    "fecha_evento"
                ],
                value[
                    "id_lead"
                ],
            )
        )

        item[
            "candidates"
        ] = candidates[:8]

        income_list.append(
            item
        )

    return {
        "ok": True,
        "brand":
            brand_norm,
        "incomes":
            income_list,
        "events":
            event_list,
    }


@router.post(
    "/pnl-v4/reconciliation/allocate"
)
def pnl_reconciliation_allocate_v4(
    body: dict = Body(...),
    me: dict = Depends(
        legacy.get_current_user
    ),
    db=Depends(get_db),
):
    _role_pnl(me)

    try:
        movement_id = int(
            body.get(
                "id_bank_movement"
            )
        )

        lead_id = int(
            body.get(
                "id_lead"
            )
        )

        amount = round(
            float(
                body.get(
                    "amount"
                )
            ),
            2,
        )

    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail="Datos de conciliacion invalidos.",
        ) from exc

    if amount <= 0:
        raise HTTPException(
            status_code=400,
            detail="El monto debe ser mayor a cero.",
        )

    note = str(
        body.get(
            "note"
        )
        or ""
    ).strip()[:500]

    movement = db.execute(
        text(
            """
            SELECT
              bm.id_bank_movement,
              bm.amount,

              ba.id_legal_entity,

              e.is_group_member,
              e.is_holding,
              e.is_rolfi

            FROM public.fin_bank_movements bm

            JOIN public.fin_bank_accounts ba
              ON ba.id_bank_account=
                 bm.id_bank_account

            LEFT JOIN public.fin_legal_entities e
              ON e.id_legal_entity=
                 ba.id_legal_entity

            WHERE bm.id_bank_movement=:id

            FOR UPDATE OF bm
            """
        ),
        {
            "id":
                movement_id
        },
    ).mappings().first()

    if not movement:
        raise HTTPException(
            status_code=404,
            detail="Deposito no existe.",
        )

    movement_amount = _fin_num(
        movement[
            "amount"
        ]
    )

    if movement_amount <= 0:
        raise HTTPException(
            status_code=400,
            detail=(
                "Solo los abonos positivos "
                "pueden conciliarse."
            ),
        )

    used = _fin_num(
        db.execute(
            text(
                """
                SELECT
                  COALESCE(
                    SUM(amount),
                    0
                  )
                FROM
                  public.fin_bank_event_allocations
                WHERE id_bank_movement=:id
                  AND is_active=TRUE
                """
            ),
            {
                "id":
                    movement_id
            },
        ).scalar()
    )

    available = (
        movement_amount
        - used
    )

    if amount > available + 0.005:
        raise HTTPException(
            status_code=400,
            detail=(
                "El deposito no tiene saldo "
                "suficiente para esa asignacion."
            ),
        )

    confirmed_id = _fin_confirmed_state_id(
        db
    )

    event = db.execute(
        text(
            """
            SELECT
              l.id_lead,
              l.cliente,
              l.id_estado,

              e.id_legal_entity,
              e.is_group_member,
              e.is_holding,
              e.is_rolfi,

              COALESCE(
                NULLIF(
                  c.total,
                  0
                ),
                NULLIF(
                  fe.monto_bruto,
                  0
                ),
                l.monto_cotizado,
                0
              )::numeric
                AS total_evento

            FROM public.leads l

            JOIN public.marcas m
              ON m.id_marca=
                 l.id_marca
             AND m.finance_enabled=TRUE
             AND m.id_legal_entity
                 IS NOT NULL

            JOIN public.fin_legal_entities e
              ON e.id_legal_entity=
                 m.id_legal_entity

            LEFT JOIN public.cotizaciones c
              ON c.id_cotizacion=
                 l.id_cotizacion_vigente

            LEFT JOIN LATERAL (
              SELECT
                f.monto_bruto
              FROM public.fin_eventos f
              WHERE f.id_lead=
                    l.id_lead
              ORDER BY
                f.id_evento DESC
              LIMIT 1
            ) fe ON TRUE

            WHERE l.id_lead=:id

            LIMIT 1
            """
        ),
        {
            "id":
                lead_id
        },
    ).mappings().first()

    if not event:
        raise HTTPException(
            status_code=404,
            detail="Evento no existe.",
        )

    if int(
        event[
            "id_estado"
        ]
        or 0
    ) != int(
        confirmed_id
    ):
        raise HTTPException(
            status_code=400,
            detail=(
                "Solo se puede conciliar "
                "contra eventos confirmados."
            ),
        )

    bank_entity = {
        "id_legal_entity":
            movement.get(
                "id_legal_entity"
            ),
        "is_group_member":
            movement.get(
                "is_group_member"
            ),
        "is_holding":
            movement.get(
                "is_holding"
            ),
        "is_rolfi":
            movement.get(
                "is_rolfi"
            ),
    }

    event_entity = {
        "id_legal_entity":
            event.get(
                "id_legal_entity"
            ),
        "is_group_member":
            event.get(
                "is_group_member"
            ),
        "is_holding":
            event.get(
                "is_holding"
            ),
        "is_rolfi":
            event.get(
                "is_rolfi"
            ),
    }

    if not _fin_v4_scope_allowed(
        bank_entity,
        event_entity,
    ):
        raise HTTPException(
            status_code=400,
            detail=(
                "Ese deposito pertenece a "
                "otra sociedad/grupo."
            ),
        )

    already_collected = _fin_num(
        db.execute(
            text(
                """
                SELECT
                  COALESCE(
                    SUM(amount),
                    0
                  )
                FROM
                  public.fin_bank_event_allocations
                WHERE id_lead=:id
                  AND is_active=TRUE
                """
            ),
            {
                "id":
                    lead_id
            },
        ).scalar()
    )

    event_total = _fin_num(
        event[
            "total_evento"
        ]
    )

    pending = max(
        event_total
        - already_collected,
        0,
    )

    if amount > pending + 0.005:
        raise HTTPException(
            status_code=400,
            detail=(
                "El monto supera el saldo "
                "pendiente del evento. "
                "Asigna solamente hasta completar "
                "el 100%; el resto del deposito "
                "queda disponible."
            ),
        )

    allocation_id = db.execute(
        text(
            """
            INSERT INTO
              public.fin_bank_event_allocations(
                id_bank_movement,
                id_lead,
                amount,
                note,
                allocated_by,
                is_active
              )
            VALUES(
              :movement,
              :lead,
              :amount,
              :note,
              :actor,
              TRUE
            )
            RETURNING
              id_bank_event_allocation
            """
        ),
        {
            "movement":
                movement_id,
            "lead":
                lead_id,
            "amount":
                amount,
            "note":
                note or None,
            "actor":
                _actor(me),
        },
    ).scalar()

    db.commit()

    new_collected = (
        already_collected
        + amount
    )

    pct = (
        (
            new_collected
            / event_total
        )
        * 100.0
        if event_total > 0
        else 0.0
    )

    return {
        "ok": True,
        "id_bank_event_allocation":
            int(allocation_id),
        "event_total":
            event_total,
        "collected":
            new_collected,
        "pending":
            max(
                event_total
                - new_collected,
                0,
            ),
        "pct_cobrado":
            min(
                pct,
                100.0,
            ),
        "deposit_available":
            max(
                available
                - amount,
                0,
            ),
    }


@router.post(
    "/pnl-v4/reconciliation/unallocate"
)
def pnl_reconciliation_unallocate_v4(
    body: dict = Body(...),
    me: dict = Depends(
        legacy.get_current_user
    ),
    db=Depends(get_db),
):
    _role_pnl(me)

    try:
        allocation_id = int(
            body.get(
                "id_bank_event_allocation"
            )
        )
    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail="Asignacion invalida.",
        ) from exc

    exists = db.execute(
        text(
            """
            SELECT 1
            FROM
              public.fin_bank_event_allocations
            WHERE id_bank_event_allocation=:id
              AND is_active=TRUE
            """
        ),
        {
            "id":
                allocation_id
        },
    ).scalar()

    if not exists:
        raise HTTPException(
            status_code=404,
            detail="Asignacion no existe.",
        )

    db.execute(
        text(
            """
            UPDATE
              public.fin_bank_event_allocations

            SET
              is_active=FALSE,
              voided_by=:actor,
              voided_at=now()

            WHERE id_bank_event_allocation=:id
            """
        ),
        {
            "id":
                allocation_id,
            "actor":
                _actor(me),
        },
    )

    db.commit()

    return {
        "ok": True
    }


@router.get("/pnl")
def pnl_simple(
    year: int = Query(..., ge=2020, le=2100),
    brand: str = Query("CONSOLIDADO"),
    me: dict = Depends(legacy.get_current_user),
):
    _role_pnl(me)
    return _build_matrix(year, brand, me)


@router.get("/pnl/export")
def pnl_export(
    year: int = Query(..., ge=2020, le=2100),
    brand: str = Query("CONSOLIDADO"),
    me: dict = Depends(legacy.get_current_user),
):
    _role_pnl(me)

    data = _build_matrix(year, brand, me)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "P&L"

    months = [
        "ENE",
        "FEB",
        "MAR",
        "ABR",
        "MAY",
        "JUN",
        "JUL",
        "AGO",
        "SEP",
        "OCT",
        "NOV",
        "DIC",
    ]

    sheet.append(
        ["P&L", *months, "TOTAL"]
    )

    for row in data["rows"]:
        values = [
            row["label"],
            *[
                float(row["values"].get(month, 0))
                for month in range(1, 13)
            ],
            float(row["total"]),
        ]

        sheet.append(values)

    for cell in sheet[1]:
        cell.font = cell.font.copy(bold=True)

    for row in sheet.iter_rows(
        min_row=2,
        min_col=2,
        max_col=14,
    ):
        for cell in row:
            cell.number_format = '$#,##0;[Red]-$#,##0'

    sheet.freeze_panes = "B2"
    sheet.column_dimensions["A"].width = 42

    for column in "BCDEFGHIJKLMN":
        sheet.column_dimensions[column].width = 15

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)

    safe_brand = re.sub(r"[^A-Z0-9_-]", "_", data["brand"])
    filename = f"PL_{year}_{safe_brand}.xlsx"

    return StreamingResponse(
        output,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"'
        },
    )


# === FINANCE_TREASURY_V5_BEGIN ===

def _treasury_v5_role(me: dict) -> str:
    raw = str((me or {}).get("rol") or (me or {}).get("role") or "").strip().upper()
    normalized = re.sub(r"[ _-]+", "", raw)
    allowed = {"ADMIN", "SUPERADMIN", "FINANZAS", "CONTROLDEGESTION"}
    if normalized not in allowed:
        raise HTTPException(status_code=403, detail="No autorizado para Tesoreria / Bancos.")
    return normalized


def _treasury_v5_period(year: int, month: int) -> tuple[date, date]:
    if year < 2000 or year > 2100:
        raise HTTPException(status_code=400, detail="Ano invalido.")
    if month < 1 or month > 12:
        raise HTTPException(status_code=400, detail="Mes invalido.")
    start = date(year, month, 1)
    next_start = date(year + 1, 1, 1) if month == 12 else date(year, month + 1, 1)
    return start, next_start


def _treasury_v5_num(value):
    if value is None:
        return None
    return float(value)


@router.get("/treasury-v5")
def treasury_v5(
    year: int = Query(default=date.today().year, ge=2000, le=2100),
    month: int = Query(default=date.today().month, ge=1, le=12),
    me: dict = Depends(legacy.get_current_user),
):
    _treasury_v5_role(me)
    start, next_start = _treasury_v5_period(year, month)

    with legacy.get_connection() as conn:
        entities = conn.execute(
            text(
                """
                SELECT id_legal_entity, legal_code, legal_name, rut,
                       is_group_member, is_holding, is_rolfi
                FROM public.fin_legal_entities
                WHERE COALESCE(is_active, TRUE) IS TRUE
                ORDER BY is_rolfi ASC, legal_name ASC
                """
            )
        ).mappings().all()

        accounts = conn.execute(
            text(
                """
                SELECT ba.id_bank_account, ba.bank_name, ba.label, ba.currency,
                       ba.id_legal_entity,
                       e.legal_code, e.legal_name, e.rut,
                       COALESCE(e.is_group_member, FALSE) AS is_group_member,
                       COALESCE(e.is_rolfi, FALSE) AS is_rolfi
                FROM public.fin_bank_accounts ba
                LEFT JOIN public.fin_legal_entities e
                  ON e.id_legal_entity=ba.id_legal_entity
                WHERE COALESCE(ba.is_active, TRUE) IS TRUE
                ORDER BY COALESCE(e.is_rolfi, FALSE) ASC,
                         COALESCE(e.legal_name, 'ZZZ') ASC,
                         ba.bank_name ASC, ba.label ASC,
                         ba.id_bank_account ASC
                """
            )
        ).mappings().all()

        out = []
        group_current = 0.0
        rolfi_current = 0.0
        latest_date = None
        group_latest_date = None
        rolfi_latest_date = None
        accounts_with_balance = 0

        for account in accounts:
            account_id = int(account["id_bank_account"])
            params = {"account": account_id, "start": start, "next_start": next_start}

            pre = conn.execute(
                text(
                    """
                    SELECT tx_date, amount, balance, id_bank_movement
                    FROM public.fin_bank_movements
                    WHERE id_bank_account=:account
                      AND tx_date < :start
                      AND balance IS NOT NULL
                    ORDER BY tx_date DESC, id_bank_movement DESC
                    LIMIT 1
                    """
                ),
                params,
            ).mappings().first()

            first_period = conn.execute(
                text(
                    """
                    SELECT tx_date, amount, balance, id_bank_movement
                    FROM public.fin_bank_movements
                    WHERE id_bank_account=:account
                      AND tx_date >= :start
                      AND tx_date < :next_start
                      AND balance IS NOT NULL
                    ORDER BY tx_date ASC, id_bank_movement ASC
                    LIMIT 1
                    """
                ),
                params,
            ).mappings().first()

            last_period = conn.execute(
                text(
                    """
                    SELECT tx_date, amount, balance, id_bank_movement
                    FROM public.fin_bank_movements
                    WHERE id_bank_account=:account
                      AND tx_date >= :start
                      AND tx_date < :next_start
                      AND balance IS NOT NULL
                    ORDER BY tx_date DESC, id_bank_movement DESC
                    LIMIT 1
                    """
                ),
                params,
            ).mappings().first()

            latest = conn.execute(
                text(
                    """
                    SELECT tx_date, amount, balance, id_bank_movement
                    FROM public.fin_bank_movements
                    WHERE id_bank_account=:account
                      AND balance IS NOT NULL
                    ORDER BY tx_date DESC, id_bank_movement DESC
                    LIMIT 1
                    """
                ),
                {"account": account_id},
            ).mappings().first()

            sums = conn.execute(
                text(
                    """
                    SELECT
                      COALESCE(SUM(CASE WHEN amount > 0 THEN amount ELSE 0 END), 0) AS credits,
                      COALESCE(SUM(CASE WHEN amount < 0 THEN -amount ELSE 0 END), 0) AS debits,
                      COALESCE(SUM(amount), 0) AS net_flow,
                      COUNT(*)::int AS movement_count
                    FROM public.fin_bank_movements
                    WHERE id_bank_account=:account
                      AND tx_date >= :start
                      AND tx_date < :next_start
                    """
                ),
                params,
            ).mappings().one()

            opening = None
            derived_opening = None
            if pre and pre.get("balance") is not None:
                opening = pre["balance"]
            elif first_period and first_period.get("balance") is not None:
                derived_opening = first_period["balance"] - first_period["amount"]
                opening = derived_opening

            if first_period and first_period.get("balance") is not None:
                derived_opening = first_period["balance"] - first_period["amount"]

            continuity_diff = None
            if pre and pre.get("balance") is not None and derived_opening is not None:
                continuity_diff = derived_opening - pre["balance"]

            net_flow = sums["net_flow"]
            calculated_close = (opening + net_flow) if opening is not None else None

            if last_period and last_period.get("balance") is not None:
                reported_close = last_period["balance"]
                period_balance_date = last_period["tx_date"]
            elif pre and pre.get("balance") is not None:
                reported_close = pre["balance"]
                period_balance_date = pre["tx_date"]
            else:
                reported_close = None
                period_balance_date = None

            reconciliation_diff = None
            if reported_close is not None and calculated_close is not None:
                reconciliation_diff = reported_close - calculated_close

            current_balance = latest["balance"] if latest else None
            current_balance_date = latest["tx_date"] if latest else None

            status = "SIN_CARTOLA"
            if current_balance is not None:
                status = "OK"
                if (
                    reconciliation_diff is not None
                    and abs(float(reconciliation_diff)) > 0.5
                ) or (
                    continuity_diff is not None
                    and abs(float(continuity_diff)) > 0.5
                ):
                    status = "REVISAR"

                accounts_with_balance += 1
                current_float = float(current_balance)
                if bool(account.get("is_rolfi")):
                    rolfi_current += current_float
                    if current_balance_date is not None and (
                        rolfi_latest_date is None or current_balance_date > rolfi_latest_date
                    ):
                        rolfi_latest_date = current_balance_date
                elif bool(account.get("is_group_member")):
                    group_current += current_float
                    if current_balance_date is not None and (
                        group_latest_date is None or current_balance_date > group_latest_date
                    ):
                        group_latest_date = current_balance_date

                if current_balance_date is not None and (
                    latest_date is None or current_balance_date > latest_date
                ):
                    latest_date = current_balance_date

            out.append(
                {
                    "id_bank_account": account_id,
                    "bank_name": account["bank_name"],
                    "label": account["label"],
                    "currency": account["currency"],
                    "id_legal_entity": account.get("id_legal_entity"),
                    "legal_code": account.get("legal_code"),
                    "legal_name": account.get("legal_name") or "SIN RAZON SOCIAL",
                    "rut": account.get("rut") or "",
                    "is_group_member": bool(account.get("is_group_member")),
                    "is_rolfi": bool(account.get("is_rolfi")),
                    "opening_balance": _treasury_v5_num(opening),
                    "credits": _treasury_v5_num(sums["credits"]),
                    "debits": _treasury_v5_num(sums["debits"]),
                    "net_flow": _treasury_v5_num(net_flow),
                    "movement_count": int(sums["movement_count"] or 0),
                    "period_close": _treasury_v5_num(reported_close),
                    "period_balance_date": period_balance_date,
                    "calculated_close": _treasury_v5_num(calculated_close),
                    "reconciliation_diff": _treasury_v5_num(reconciliation_diff),
                    "continuity_diff": _treasury_v5_num(continuity_diff),
                    "current_balance": _treasury_v5_num(current_balance),
                    "current_balance_date": current_balance_date,
                    "status": status,
                }
            )

    return {
        "ok": True,
        "year": year,
        "month": month,
        "period_start": start,
        "next_period_start": next_start,
        "group_current_balance": group_current,
        "rolfi_current_balance": rolfi_current,
        "accounts_with_balance": accounts_with_balance,
        "active_accounts": len(out),
        "latest_balance_date": latest_date,
        "group_latest_balance_date": group_latest_date,
        "rolfi_latest_balance_date": rolfi_latest_date,
        "entities": [dict(row) for row in entities],
        "accounts": out,
    }


@router.get("/treasury-v5/accounts/{id_bank_account}/movements")
def treasury_v5_account_movements(
    id_bank_account: int,
    year: int = Query(..., ge=2000, le=2100),
    month: int = Query(..., ge=1, le=12),
    me: dict = Depends(legacy.get_current_user),
):
    _treasury_v5_role(me)
    start, next_start = _treasury_v5_period(year, month)

    with legacy.get_connection() as conn:
        account = conn.execute(
            text(
                """
                SELECT ba.id_bank_account, ba.bank_name, ba.label, ba.currency,
                       e.legal_name, e.rut
                FROM public.fin_bank_accounts ba
                LEFT JOIN public.fin_legal_entities e
                  ON e.id_legal_entity=ba.id_legal_entity
                WHERE ba.id_bank_account=:id
                  AND COALESCE(ba.is_active, TRUE) IS TRUE
                LIMIT 1
                """
            ),
            {"id": id_bank_account},
        ).mappings().first()

        if not account:
            raise HTTPException(status_code=404, detail="Cuenta bancaria no encontrada.")

        pre = conn.execute(
            text(
                """
                SELECT tx_date, balance
                FROM public.fin_bank_movements
                WHERE id_bank_account=:id
                  AND tx_date < :start
                  AND balance IS NOT NULL
                ORDER BY tx_date DESC, id_bank_movement DESC
                LIMIT 1
                """
            ),
            {"id": id_bank_account, "start": start},
        ).mappings().first()

        rows = conn.execute(
            text(
                """
                SELECT id_bank_movement, tx_date, description, amount,
                       reference, balance, status, cuenta_code,
                       descripcion_interna
                FROM public.fin_bank_movements
                WHERE id_bank_account=:id
                  AND tx_date >= :start
                  AND tx_date < :next_start
                ORDER BY tx_date ASC, id_bank_movement ASC
                """
            ),
            {"id": id_bank_account, "start": start, "next_start": next_start},
        ).mappings().all()

    items = []
    for row in rows:
        item = dict(row)
        item["amount"] = _treasury_v5_num(item.get("amount"))
        item["balance"] = _treasury_v5_num(item.get("balance"))
        item["kind"] = "ABONO" if float(item.get("amount") or 0) >= 0 else "CARGO"
        items.append(item)

    return {
        "ok": True,
        "year": year,
        "month": month,
        "account": dict(account),
        "opening_balance": _treasury_v5_num(pre.get("balance")) if pre else None,
        "items": items,
    }


@router.post("/treasury-bank-accounts-v5")
def treasury_v5_create_bank_account(
    body: dict = Body(...),
    me: dict = Depends(legacy.get_current_user),
):
    _treasury_v5_role(me)

    try:
        entity_id = int(body.get("id_legal_entity"))
    except Exception:
        raise HTTPException(status_code=400, detail="Razon social invalida.")

    bank_name = str(body.get("bank_name") or "").strip().upper()
    label = str(body.get("label") or "").strip().upper()
    currency = str(body.get("currency") or "CLP").strip().upper()[:3]

    if not bank_name or not label:
        raise HTTPException(status_code=400, detail="Banco y nombre de cuenta son requeridos.")
    if len(bank_name) > 120 or len(label) > 120:
        raise HTTPException(status_code=400, detail="Banco o nombre de cuenta demasiado largo.")

    with legacy.get_connection() as conn:
        entity = conn.execute(
            text(
                """
                SELECT id_legal_entity, legal_name, rut
                FROM public.fin_legal_entities
                WHERE id_legal_entity=:id
                  AND COALESCE(is_active, TRUE) IS TRUE
                LIMIT 1
                """
            ),
            {"id": entity_id},
        ).mappings().first()
        if not entity:
            raise HTTPException(status_code=404, detail="Razon social no encontrada.")

        existing = conn.execute(
            text(
                """
                SELECT id_bank_account, is_active
                FROM public.fin_bank_accounts
                WHERE id_legal_entity=:entity
                  AND upper(bank_name)=:bank
                  AND upper(label)=:label
                ORDER BY id_bank_account ASC
                LIMIT 1
                """
            ),
            {"entity": entity_id, "bank": bank_name, "label": label},
        ).mappings().first()

        if existing:
            account_id = int(existing["id_bank_account"])
            if existing.get("is_active") is False:
                conn.execute(
                    text(
                        """
                        UPDATE public.fin_bank_accounts
                        SET is_active=TRUE, currency=:currency
                        WHERE id_bank_account=:id
                        """
                    ),
                    {"currency": currency, "id": account_id},
                )
                conn.commit()
            return {
                "ok": True,
                "id_bank_account": account_id,
                "duplicate": True,
                "legal_name": entity["legal_name"],
                "rut": entity["rut"],
            }

        account_id = conn.execute(
            text(
                """
                INSERT INTO public.fin_bank_accounts(
                  bank_name, label, currency, is_active,
                  created_at, id_legal_entity
                )
                VALUES(
                  :bank, :label, :currency, TRUE,
                  now(), :entity
                )
                RETURNING id_bank_account
                """
            ),
            {
                "bank": bank_name,
                "label": label,
                "currency": currency,
                "entity": entity_id,
            },
        ).scalar_one()
        conn.commit()

    return {
        "ok": True,
        "id_bank_account": int(account_id),
        "duplicate": False,
        "legal_name": entity["legal_name"],
        "rut": entity["rut"],
    }

# === FINANCE_TREASURY_V5_END ===


# ==========================================================
# P&L V6
# UX de clasificación por movimiento / PUC / responsables
# ==========================================================

def _pnl_v6_user_id(me: dict) -> int:
    for key in (
        "id_usuario",
        "user_id",
        "id",
    ):
        value = me.get(key)

        if value is None:
            continue

        try:
            return int(value)
        except Exception:
            continue

    raise HTTPException(
        status_code=401,
        detail="No fue posible identificar al usuario.",
    )


def _pnl_v6_admin(me: dict) -> str:
    role = str(_role(me) or "").strip().upper()

    normalized = (
        role.replace(" ", "")
        .replace("_", "")
        .replace("-", "")
    )

    if normalized not in {
        "ADMIN",
        "SUPERADMIN",
        "FINANZAS",
        "CONTROLDEGESTION",
    }:
        raise HTTPException(
            status_code=403,
            detail=(
                "Sólo Administración o Finanzas puede "
                "cambiar responsables o el PUC."
            ),
        )

    return role


@router.get("/pnl-v6/meta")
def pnl_v6_meta(
    me: dict = Depends(legacy.get_current_user),
):
    role = _role(me)

    normalized = (
        str(role or "")
        .upper()
        .replace(" ", "")
        .replace("_", "")
        .replace("-", "")
    )

    can_manage = normalized in {
        "ADMIN",
        "SUPERADMIN",
        "FINANZAS",
        "CONTROLDEGESTION",
    }

    with legacy.get_connection() as conn:
        accounts = conn.execute(
            text(
                """
                SELECT
                  code,
                  name AS label,
                  type,
                  classification AS "group",
                  parent_code,
                  description
                FROM public.plan_cuentas
                WHERE COALESCE(is_active,TRUE)=TRUE
                  AND type='expense'
                ORDER BY
                  CASE
                    WHEN code ~ '^[0-9]+$'
                    THEN code::bigint
                    ELSE 999999999
                  END,
                  code
                """
            )
        ).mappings().all()

        entities = conn.execute(
            text(
                """
                SELECT
                  id_legal_entity,
                  legal_code,
                  legal_name,
                  rut,
                  is_group_member,
                  is_holding,
                  is_rolfi
                FROM public.fin_legal_entities
                WHERE is_active=TRUE
                ORDER BY
                  is_rolfi,
                  legal_name
                """
            )
        ).mappings().all()

        responsibles = conn.execute(
            text(
                """
                SELECT
                  u.id_usuario,
                  u.nombre,
                  u.rol,
                  u.cargo,
                  COALESCE(
                    (
                      SELECT rs.centro_costo
                      FROM public.rrhh_staff rs
                      WHERE rs.id_usuario=u.id_usuario
                        AND COALESCE(rs.is_active,TRUE)=TRUE
                      ORDER BY rs.id_staff DESC
                      LIMIT 1
                    ),
                    ''
                  ) AS centro_costo
                FROM public.usuarios u
                JOIN public.fin_expense_responsibles fr
                  ON fr.id_usuario=u.id_usuario
                 AND fr.is_enabled=TRUE
                WHERE u.is_active=TRUE
                ORDER BY
                  CASE
                    WHEN upper(
                      COALESCE(
                        (
                          SELECT rs2.centro_costo
                          FROM public.rrhh_staff rs2
                          WHERE rs2.id_usuario=u.id_usuario
                            AND COALESCE(rs2.is_active,TRUE)=TRUE
                          ORDER BY rs2.id_staff DESC
                          LIMIT 1
                        ),
                        ''
                      )
                    )='ROLFI'
                    THEN 0
                    ELSE 1
                  END,
                  u.nombre
                """
            )
        ).mappings().all()

        candidates = []

        if can_manage:
            candidates = conn.execute(
                text(
                    """
                    SELECT
                      u.id_usuario,
                      u.nombre,
                      u.rol,
                      u.cargo,

                      COALESCE(
                        (
                          SELECT rs.centro_costo
                          FROM public.rrhh_staff rs
                          WHERE rs.id_usuario=u.id_usuario
                            AND COALESCE(rs.is_active,TRUE)=TRUE
                          ORDER BY rs.id_staff DESC
                          LIMIT 1
                        ),
                        ''
                      ) AS centro_costo,

                      COALESCE(
                        fr.is_enabled,
                        FALSE
                      ) AS enabled

                    FROM public.usuarios u

                    LEFT JOIN public.fin_expense_responsibles fr
                      ON fr.id_usuario=u.id_usuario

                    WHERE u.is_active=TRUE

                    ORDER BY
                      u.nombre
                    """
                )
            ).mappings().all()

    return {
        "ok": True,
        "accounts": [
            dict(row)
            for row in accounts
        ],
        "entities": [
            dict(row)
            for row in entities
        ],
        "responsibles": [
            dict(row)
            for row in responsibles
        ],
        "responsible_candidates": [
            dict(row)
            for row in candidates
        ],
        "can_manage": can_manage,
    }


@router.get("/pnl-v6/movements")
def pnl_v6_movements(
    q: str = Query(
        "",
        max_length=120,
    ),
    movement_id: int | None = Query(
        None
    ),
    limit: int = Query(
        500,
        ge=1,
        le=1000,
    ),
    me: dict = Depends(
        legacy.get_current_user
    ),
):
    _role_classify(me)

    search = (
        "%"
        + str(q or "").strip()
        + "%"
    )

    with legacy.get_connection() as conn:
        rows = conn.execute(
            text(
                """
                SELECT
                  bm.id_bank_movement,
                  bm.tx_date,
                  bm.description,
                  bm.amount,
                  bm.reference,
                  bm.balance,
                  bm.status,

                  ba.id_bank_account,
                  ba.bank_name,
                  ba.label AS bank_account,
                  ba.id_legal_entity,

                  e.legal_name,
                  e.rut,
                  e.legal_code,
                  e.is_rolfi

                FROM public.fin_bank_movements bm

                JOIN public.fin_bank_accounts ba
                  ON ba.id_bank_account=
                     bm.id_bank_account

                LEFT JOIN public.fin_legal_entities e
                  ON e.id_legal_entity=
                     ba.id_legal_entity

                WHERE bm.status='PENDING'
                  AND bm.amount < 0

                  AND (
                    CAST(:movement_id AS bigint) IS NULL
                    OR bm.id_bank_movement=
                       CAST(:movement_id AS bigint)
                  )

                  AND (
                    :q=''
                    OR bm.description ILIKE :search
                    OR COALESCE(
                         bm.reference,
                         ''
                       ) ILIKE :search
                    OR ba.bank_name ILIKE :search
                    OR ba.label ILIKE :search
                    OR COALESCE(
                         e.legal_name,
                         ''
                       ) ILIKE :search
                    OR COALESCE(
                         e.rut,
                         ''
                       ) ILIKE :search
                  )

                ORDER BY
                  bm.tx_date DESC,
                  bm.id_bank_movement DESC

                LIMIT :limit
                """
            ),
            {
                "q": str(q or "").strip(),
                "search": search,
                "movement_id": movement_id,
                "limit": limit,
            },
        ).mappings().all()

    return {
        "ok": True,
        "items": [
            dict(row)
            for row in rows
        ],
    }


@router.post("/pnl-v6/classify")
def pnl_v6_classify(
    body: dict = Body(...),
    me: dict = Depends(
        legacy.get_current_user
    ),
):
    _role_classify(me)

    actor = _pnl_v6_user_id(me)

    try:
        movement_id = int(
            body.get(
                "id_bank_movement"
            )
        )
    except Exception:
        raise HTTPException(
            status_code=400,
            detail="Movimiento inválido.",
        )

    try:
        responsible_id = int(
            body.get(
                "responsable_id"
            )
        )
    except Exception:
        raise HTTPException(
            status_code=400,
            detail="Selecciona al responsable.",
        )

    account_code = str(
        body.get("cuenta_code")
        or ""
    ).strip()

    internal_description = str(
        body.get(
            "descripcion_interna"
        )
        or ""
    ).strip()

    if not account_code:
        raise HTTPException(
            status_code=400,
            detail=(
                "Selecciona una cuenta "
                "del Plan de Cuentas."
            ),
        )

    if not internal_description:
        raise HTTPException(
            status_code=400,
            detail=(
                "Indica brevemente a qué "
                "corresponde el gasto."
            ),
        )

    with legacy.get_connection() as conn:
        account = conn.execute(
            text(
                """
                SELECT
                  code,
                  name
                FROM public.plan_cuentas
                WHERE code=:code
                  AND type='expense'
                  AND COALESCE(
                        is_active,
                        TRUE
                      )=TRUE
                """
            ),
            {
                "code": account_code,
            },
        ).mappings().first()

        if not account:
            raise HTTPException(
                status_code=400,
                detail=(
                    "La cuenta PUC seleccionada "
                    "no es válida."
                ),
            )

        responsible = conn.execute(
            text(
                """
                SELECT
                  u.id_usuario,
                  u.nombre
                FROM public.usuarios u
                JOIN public.fin_expense_responsibles fr
                  ON fr.id_usuario=
                     u.id_usuario
                 AND fr.is_enabled=TRUE
                WHERE u.id_usuario=:id
                  AND u.is_active=TRUE
                """
            ),
            {
                "id": responsible_id,
            },
        ).mappings().first()

        if not responsible:
            raise HTTPException(
                status_code=400,
                detail=(
                    "El responsable seleccionado "
                    "no está habilitado para Finanzas."
                ),
            )

        movement = conn.execute(
            text(
                """
                SELECT
                  bm.id_bank_movement,
                  bm.tx_date,
                  bm.description,
                  bm.amount,
                  bm.reference,
                  bm.status,

                  ba.id_bank_account,
                  ba.id_legal_entity,
                  ba.bank_name,
                  ba.label AS bank_account,

                  e.legal_name,
                  e.rut

                FROM public.fin_bank_movements bm

                JOIN public.fin_bank_accounts ba
                  ON ba.id_bank_account=
                     bm.id_bank_account

                LEFT JOIN public.fin_legal_entities e
                  ON e.id_legal_entity=
                     ba.id_legal_entity

                WHERE bm.id_bank_movement=:id

                FOR UPDATE OF bm
                """
            ),
            {
                "id": movement_id,
            },
        ).mappings().first()

        if not movement:
            raise HTTPException(
                status_code=404,
                detail="Movimiento no encontrado.",
            )

        if (
            str(
                movement["status"]
                or ""
            ).upper()
            != "PENDING"
        ):
            raise HTTPException(
                status_code=409,
                detail=(
                    "Este movimiento ya fue "
                    "procesado."
                ),
            )

        amount = Decimal(
            str(
                movement["amount"]
            )
        )

        if amount >= 0:
            raise HTTPException(
                status_code=400,
                detail=(
                    "Este movimiento es un abono, "
                    "no un gasto."
                ),
            )

        entity_id = movement[
            "id_legal_entity"
        ]

        if entity_id is None:
            raise HTTPException(
                status_code=400,
                detail=(
                    "La cuenta bancaria no tiene "
                    "una razón social/RUT asignada."
                ),
            )

        expense_amount = abs(amount)

        conn.execute(
            text(
                """
                INSERT INTO public.fin_gastos(
                  fecha,
                  cuenta_code,
                  monto,
                  descripcion,
                  proveedor,
                  marca,
                  centro_costo,
                  created_at,
                  doc_num,
                  pagado,
                  fecha_pago,
                  tipo_doc,
                  is_active,
                  bank_movement_id,
                  responsable_id,
                  descripcion_interna,
                  id_legal_entity
                )
                VALUES(
                  :fecha,
                  :cuenta,
                  :monto,
                  :descripcion,
                  :proveedor,
                  'GENERAL',
                  NULL,
                  now(),
                  :doc_num,
                  TRUE,
                  :fecha_pago,
                  'CARTOLA',
                  TRUE,
                  :movement,
                  :responsable,
                  :descripcion_interna,
                  :entity
                )

                ON CONFLICT (
                  bank_movement_id
                )

                DO UPDATE SET
                  cuenta_code=
                    EXCLUDED.cuenta_code,
                  monto=
                    EXCLUDED.monto,
                  descripcion=
                    EXCLUDED.descripcion,
                  proveedor=
                    EXCLUDED.proveedor,
                  fecha=
                    EXCLUDED.fecha,
                  pagado=TRUE,
                  fecha_pago=
                    EXCLUDED.fecha_pago,
                  is_active=TRUE,
                  responsable_id=
                    EXCLUDED.responsable_id,
                  descripcion_interna=
                    EXCLUDED.descripcion_interna,
                  id_legal_entity=
                    EXCLUDED.id_legal_entity
                """
            ),
            {
                "fecha": movement[
                    "tx_date"
                ],
                "cuenta": account_code,
                "monto": expense_amount,
                "descripcion": movement[
                    "description"
                ],
                "proveedor": str(
                    movement[
                        "description"
                    ]
                    or ""
                )[:200],
                "doc_num": movement[
                    "reference"
                ],
                "fecha_pago": movement[
                    "tx_date"
                ],
                "movement": movement_id,
                "responsable":
                    responsible_id,
                "descripcion_interna":
                    internal_description,
                "entity": entity_id,
            },
        )

        conn.execute(
            text(
                """
                UPDATE public.fin_bank_movements

                SET
                  status='CLASSIFIED',
                  cuenta_code=:cuenta,
                  marca='GENERAL',
                  responsable_id=:responsable,
                  descripcion_interna=
                    :descripcion_interna,
                  classified_by=:actor,
                  classified_at=now()

                WHERE id_bank_movement=:id
                """
            ),
            {
                "cuenta": account_code,
                "responsable":
                    responsible_id,
                "descripcion_interna":
                    internal_description,
                "actor": actor,
                "id": movement_id,
            },
        )

        conn.commit()

    return {
        "ok": True,
        "classified": 1,
        "id_bank_movement":
            movement_id,
        "id_legal_entity":
            int(entity_id),
        "legal_name":
            movement["legal_name"],
        "rut":
            movement["rut"],
        "cuenta_code":
            account_code,
        "responsable":
            responsible[
                "nombre"
            ],
    }


@router.post("/pnl-v6/ignore")
def pnl_v6_ignore(
    body: dict = Body(...),
    me: dict = Depends(
        legacy.get_current_user
    ),
):
    _role_classify(me)

    actor = _pnl_v6_user_id(me)

    try:
        movement_id = int(
            body.get(
                "id_bank_movement"
            )
        )
    except Exception:
        raise HTTPException(
            status_code=400,
            detail="Movimiento inválido.",
        )

    with legacy.get_connection() as conn:
        result = conn.execute(
            text(
                """
                UPDATE public.fin_bank_movements

                SET
                  status='IGNORED',
                  classified_by=:actor,
                  classified_at=now()

                WHERE id_bank_movement=:id
                  AND status='PENDING'
                """
            ),
            {
                "actor": actor,
                "id": movement_id,
            },
        )

        conn.commit()

    if int(result.rowcount or 0) != 1:
        raise HTTPException(
            status_code=409,
            detail=(
                "El movimiento ya fue "
                "procesado."
            ),
        )

    return {
        "ok": True,
        "ignored": 1,
    }


@router.post("/pnl-v6/responsibles")
def pnl_v6_responsible_save(
    body: dict = Body(...),
    me: dict = Depends(
        legacy.get_current_user
    ),
):
    _pnl_v6_admin(me)

    actor = _pnl_v6_user_id(me)

    try:
        user_id = int(
            body.get("id_usuario")
        )
    except Exception:
        raise HTTPException(
            status_code=400,
            detail="Usuario inválido.",
        )

    enabled = bool(
        body.get("enabled")
    )

    with legacy.get_connection() as conn:
        user = conn.execute(
            text(
                """
                SELECT id_usuario
                FROM public.usuarios
                WHERE id_usuario=:id
                  AND is_active=TRUE
                """
            ),
            {
                "id": user_id,
            },
        ).first()

        if not user:
            raise HTTPException(
                status_code=404,
                detail="Usuario no encontrado.",
            )

        conn.execute(
            text(
                """
                INSERT INTO
                  public.fin_expense_responsibles(
                    id_usuario,
                    is_enabled,
                    updated_by,
                    updated_at
                  )

                VALUES(
                  :id,
                  :enabled,
                  :actor,
                  now()
                )

                ON CONFLICT (
                  id_usuario
                )

                DO UPDATE SET
                  is_enabled=
                    EXCLUDED.is_enabled,
                  updated_by=
                    EXCLUDED.updated_by,
                  updated_at=now()
                """
            ),
            {
                "id": user_id,
                "enabled": enabled,
                "actor": actor,
            },
        )

        conn.commit()

    return {
        "ok": True,
        "id_usuario": user_id,
        "enabled": enabled,
    }


@router.post("/pnl-v6/puc")
def pnl_v6_puc_create(
    body: dict = Body(...),
    me: dict = Depends(
        legacy.get_current_user
    ),
):
    _pnl_v6_admin(me)

    code = str(
        body.get("code")
        or ""
    ).strip()

    name = str(
        body.get("name")
        or ""
    ).strip()

    account_type = str(
        body.get("type")
        or "expense"
    ).strip().lower()

    classification = str(
        body.get(
            "classification"
        )
        or "Operacional"
    ).strip()

    parent_code = str(
        body.get("parent_code")
        or ""
    ).strip() or None

    description = str(
        body.get("description")
        or ""
    ).strip() or None

    if not code or not name:
        raise HTTPException(
            status_code=400,
            detail=(
                "Código y nombre son "
                "obligatorios."
            ),
        )

    if account_type not in {
        "expense",
        "revenue",
        "asset",
        "liability",
        "equity",
    }:
        raise HTTPException(
            status_code=400,
            detail="Tipo de cuenta inválido.",
        )

    with legacy.get_connection() as conn:
        existing = conn.execute(
            text(
                """
                SELECT code
                FROM public.plan_cuentas
                WHERE code=:code
                """
            ),
            {
                "code": code,
            },
        ).first()

        if existing:
            raise HTTPException(
                status_code=409,
                detail=(
                    "Ya existe una cuenta "
                    "con ese código."
                ),
            )

        if parent_code:
            parent = conn.execute(
                text(
                    """
                    SELECT code
                    FROM public.plan_cuentas
                    WHERE code=:code
                      AND COALESCE(
                            is_active,
                            TRUE
                          )=TRUE
                    """
                ),
                {
                    "code": parent_code,
                },
            ).first()

            if not parent:
                raise HTTPException(
                    status_code=400,
                    detail=(
                        "La cuenta padre "
                        "no existe."
                    ),
                )

        conn.execute(
            text(
                """
                INSERT INTO public.plan_cuentas(
                  code,
                  name,
                  type,
                  classification,
                  description,
                  parent_code,
                  is_active,
                  created_at,
                  updated_at
                )

                VALUES(
                  :code,
                  :name,
                  :type,
                  :classification,
                  :description,
                  :parent_code,
                  TRUE,
                  now(),
                  now()
                )
                """
            ),
            {
                "code": code,
                "name": name,
                "type": account_type,
                "classification":
                    classification,
                "description":
                    description,
                "parent_code":
                    parent_code,
            },
        )

        conn.commit()

    return {
        "ok": True,
        "code": code,
        "name": name,
        "parent_code":
            parent_code,
    }

# === GD PNL V62 MANAGEMENT START ===

def _pnl_v62_norm_brand(value):
    raw = str(value or "").strip().upper()

    replacements = (
        ("Á", "A"),
        ("É", "E"),
        ("Í", "I"),
        ("Ó", "O"),
        ("Ú", "U"),
        ("Ñ", "N"),
        ("_", " "),
        ("-", " "),
    )

    for old, new in replacements:
        raw = raw.replace(old, new)

    return " ".join(raw.split())


def _pnl_v62_find_brand(
    available,
    aliases,
):
    aliases_norm = {
        _pnl_v62_norm_brand(value)
        for value in aliases
    }

    for brand in available:
        if (
            _pnl_v62_norm_brand(brand)
            in aliases_norm
        ):
            return str(brand)

    return None


def _pnl_v62_month_value(
    matrix,
    key,
    month,
):
    for row in (matrix or {}).get("rows") or []:
        if str(row.get("key") or "") != str(key):
            continue

        values = row.get("values") or {}

        value = values.get(
            month,
            values.get(
                str(month),
                0,
            ),
        )

        try:
            return float(value or 0)
        except Exception:
            return 0.0

    return 0.0


def _pnl_v62_summary(
    matrix,
    month,
):
    ingresos = _pnl_v62_month_value(
        matrix,
        "INGRESOS",
        month,
    )

    gastos = _pnl_v62_month_value(
        matrix,
        "TOTAL_GASTOS",
        month,
    )

    resultado = _pnl_v62_month_value(
        matrix,
        "RESULTADO",
        month,
    )

    return {
        "ingresos": ingresos,
        "gastos": gastos,
        "resultado": resultado,
    }


@router.get("/pnl-v6/management")
def pnl_v62_management(
    year: int = Query(
        ...,
        ge=2020,
        le=2100,
    ),
    month: int = Query(
        ...,
        ge=1,
        le=12,
    ),
    me: dict = Depends(
        legacy.get_current_user
    ),
):
    _role_pnl(me)

    with legacy.get_connection() as conn:
        available_brands = [
            str(value)
            for value in _brands(conn)
        ]

        holding_rows = conn.execute(
            text(
                """
                SELECT
                    ma.recurring_code,
                    ma.unit_code,
                    t.name,
                    t.cuenta_code,

                    COALESCE(
                        (
                            SELECT
                                rm.amount
                            FROM
                                public.fin_recurring_expense_months rm
                            WHERE
                                rm.id_recurring_template =
                                    t.id_recurring_template
                                AND rm.year = :year
                                AND rm.month = :month
                                AND rm.is_active IS TRUE
                            ORDER BY
                                rm.id_recurring_month DESC
                            LIMIT 1
                        ),
                        t.monthly_amount
                    )::numeric AS amount

                FROM
                    public.fin_management_recurring_assignments ma

                JOIN
                    public.fin_recurring_expense_templates t
                  ON t.recurring_code =
                     ma.recurring_code

                WHERE
                    ma.unit_code = 'HOLDING'
                    AND ma.is_active IS TRUE
                    AND t.is_active IS TRUE

                ORDER BY
                    t.recurring_code
                """
            ),
            {
                "year": int(year),
                "month": int(month),
            },
        ).mappings().all()

    definitions = [
        {
            "code": "CAMALEON",
            "label": "CAMALEÓN",
            "aliases": (
                "CAMALEON",
                "CARRITOS CAMALEON",
            ),
        },
        {
            "code": "GOURMET",
            "label": "GOURMET",
            "aliases": (
                "GOURMET",
                "CARRITOS GOURMET",
            ),
        },
        {
            "code": "EXPRESS",
            "label": "EXPRESS",
            "aliases": (
                "EXPRESS",
                "CARRITOS EXPRESS",
            ),
        },
        {
            "code": "DEL_SABOR",
            "label": "DEL SABOR",
            "aliases": (
                "DEL SABOR",
                "CARRITOS DEL SABOR",
            ),
        },
    ]

    units = []

    for definition in definitions:
        real_brand = _pnl_v62_find_brand(
            available_brands,
            definition["aliases"],
        )

        if real_brand:
            matrix = _build_matrix(
                int(year),
                real_brand,
                me,
            )
        else:
            matrix = {
                "ok": True,
                "year": int(year),
                "brand": None,
                "months": list(range(1, 13)),
                "rows": [],
                "note": (
                    "Marca no encontrada "
                    "en catálogo financiero."
                ),
            }

        units.append(
            {
                "code": definition["code"],
                "label": definition["label"],
                "source_brand": real_brand,
                "summary": _pnl_v62_summary(
                    matrix,
                    int(month),
                ),
                "matrix": matrix,
            }
        )

    consolidated = _build_matrix(
        int(year),
        "CONSOLIDADO",
        me,
    )

    holding_items = [
        {
            "recurring_code":
                str(row["recurring_code"]),
            "name":
                str(row["name"] or ""),
            "cuenta_code":
                str(row["cuenta_code"] or ""),
            "amount":
                float(row["amount"] or 0),
        }
        for row in holding_rows
    ]

    holding_total = sum(
        item["amount"]
        for item in holding_items
    )

    return {
        "ok": True,

        "period": {
            "year": int(year),
            "month": int(month),
        },

        "units": units,

        "holding": {
            "code": "HOLDING",
            "label": "HOLDING",
            "amount": holding_total,
            "items": holding_items,
            "rule": (
                "100% HOLDING. "
                "No se prorratea entre marcas."
            ),
        },

        "consolidated": {
            "summary": _pnl_v62_summary(
                consolidated,
                int(month),
            ),
            "matrix": consolidated,
        },

        "available_brands":
            available_brands,

        "missing_group_brands": [
            item["code"]
            for item in units
            if not item["source_brand"]
        ],
    }


# === GD PNL V62 MANAGEMENT END ===


# === GD PNL V62 UNIT CLASSIFY START ===

PNL_V62_MANAGEMENT_UNITS = (
    "CAMALEON",
    "GOURMET",
    "EXPRESS",
    "DEL SABOR",
    "BRONTOS",
    "PETRAS",
    "MAS FLOW",
    "CORPORATIVO",
    "HOLDING",
)


def _pnl_v62_management_unit(value):
    unit = (
        str(value or "")
        .strip()
        .upper()
    )

    if unit not in PNL_V62_MANAGEMENT_UNITS:
        raise HTTPException(
            status_code=400,
            detail="Selecciona un P&L destino válido.",
        )

    return unit


@router.get("/pnl-v6/management-units")
def pnl_v62_management_units(
    me: dict = Depends(
        legacy.get_current_user
    ),
):
    _role_classify(me)

    labels = {
        "CAMALEON": "CAMALEÓN",
        "GOURMET": "GOURMET",
        "EXPRESS": "EXPRESS",
        "DEL SABOR": "DEL SABOR",
        "BRONTOS": "BRONTOS",
        "PETRAS": "PETRAS",
        "MAS FLOW": "MAS FLOW",
        "CORPORATIVO": "CORPORATIVO",
        "HOLDING": "HOLDING",
    }

    return {
        "ok": True,
        "items": [
            {
                "code": code,
                "label": labels[code],
            }
            for code in PNL_V62_MANAGEMENT_UNITS
        ],
    }


@router.post("/pnl-v6/classify-v62")
def pnl_v62_classify(
    body: dict = Body(...),
    me: dict = Depends(
        legacy.get_current_user
    ),
):
    _role_classify(me)

    unit = _pnl_v62_management_unit(
        body.get("marca")
    )

    movement_id = int(
        body.get("id_bank_movement")
        or 0
    )

    if movement_id <= 0:
        raise HTTPException(
            status_code=400,
            detail="Movimiento inválido.",
        )

    # Mantiene toda la lógica V6 actual:
    # PUC, responsable, sociedad, etc.
    result = pnl_v6_classify(
        body=body,
        me=me,
    )

    # Agregamos la dimensión de gestión.
    with legacy.get_connection() as conn:

        conn.execute(
            text(
                """
                UPDATE public.fin_bank_movements
                SET marca=:unit
                WHERE id_bank_movement=:movement_id
                """
            ),
            {
                "unit": unit,
                "movement_id": movement_id,
            },
        )

        conn.execute(
            text(
                """
                UPDATE public.fin_gastos
                SET marca=:unit
                WHERE bank_movement_id=:movement_id
                  AND COALESCE(is_active,TRUE)
                """
            ),
            {
                "unit": unit,
                "movement_id": movement_id,
            },
        )

        conn.commit()

    payload = (
        dict(result)
        if isinstance(result, dict)
        else {"ok": True}
    )

    payload["marca"] = unit
    payload["management_unit"] = unit

    return payload


# === GD PNL V62 UNIT CLASSIFY END ===


# === GD PNL V62 INCOME CLASSIFIER START ===


def _pnl_v62_valid_unit(
    unit: str,
) -> str:

    value = str(
        unit or ""
    ).strip().upper()

    with legacy.get_connection() as conn:
        valid = {
            str(item or "").strip().upper()
            for item in _brands(conn)
        }

    valid.update({
        "CORPORATIVO",
        "HOLDING",
    })

    if value not in valid:
        raise HTTPException(
            status_code=400,
            detail=(
                "Selecciona un P&L "
                "destino válido."
            ),
        )

    return value


@router.get("/pnl-v6/income-accounts")
def pnl_v62_income_accounts(
    me: dict = Depends(
        legacy.get_current_user
    ),
):
    _role_classify(me)

    with legacy.get_connection() as conn:

        rows = conn.execute(
            text(
                """
                SELECT
                  code,
                  name AS label,
                  type,
                  classification,
                  parent_code

                FROM public.plan_cuentas

                WHERE COALESCE(
                        is_active,
                        TRUE
                      )=TRUE

                  AND type IN (
                    'revenue',
                    'asset',
                    'liability',
                    'equity'
                  )

                  -- GD V62 SALES REVENUE GUARD
                  -- Las ventas normales nacen del
                  -- circuito comercial / conciliacion.
                  -- No deben volver a registrarse
                  -- desde la cartola.
                  AND code NOT IN (
                    '4000',
                    '4110',
                    '4120',
                    '4130',
                    '4140',
                    '4150'
                  )

                ORDER BY code
                """
            )
        ).mappings().all()

    return {
        "ok": True,
        "items": [
            dict(row)
            for row in rows
        ],
    }


@router.get("/pnl-v6/income-movements")
def pnl_v62_income_movements(
    q: str = Query(
        "",
        max_length=120,
    ),
    movement_id: int | None = Query(
        None
    ),
    limit: int = Query(
        500,
        ge=1,
        le=1000,
    ),
    me: dict = Depends(
        legacy.get_current_user
    ),
):
    _role_classify(me)

    needle = str(
        q or ""
    ).strip()

    search = (
        "%"
        + needle
        + "%"
    )

    with legacy.get_connection() as conn:

        rows = conn.execute(
            text(
                """
                SELECT
                  bm.id_bank_movement,
                  bm.tx_date,
                  bm.description,
                  bm.amount,
                  bm.reference,
                  bm.balance,
                  bm.status,

                  ba.id_bank_account,
                  ba.bank_name,
                  ba.label
                    AS bank_account,
                  ba.id_legal_entity,

                  e.legal_name,
                  e.rut,
                  e.legal_code,
                  e.is_rolfi

                FROM public.fin_bank_movements bm

                JOIN public.fin_bank_accounts ba
                  ON ba.id_bank_account=
                     bm.id_bank_account

                LEFT JOIN public.fin_legal_entities e
                  ON e.id_legal_entity=
                     ba.id_legal_entity

                WHERE bm.status='PENDING'
                  AND bm.amount > 0

                  -- Conciliacion antigua
                  AND bm.associated_lead_id
                      IS NULL

                  -- Conciliacion nueva
                  AND NOT EXISTS (
                    SELECT 1
                    FROM public.fin_bank_event_allocations a
                    WHERE
                      a.id_bank_movement=
                        bm.id_bank_movement
                      AND a.is_active IS TRUE
                  )

                  AND (
                    CAST(
                      :movement_id
                      AS bigint
                    ) IS NULL

                    OR bm.id_bank_movement=
                       CAST(
                         :movement_id
                         AS bigint
                       )
                  )

                  AND (
                    :q=''

                    OR bm.description
                       ILIKE :search

                    OR COALESCE(
                         bm.reference,
                         ''
                       ) ILIKE :search

                    OR ba.bank_name
                       ILIKE :search

                    OR ba.label
                       ILIKE :search

                    OR COALESCE(
                         e.legal_name,
                         ''
                       ) ILIKE :search

                    OR COALESCE(
                         e.rut,
                         ''
                       ) ILIKE :search
                  )

                ORDER BY
                  bm.tx_date DESC,
                  bm.id_bank_movement DESC

                LIMIT :limit
                """
            ),
            {
                "q": needle,
                "search": search,
                "movement_id":
                    movement_id,
                "limit": limit,
            },
        ).mappings().all()

    return {
        "ok": True,
        "items": [
            dict(row)
            for row in rows
        ],
    }


@router.post("/pnl-v6/classify-income")
def pnl_v62_classify_income(
    body: dict = Body(...),
    me: dict = Depends(
        legacy.get_current_user
    ),
):
    _role_classify(me)

    actor = _pnl_v6_user_id(me)

    try:
        movement_id = int(
            body.get(
                "id_bank_movement"
            )
        )
    except Exception:
        raise HTTPException(
            status_code=400,
            detail="Movimiento inválido.",
        )

    account_code = str(
        body.get(
            "cuenta_code"
        )
        or ""
    ).strip()

    unit = _pnl_v62_valid_unit(
        body.get("marca")
    )

    internal_description = str(
        body.get(
            "descripcion_interna"
        )
        or ""
    ).strip()

    if not account_code:
        raise HTTPException(
            status_code=400,
            detail=(
                "Selecciona una cuenta "
                "del PUC."
            ),
        )

    if not internal_description:
        raise HTTPException(
            status_code=400,
            detail=(
                "Indica brevemente a qué "
                "corresponde el ingreso."
            ),
        )

    with legacy.get_connection() as conn:

        account = conn.execute(
            text(
                """
                SELECT
                  code,
                  name,
                  type

                FROM public.plan_cuentas

                WHERE code=:code

                  AND type IN (
                    'revenue',
                    'asset',
                    'liability',
                    'equity'
                  )

                  -- GD V62 SALES REVENUE GUARD
                  -- Las ventas normales nacen del
                  -- circuito comercial / conciliacion.
                  -- No deben volver a registrarse
                  -- desde la cartola.
                  AND code NOT IN (
                    '4000',
                    '4110',
                    '4120',
                    '4130',
                    '4140',
                    '4150'
                  )

                  AND COALESCE(
                        is_active,
                        TRUE
                      )=TRUE
                """
            ),
            {
                "code": account_code,
            },
        ).mappings().first()

        if not account:
            raise HTTPException(
                status_code=400,
                detail=(
                    "La cuenta seleccionada "
                    "no es válida para "
                    "clasificar un abono."
                ),
            )

        movement = conn.execute(
            text(
                """
                SELECT
                  bm.id_bank_movement,
                  bm.tx_date,
                  bm.description,
                  bm.amount,
                  bm.reference,
                  bm.status,
                  bm.associated_lead_id,

                  ba.id_bank_account,
                  ba.id_legal_entity,
                  ba.bank_name,
                  ba.label AS bank_account,

                  e.legal_name,
                  e.rut

                FROM public.fin_bank_movements bm

                JOIN public.fin_bank_accounts ba
                  ON ba.id_bank_account=
                     bm.id_bank_account

                LEFT JOIN public.fin_legal_entities e
                  ON e.id_legal_entity=
                     ba.id_legal_entity

                WHERE bm.id_bank_movement=:id

                FOR UPDATE OF bm
                """
            ),
            {
                "id": movement_id,
            },
        ).mappings().first()

        if not movement:
            raise HTTPException(
                status_code=404,
                detail=(
                    "Movimiento no encontrado."
                ),
            )

        if str(
            movement["status"]
            or ""
        ).upper() != "PENDING":
            raise HTTPException(
                status_code=409,
                detail=(
                    "Este movimiento ya fue "
                    "procesado."
                ),
            )

        amount = Decimal(
            str(
                movement["amount"]
            )
        )

        if amount <= 0:
            raise HTTPException(
                status_code=400,
                detail=(
                    "Este movimiento no es "
                    "un abono."
                ),
            )

        if movement[
            "associated_lead_id"
        ] is not None:
            raise HTTPException(
                status_code=409,
                detail=(
                    "Este depósito ya está "
                    "asociado a una venta. "
                    "Usa Conciliar depósitos."
                ),
            )

        allocated = conn.execute(
            text(
                """
                SELECT 1
                FROM public.fin_bank_event_allocations
                WHERE
                  id_bank_movement=:id
                  AND is_active IS TRUE
                LIMIT 1
                """
            ),
            {
                "id": movement_id,
            },
        ).first()

        if allocated:
            raise HTTPException(
                status_code=409,
                detail=(
                    "Este depósito ya tiene "
                    "conciliación activa."
                ),
            )

        conn.execute(
            text(
                """
                UPDATE public.fin_bank_movements

                SET
                  status='CLASSIFIED',
                  cuenta_code=:cuenta,
                  marca=:marca,
                  descripcion_interna=
                    :descripcion,
                  classified_by=:actor,
                  classified_at=now()

                WHERE id_bank_movement=:id
                """
            ),
            {
                "cuenta":
                    account_code,
                "marca":
                    unit,
                "descripcion":
                    internal_description,
                "actor":
                    actor,
                "id":
                    movement_id,
            },
        )

        conn.commit()

    account_type = str(
        account["type"]
        or ""
    ).lower()

    return {
        "ok": True,
        "classified": 1,
        "id_bank_movement":
            movement_id,
        "cuenta_code":
            account_code,
        "account_type":
            account_type,
        "marca":
            unit,
        "affects_pnl":
            account_type == "revenue",
    }


# === GD PNL V62 INCOME CLASSIFIER END ===



# === GD FINANCE V64 DETAIL BATCH START ===


def _pnl_v64_expense_filters(
    *,
    q: str = "",
    bank_account_id: int | None = None,
    legal_entity_id: int | None = None,
    bank_import_id: int | None = None,
):
    clauses = [
        "bm.status='PENDING'",
        "bm.amount < 0",
        """NOT EXISTS (
          SELECT 1
          FROM public.fin_payable_bank_allocations allocation
          WHERE allocation.bank_movement_id=bm.id_bank_movement
            AND allocation.status='CONFIRMED'
        )""",
    ]

    params = {}

    needle = str(
        q or ""
    ).strip()

    if needle:
        params["search"] = (
            "%"
            + needle
            + "%"
        )

        clauses.append(
            """
            (
              bm.description ILIKE :search
              OR COALESCE(
                   bm.reference,
                   ''
                 ) ILIKE :search
              OR ba.bank_name ILIKE :search
              OR ba.label ILIKE :search
              OR COALESCE(
                   e.legal_name,
                   ''
                 ) ILIKE :search
              OR COALESCE(
                   e.rut,
                   ''
                 ) ILIKE :search
              OR COALESCE(
                   bi.original_name,
                   ''
                 ) ILIKE :search
            )
            """
        )

    if bank_account_id is not None:
        clauses.append(
            "bm.id_bank_account=:bank_account_id"
        )
        params["bank_account_id"] = int(
            bank_account_id
        )

    if legal_entity_id is not None:
        clauses.append(
            "ba.id_legal_entity=:legal_entity_id"
        )
        params["legal_entity_id"] = int(
            legal_entity_id
        )

    if bank_import_id is not None:
        clauses.append(
            "bm.id_bank_import=:bank_import_id"
        )
        params["bank_import_id"] = int(
            bank_import_id
        )

    return (
        "\n AND ".join(clauses),
        params,
    )


@router.get("/pnl-v6/movements-v64")
def pnl_v64_movements(
    q: str = Query(
        "",
        max_length=120,
    ),
    bank_account_id: int | None = Query(
        None
    ),
    legal_entity_id: int | None = Query(
        None
    ),
    bank_import_id: int | None = Query(
        None
    ),
    page: int = Query(
        1,
        ge=1,
    ),
    page_size: int = Query(
        25,
        ge=10,
        le=100,
    ),
    me: dict = Depends(
        legacy.get_current_user
    ),
):
    _role_classify(me)

    where_sql, params = (
        _pnl_v64_expense_filters(
            q=q,
            bank_account_id=
                bank_account_id,
            legal_entity_id=
                legal_entity_id,
            bank_import_id=
                bank_import_id,
        )
    )

    offset = (
        int(page) - 1
    ) * int(page_size)

    with legacy.get_connection() as conn:

        summary = conn.execute(
            text(
                f"""
                SELECT
                  COUNT(*)::int AS total,

                  COALESCE(
                    SUM(
                      ABS(bm.amount)
                    ),
                    0
                  )::numeric AS total_amount

                FROM public.fin_bank_movements bm

                JOIN public.fin_bank_accounts ba
                  ON ba.id_bank_account=
                     bm.id_bank_account

                LEFT JOIN public.fin_legal_entities e
                  ON e.id_legal_entity=
                     ba.id_legal_entity

                LEFT JOIN public.fin_bank_imports bi
                  ON bi.id_bank_import=
                     bm.id_bank_import

                WHERE
                  {where_sql}
                """
            ),
            params,
        ).mappings().first()

        total = int(
            summary["total"]
            or 0
        )

        rows = conn.execute(
            text(
                f"""
                SELECT
                  bm.id_bank_movement,
                  bm.id_bank_import,
                  bm.tx_date,
                  bm.description,
                  bm.amount,
                  bm.reference,
                  bm.balance,
                  bm.status,

                  ba.id_bank_account,
                  ba.bank_name,
                  ba.label AS bank_account,
                  ba.id_legal_entity,

                  e.legal_name,
                  e.rut,
                  e.legal_code,

                  bi.original_name
                    AS cartola

                FROM public.fin_bank_movements bm

                JOIN public.fin_bank_accounts ba
                  ON ba.id_bank_account=
                     bm.id_bank_account

                LEFT JOIN public.fin_legal_entities e
                  ON e.id_legal_entity=
                     ba.id_legal_entity

                LEFT JOIN public.fin_bank_imports bi
                  ON bi.id_bank_import=
                     bm.id_bank_import

                WHERE
                  {where_sql}

                ORDER BY
                  bm.tx_date DESC,
                  bm.id_bank_movement DESC

                LIMIT :page_size
                OFFSET :offset
                """
            ),
            {
                **params,
                "page_size":
                    int(page_size),
                "offset":
                    int(offset),
            },
        ).mappings().all()

        accounts = conn.execute(
            text(
                """
                SELECT
                  ba.id_bank_account,
                  ba.bank_name,
                  ba.label,
                  ba.id_legal_entity,
                  e.legal_name,
                  e.rut,
                  COUNT(bm.id_bank_movement)::int
                    AS pending_count

                FROM public.fin_bank_accounts ba

                LEFT JOIN public.fin_legal_entities e
                  ON e.id_legal_entity=
                     ba.id_legal_entity

                LEFT JOIN public.fin_bank_movements bm
                  ON bm.id_bank_account=
                     ba.id_bank_account
                 AND bm.status='PENDING'
                 AND bm.amount < 0
                 AND NOT EXISTS (
                   SELECT 1
                   FROM public.fin_payable_bank_allocations allocation
                   WHERE allocation.bank_movement_id=bm.id_bank_movement
                     AND allocation.status='CONFIRMED'
                 )

                WHERE ba.is_active IS TRUE

                GROUP BY
                  ba.id_bank_account,
                  ba.bank_name,
                  ba.label,
                  ba.id_legal_entity,
                  e.legal_name,
                  e.rut

                ORDER BY
                  ba.bank_name,
                  ba.label
                """
            )
        ).mappings().all()

        entities = conn.execute(
            text(
                """
                SELECT
                  e.id_legal_entity,
                  e.legal_name,
                  e.rut,
                  e.legal_code,
                  COUNT(bm.id_bank_movement)::int
                    AS pending_count

                FROM public.fin_legal_entities e

                LEFT JOIN public.fin_bank_accounts ba
                  ON ba.id_legal_entity=
                     e.id_legal_entity
                 AND ba.is_active IS TRUE

                LEFT JOIN public.fin_bank_movements bm
                  ON bm.id_bank_account=
                     ba.id_bank_account
                 AND bm.status='PENDING'
                 AND bm.amount < 0
                 AND NOT EXISTS (
                   SELECT 1
                   FROM public.fin_payable_bank_allocations allocation
                   WHERE allocation.bank_movement_id=bm.id_bank_movement
                     AND allocation.status='CONFIRMED'
                 )

                WHERE e.is_active IS TRUE

                GROUP BY
                  e.id_legal_entity,
                  e.legal_name,
                  e.rut,
                  e.legal_code

                ORDER BY e.legal_name
                """
            )
        ).mappings().all()

        imports = conn.execute(
            text(
                """
                SELECT
                  bi.id_bank_import,
                  bi.id_bank_account,
                  bi.original_name,
                  bi.created_at,
                  ba.bank_name,
                  ba.label AS bank_account,
                  COUNT(bm.id_bank_movement)::int
                    AS pending_count

                FROM public.fin_bank_imports bi

                JOIN public.fin_bank_accounts ba
                  ON ba.id_bank_account=
                     bi.id_bank_account

                LEFT JOIN public.fin_bank_movements bm
                  ON bm.id_bank_import=
                     bi.id_bank_import
                 AND bm.status='PENDING'
                 AND bm.amount < 0
                 AND NOT EXISTS (
                   SELECT 1
                   FROM public.fin_payable_bank_allocations allocation
                   WHERE allocation.bank_movement_id=bm.id_bank_movement
                     AND allocation.status='CONFIRMED'
                 )

                GROUP BY
                  bi.id_bank_import,
                  bi.id_bank_account,
                  bi.original_name,
                  bi.created_at,
                  ba.bank_name,
                  ba.label

                HAVING COUNT(
                  bm.id_bank_movement
                ) > 0

                ORDER BY
                  bi.created_at DESC,
                  bi.id_bank_import DESC
                """
            )
        ).mappings().all()

    pages = (
        (
            total
            + int(page_size)
            - 1
        )
        // int(page_size)
        if total
        else 0
    )

    return {
        "ok": True,
        "page": int(page),
        "page_size": int(page_size),
        "pages": int(pages),
        "total": total,
        "total_amount": float(
            summary["total_amount"]
            or 0
        ),
        "items": [
            dict(row)
            for row in rows
        ],
        "filters": {
            "bank_accounts": [
                dict(row)
                for row in accounts
            ],
            "legal_entities": [
                dict(row)
                for row in entities
            ],
            "imports": [
                dict(row)
                for row in imports
            ],
        },
    }


@router.get("/pnl-v6/movements-v64/ids")
def pnl_v64_movement_ids(
    q: str = Query(
        "",
        max_length=120,
    ),
    bank_account_id: int | None = Query(
        None
    ),
    legal_entity_id: int | None = Query(
        None
    ),
    bank_import_id: int | None = Query(
        None
    ),
    me: dict = Depends(
        legacy.get_current_user
    ),
):
    _role_classify(me)

    where_sql, params = (
        _pnl_v64_expense_filters(
            q=q,
            bank_account_id=
                bank_account_id,
            legal_entity_id=
                legal_entity_id,
            bank_import_id=
                bank_import_id,
        )
    )

    with legacy.get_connection() as conn:

        rows = conn.execute(
            text(
                f"""
                SELECT
                  bm.id_bank_movement

                FROM public.fin_bank_movements bm

                JOIN public.fin_bank_accounts ba
                  ON ba.id_bank_account=
                     bm.id_bank_account

                LEFT JOIN public.fin_legal_entities e
                  ON e.id_legal_entity=
                     ba.id_legal_entity

                LEFT JOIN public.fin_bank_imports bi
                  ON bi.id_bank_import=
                     bm.id_bank_import

                WHERE
                  {where_sql}

                ORDER BY
                  bm.tx_date DESC,
                  bm.id_bank_movement DESC

                LIMIT 5000
                """
            ),
            params,
        ).scalars().all()

    return {
        "ok": True,
        "total": len(rows),
        "movement_ids": [
            int(value)
            for value in rows
        ],
        "truncated":
            len(rows) >= 5000,
    }


@router.get("/pnl-v6/detail-v64")
def pnl_v64_detail(
    year: int = Query(
        ...,
        ge=2020,
        le=2100,
    ),
    month: int = Query(
        ...,
        ge=1,
        le=12,
    ),
    unit: str = Query(...),
    account_code: str = Query(""),
    group: str = Query(""),
    me: dict = Depends(
        legacy.get_current_user
    ),
):
    _role_pnl(me)

    unit = _pnl_v62_management_unit(
        unit
    )

    account_code = str(
        account_code or ""
    ).strip()

    group = str(
        group or ""
    ).strip()

    if account_code:
        codes = [
            account_code
        ]

    elif group:
        codes = [
            str(code)
            for code, _label, row_group
            in SIMPLE_ACCOUNTS
            if str(row_group) == group
        ]

    else:
        raise HTTPException(
            status_code=400,
            detail=(
                "Indica cuenta o grupo."
            ),
        )

    if not codes:
        raise HTTPException(
            status_code=404,
            detail=(
                "No hay cuentas para "
                "ese detalle."
            ),
        )

    start, end = _fin_month_bounds(
        int(year),
        int(month),
    )

    with legacy.get_connection() as conn:

        actual = conn.execute(
            text(
                """
                SELECT
                  g.id_gasto,
                  g.fecha,
                  g.cuenta_code,
                  pc.name AS cuenta,
                  g.monto,
                  g.descripcion,
                  g.descripcion_interna,
                  g.proveedor,
                  g.doc_num,
                  g.responsable_id,

                  u.nombre AS responsable,

                  bm.id_bank_movement,
                  bm.description
                    AS movimiento_banco,
                  bm.reference
                    AS referencia_banco,

                  ba.id_bank_account,
                  ba.bank_name,
                  ba.label AS bank_account,

                  e.legal_name,
                  e.rut

                FROM public.fin_gastos g

                LEFT JOIN public.plan_cuentas pc
                  ON pc.code=g.cuenta_code

                LEFT JOIN public.usuarios u
                  ON u.id_usuario=
                     g.responsable_id

                LEFT JOIN public.fin_bank_movements bm
                  ON bm.id_bank_movement=
                     g.bank_movement_id

                LEFT JOIN public.fin_bank_accounts ba
                  ON ba.id_bank_account=
                     bm.id_bank_account

                LEFT JOIN public.fin_legal_entities e
                  ON e.id_legal_entity=
                     g.id_legal_entity

                WHERE
                  COALESCE(
                    g.is_active,
                    TRUE
                  ) IS TRUE

                  AND g.fecha >= :start
                  AND g.fecha < :end

                  AND UPPER(
                    BTRIM(
                      COALESCE(
                        g.marca,
                        ''
                      )
                    )
                  )=:unit

                  AND g.cuenta_code=
                      ANY(
                        CAST(
                          :codes
                          AS text[]
                        )
                      )

                ORDER BY
                  g.fecha DESC,
                  g.id_gasto DESC
                """
            ),
            {
                "start": start,
                "end": end,
                "unit": unit,
                "codes": codes,
            },
        ).mappings().all()

        recurring = []

        if unit == "HOLDING":

            recurring = conn.execute(
                text(
                    """
                    SELECT
                      t.cuenta_code,
                      pc.name AS cuenta,
                      ma.recurring_code
                        AS referencia,

                      t.name AS descripcion,

                      COALESCE(
                        (
                          SELECT rm.amount
                          FROM public.fin_recurring_expense_months rm
                          WHERE
                            rm.id_recurring_template=
                              t.id_recurring_template
                            AND rm.year=:year
                            AND rm.month=:month
                            AND rm.is_active IS TRUE
                          ORDER BY
                            rm.id_recurring_month DESC
                          LIMIT 1
                        ),
                        t.monthly_amount
                      )::numeric AS monto

                    FROM public.fin_management_recurring_assignments ma

                    JOIN public.fin_recurring_expense_templates t
                      ON t.recurring_code=
                         ma.recurring_code
                     AND t.is_active IS TRUE

                    LEFT JOIN public.plan_cuentas pc
                      ON pc.code=t.cuenta_code

                    WHERE
                      ma.unit_code='HOLDING'
                      AND ma.is_active IS TRUE

                      AND t.cuenta_code=
                          ANY(
                            CAST(
                              :codes
                              AS text[]
                            )
                          )

                    ORDER BY
                      t.cuenta_code,
                      ma.recurring_code
                    """
                ),
                {
                    "year": int(year),
                    "month": int(month),
                    "codes": codes,
                },
            ).mappings().all()

        else:

            legal_code = (
                PNL_V63_ENTITY_BY_UNIT.get(
                    unit
                )
            )

            if legal_code:

                recurring = conn.execute(
                    text(
                        """
                        SELECT
                          t.cuenta_code,
                          pc.name AS cuenta,
                          a.allocation_code
                            AS referencia,

                          COALESCE(
                            NULLIF(
                              a.note,
                              ''
                            ),
                            t.name
                          ) AS descripcion,

                          a.monthly_amount::numeric
                            AS monto,

                          e.legal_name,
                          e.rut

                        FROM public.fin_recurring_expense_allocations a

                        JOIN public.fin_recurring_expense_templates t
                          ON t.id_recurring_template=
                             a.id_recurring_template
                         AND t.is_active IS TRUE

                        JOIN public.fin_legal_entities e
                          ON e.id_legal_entity=
                             a.id_legal_entity
                         AND e.is_active IS TRUE

                        LEFT JOIN public.plan_cuentas pc
                          ON pc.code=t.cuenta_code

                        WHERE
                          a.is_active IS TRUE
                          AND e.legal_code=:legal_code

                          AND t.cuenta_code=
                              ANY(
                                CAST(
                                  :codes
                                  AS text[]
                                )
                              )

                        ORDER BY
                          t.cuenta_code,
                          a.allocation_code
                        """
                    ),
                    {
                        "legal_code":
                            legal_code,
                        "codes":
                            codes,
                    },
                ).mappings().all()

    actual_by_code = {
        code: 0.0
        for code in codes
    }

    recurring_by_code = {
        code: 0.0
        for code in codes
    }

    items = []

    for row in actual:

        code = str(
            row["cuenta_code"]
        )

        amount = float(
            row["monto"]
            or 0
        )

        actual_by_code[code] = (
            actual_by_code.get(
                code,
                0.0,
            )
            + amount
        )

        item = dict(row)
        item["source"] = "CARTOLA"
        items.append(item)

    recurring_items = []

    for row in recurring:

        code = str(
            row["cuenta_code"]
        )

        amount = float(
            row["monto"]
            or 0
        )

        recurring_by_code[code] = (
            recurring_by_code.get(
                code,
                0.0,
            )
            + amount
        )

        item = dict(row)
        item["source"] = "RECURRENTE"
        recurring_items.append(item)

    displayed_by_code = {
        code: max(
            actual_by_code.get(
                code,
                0.0,
            ),
            recurring_by_code.get(
                code,
                0.0,
            ),
        )
        for code in codes
    }

    actual_total = sum(
        actual_by_code.values()
    )

    recurring_total = sum(
        recurring_by_code.values()
    )

    displayed_total = sum(
        displayed_by_code.values()
    )

    return {
        "ok": True,
        "year": int(year),
        "month": int(month),
        "unit": unit,
        "account_code":
            account_code or None,
        "group":
            group or None,
        "codes": codes,

        "actual_total":
            actual_total,

        "recurring_total":
            recurring_total,

        "displayed_total":
            displayed_total,

        "actual_count":
            len(items),

        "recurring_count":
            len(recurring_items),

        "items":
            items,

        "recurring_items":
            recurring_items,
    }


@router.post("/pnl-v6/classify-batch-v64")
def pnl_v64_classify_batch(
    body: dict = Body(...),
    me: dict = Depends(
        legacy.get_current_user
    ),
):
    _role_classify(me)

    actor = _pnl_v6_user_id(me)

    raw_ids = (
        body.get("movement_ids")
        or []
    )

    try:
        movement_ids = sorted({
            int(value)
            for value in raw_ids
            if int(value) > 0
        })
    except Exception:
        raise HTTPException(
            status_code=400,
            detail=(
                "Lista de movimientos "
                "inválida."
            ),
        )

    if not movement_ids:
        raise HTTPException(
            status_code=400,
            detail=(
                "Selecciona al menos "
                "un movimiento."
            ),
        )

    if len(movement_ids) > 1000:
        raise HTTPException(
            status_code=400,
            detail=(
                "Máximo 1000 movimientos "
                "por clasificación."
            ),
        )

    account_code = str(
        body.get("cuenta_code")
        or ""
    ).strip()

    internal_description = str(
        body.get(
            "descripcion_interna"
        )
        or ""
    ).strip()

    try:
        responsible_id = int(
            body.get(
                "responsable_id"
            )
        )
    except Exception:
        raise HTTPException(
            status_code=400,
            detail=(
                "Responsable inválido."
            ),
        )

    unit = _pnl_v62_management_unit(
        body.get("marca")
    )

    if not account_code:
        raise HTTPException(
            status_code=400,
            detail=(
                "Selecciona una cuenta "
                "PUC."
            ),
        )

    if not internal_description:
        raise HTTPException(
            status_code=400,
            detail=(
                "Indica una descripción "
                "interna."
            ),
        )

    with legacy.get_connection() as conn:

        account = conn.execute(
            text(
                """
                SELECT
                  code,
                  name,
                  type

                FROM public.plan_cuentas

                WHERE code=:code
                  AND type='expense'
                  AND COALESCE(
                        is_active,
                        TRUE
                      ) IS TRUE
                """
            ),
            {
                "code":
                    account_code,
            },
        ).mappings().first()

        if not account:
            raise HTTPException(
                status_code=400,
                detail=(
                    "La cuenta PUC no es "
                    "una cuenta de gasto válida."
                ),
            )

        responsible = conn.execute(
            text(
                """
                SELECT
                  u.id_usuario,
                  u.nombre

                FROM public.usuarios u

                JOIN public.fin_expense_responsibles fr
                  ON fr.id_usuario=
                     u.id_usuario
                 AND fr.is_enabled IS TRUE

                WHERE
                  u.id_usuario=:id
                  AND u.is_active IS TRUE
                """
            ),
            {
                "id":
                    responsible_id,
            },
        ).mappings().first()

        if not responsible:
            raise HTTPException(
                status_code=400,
                detail=(
                    "El responsable no está "
                    "habilitado para Finanzas."
                ),
            )

        movements = conn.execute(
            text(
                """
                SELECT
                  bm.id_bank_movement,
                  bm.tx_date,
                  bm.description,
                  bm.amount,
                  bm.reference,
                  bm.status,

                  ba.id_bank_account,
                  ba.id_legal_entity,

                  e.legal_name,
                  e.rut,

                  EXISTS (
                    SELECT 1
                    FROM public.fin_payable_bank_allocations allocation
                    WHERE allocation.bank_movement_id=bm.id_bank_movement
                      AND allocation.status='CONFIRMED'
                  ) AS reconciled_with_payable

                FROM public.fin_bank_movements bm

                JOIN public.fin_bank_accounts ba
                  ON ba.id_bank_account=
                     bm.id_bank_account

                LEFT JOIN public.fin_legal_entities e
                  ON e.id_legal_entity=
                     ba.id_legal_entity

                WHERE
                  bm.id_bank_movement=
                    ANY(
                      CAST(
                        :ids
                        AS bigint[]
                      )
                    )

                ORDER BY
                  bm.id_bank_movement

                FOR UPDATE OF bm
                """
            ),
            {
                "ids":
                    movement_ids,
            },
        ).mappings().all()

        if len(movements) != len(
            movement_ids
        ):
            raise HTTPException(
                status_code=409,
                detail=(
                    "Uno o más movimientos "
                    "ya no existen."
                ),
            )

        errors = []

        reconciled_ids = [
            int(movement["id_bank_movement"])
            for movement in movements
            if movement["reconciled_with_payable"]
        ]

        if reconciled_ids:
            raise HTTPException(
                status_code=409,
                detail={
                    "code": "FINANCE_MOVEMENT_RECONCILED",
                    "message": (
                        "No se pudo clasificar el movimiento porque ya está "
                        "conciliado con una factura."
                    ),
                },
            )

        for movement in movements:

            if str(
                movement["status"]
                or ""
            ).upper() != "PENDING":
                errors.append(
                    f'{movement["id_bank_movement"]}: '
                    "ya procesado"
                )
                continue

            amount = Decimal(
                str(
                    movement["amount"]
                )
            )

            if amount >= 0:
                errors.append(
                    f'{movement["id_bank_movement"]}: '
                    "no es cargo"
                )

            if movement[
                "id_legal_entity"
            ] is None:
                errors.append(
                    f'{movement["id_bank_movement"]}: '
                    "cuenta sin RUT/empresa"
                )

        if errors:
            raise HTTPException(
                status_code=409,
                detail=(
                    "No se clasificó nada. "
                    + " | ".join(
                        errors[:10]
                    )
                ),
            )

        gasto_rows = []

        total_amount = Decimal("0")

        for movement in movements:

            amount = abs(
                Decimal(
                    str(
                        movement["amount"]
                    )
                )
            )

            total_amount += amount

            gasto_rows.append(
                {
                    "fecha":
                        movement["tx_date"],

                    "cuenta":
                        account_code,

                    "monto":
                        amount,

                    "descripcion":
                        movement["description"],

                    "proveedor":
                        str(
                            movement[
                                "description"
                            ]
                            or ""
                        )[:200],

                    "marca":
                        unit,

                    "doc_num":
                        movement["reference"],

                    "fecha_pago":
                        movement["tx_date"],

                    "movement":
                        int(
                            movement[
                                "id_bank_movement"
                            ]
                        ),

                    "responsable":
                        responsible_id,

                    "descripcion_interna":
                        internal_description,

                    "entity":
                        int(
                            movement[
                                "id_legal_entity"
                            ]
                        ),
                }
            )

        conn.execute(
            text(
                """
                INSERT INTO public.fin_gastos(
                  fecha,
                  cuenta_code,
                  monto,
                  descripcion,
                  proveedor,
                  marca,
                  centro_costo,
                  created_at,
                  doc_num,
                  pagado,
                  fecha_pago,
                  tipo_doc,
                  is_active,
                  bank_movement_id,
                  responsable_id,
                  descripcion_interna,
                  id_legal_entity
                )
                VALUES(
                  :fecha,
                  :cuenta,
                  :monto,
                  :descripcion,
                  :proveedor,
                  :marca,
                  NULL,
                  now(),
                  :doc_num,
                  TRUE,
                  :fecha_pago,
                  'CARTOLA',
                  TRUE,
                  :movement,
                  :responsable,
                  :descripcion_interna,
                  :entity
                )

                ON CONFLICT (
                  bank_movement_id
                )

                DO UPDATE SET
                  cuenta_code=
                    EXCLUDED.cuenta_code,

                  monto=
                    EXCLUDED.monto,

                  descripcion=
                    EXCLUDED.descripcion,

                  proveedor=
                    EXCLUDED.proveedor,

                  marca=
                    EXCLUDED.marca,

                  fecha=
                    EXCLUDED.fecha,

                  pagado=TRUE,

                  fecha_pago=
                    EXCLUDED.fecha_pago,

                  is_active=TRUE,

                  responsable_id=
                    EXCLUDED.responsable_id,

                  descripcion_interna=
                    EXCLUDED.descripcion_interna,

                  id_legal_entity=
                    EXCLUDED.id_legal_entity
                """
            ),
            gasto_rows,
        )

        conn.execute(
            text(
                """
                UPDATE public.fin_bank_movements

                SET
                  status='CLASSIFIED',
                  cuenta_code=:cuenta,
                  marca=:marca,
                  responsable_id=:responsable,
                  descripcion_interna=
                    :descripcion_interna,
                  classified_by=:actor,
                  classified_at=now()

                WHERE
                  id_bank_movement=
                    ANY(
                      CAST(
                        :ids
                        AS bigint[]
                      )
                    )
                """
            ),
            {
                "cuenta":
                    account_code,

                "marca":
                    unit,

                "responsable":
                    responsible_id,

                "descripcion_interna":
                    internal_description,

                "actor":
                    actor,

                "ids":
                    movement_ids,
            },
        )

        conn.commit()

    return {
        "ok": True,
        "message": (
            "✓ Gasto clasificado\n\n"
            f"Cuenta: {account_code} · {account['name']}\n"
            "Monto: " + f"${float(total_amount):,.0f}".replace(",", ".") + "\n"
            "Empresa: "
            + ", ".join(sorted({str(row["legal_name"] or row["rut"] or "Sin empresa") for row in movements}))
        ),
        "classified":
            len(movement_ids),

        "movement_ids":
            movement_ids,

        "cuenta_code":
            account_code,

        "marca":
            unit,

        "responsable_id":
            responsible_id,

        "total_amount":
            float(total_amount),

        "account_name":
            account["name"],

        "companies":
            sorted({
                str(row["legal_name"] or row["rut"] or "Sin empresa")
                for row in movements
            }),
    }


# === GD FINANCE V64 DETAIL BATCH END ===


# === GD FINANCE LOANS V1 BEGIN ===

def _ensure_fin_loans_v1(conn) -> None:
    conn.execute(text("""
      CREATE TABLE IF NOT EXISTS public.fin_loans(
        id_loan BIGSERIAL PRIMARY KEY,
        external_key TEXT NOT NULL UNIQUE,
        lender TEXT NOT NULL,
        id_legal_entity INTEGER NOT NULL REFERENCES public.fin_legal_entities(id_legal_entity),
        principal NUMERIC(18,2) NOT NULL CHECK(principal >= 0),
        installment_amount NUMERIC(18,2) NOT NULL CHECK(installment_amount >= 0),
        total_installments INTEGER NOT NULL CHECK(total_installments > 0),
        paid_installments INTEGER NOT NULL DEFAULT 0 CHECK(paid_installments >= 0),
        annual_rate NUMERIC(9,6),
        status TEXT NOT NULL DEFAULT 'ACTIVE',
        notes TEXT,
        created_at TIMESTAMPTZ NOT NULL DEFAULT now(),
        updated_at TIMESTAMPTZ NOT NULL DEFAULT now(),
        CHECK(paid_installments <= total_installments)
      )
    """))
    seeds = (
        ("GREEN_BCH_01", "Banco de Chile", 1, 72656300, 2136950, 34, 16),
        ("GREEN_BCH_02", "Banco de Chile", 1, 121888440, 3385790, 36, 2),
        ("GD_FORUM_01", "Forum", 2, 38559024, 803313, 48, 12),
        ("ROLFI_PEDIDOSYA_01", "PedidosYa", 4, 31960128, 1331672, 24, 21),
    )
    for key, lender, entity, principal, installment, total, paid in seeds:
        conn.execute(text("""
          INSERT INTO public.fin_loans(
            external_key,lender,id_legal_entity,principal,installment_amount,
            total_installments,paid_installments,status,notes
          ) VALUES (
            :key,:lender,:entity,:principal,:installment,:total,:paid,'ACTIVE',
            'Carga inicial entregada por gerencia el 24-08-2026'
          ) ON CONFLICT(external_key) DO NOTHING
        """), {
            "key": key, "lender": lender, "entity": entity,
            "principal": principal, "installment": installment,
            "total": total, "paid": paid,
        })
    conn.commit()


@router.get("/loans-v1")
def loans_v1(me: dict = Depends(legacy.get_current_user)):
    legacy._ensure_roles(me, READ_ROLES)
    with legacy.get_connection() as conn:
        _ensure_fin_loans_v1(conn)
        rows = conn.execute(text("""
          SELECT l.id_loan,l.external_key,l.lender,l.id_legal_entity,
                 e.legal_code,e.legal_name,e.rut,l.principal,l.installment_amount,
                 l.total_installments,l.paid_installments,l.annual_rate,l.status,l.notes,
                 (l.installment_amount*l.paid_installments)::numeric AS paid_amount,
                 GREATEST(0,l.principal-(l.installment_amount*l.paid_installments))::numeric AS balance
          FROM public.fin_loans l
          JOIN public.fin_legal_entities e ON e.id_legal_entity=l.id_legal_entity
          ORDER BY l.id_loan
        """)).mappings().all()
    items = [dict(row) for row in rows]
    return {
        "ok": True,
        "items": items,
        "monthly_installments": float(sum(Decimal(str(row["installment_amount"])) for row in rows)),
        "principal_total": float(sum(Decimal(str(row["principal"])) for row in rows)),
        "paid_total": float(sum(Decimal(str(row["paid_amount"])) for row in rows)),
        "balance_total": float(sum(Decimal(str(row["balance"])) for row in rows)),
    }


@router.post("/loans-v1/simulate")
def loans_v1_simulate(body: dict = Body(default_factory=dict), me: dict = Depends(legacy.get_current_user)):
    legacy._ensure_roles(me, READ_ROLES)
    try:
        principal = Decimal(str(body.get("principal") or "0"))
        annual_rate = Decimal(str(body.get("annual_rate") or "0"))
        installments = int(body.get("installments") or 0)
        monthly_insurance = Decimal(str(body.get("monthly_insurance") or "0"))
        monthly_fees = Decimal(str(body.get("monthly_fees") or "0"))
    except Exception:
        raise HTTPException(status_code=400, detail="Parámetros de simulación inválidos.")
    if principal <= 0 or principal > Decimal("10000000000"):
        raise HTTPException(status_code=400, detail="El capital debe ser mayor que cero.")
    if annual_rate < 0 or annual_rate > Decimal("200"):
        raise HTTPException(status_code=400, detail="La tasa anual debe estar entre 0% y 200%.")
    if installments < 1 or installments > 600:
        raise HTTPException(status_code=400, detail="Las cuotas deben estar entre 1 y 600.")
    if monthly_insurance < 0 or monthly_fees < 0:
        raise HTTPException(status_code=400, detail="Seguros y comisiones no pueden ser negativos.")

    monthly_rate = annual_rate / Decimal("1200")
    if monthly_rate == 0:
        base_payment = principal / installments
    else:
        factor = (Decimal("1") + monthly_rate) ** installments
        base_payment = principal * monthly_rate * factor / (factor - Decimal("1"))

    balance = principal
    schedule = []
    total_interest = Decimal("0")
    total_cost = Decimal("0")
    for number in range(1, installments + 1):
        interest = balance * monthly_rate
        amortization = base_payment - interest
        if number == installments:
            amortization = balance
        payment = amortization + interest + monthly_insurance + monthly_fees
        balance = max(Decimal("0"), balance - amortization)
        total_interest += interest
        total_cost += payment
        schedule.append({
            "installment": number,
            "payment": float(payment.quantize(Decimal("1"))),
            "principal": float(amortization.quantize(Decimal("1"))),
            "interest": float(interest.quantize(Decimal("1"))),
            "insurance": float(monthly_insurance.quantize(Decimal("1"))),
            "fees": float(monthly_fees.quantize(Decimal("1"))),
            "balance": float(balance.quantize(Decimal("1"))),
        })
    return {
        "ok": True,
        "assumption": "Sistema francés; tasa nominal anual dividida en 12 meses.",
        "principal": float(principal),
        "annual_rate": float(annual_rate),
        "monthly_rate": float(monthly_rate * 100),
        "installments": installments,
        "base_installment": float(base_payment.quantize(Decimal("1"))),
        "monthly_total": float((base_payment + monthly_insurance + monthly_fees).quantize(Decimal("1"))),
        "total_interest": float(total_interest.quantize(Decimal("1"))),
        "total_cost": float(total_cost.quantize(Decimal("1"))),
        "schedule": schedule,
    }

# === GD FINANCE LOANS V1 END ===
