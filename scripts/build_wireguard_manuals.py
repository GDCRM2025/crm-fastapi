from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs"
OUT.mkdir(exist_ok=True)

NAVY = RGBColor(11, 37, 69)
BLUE = RGBColor(31, 77, 120)
GREEN = RGBColor(23, 143, 82)
MUTED = RGBColor(90, 107, 122)
RED = RGBColor(155, 28, 28)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)


def configure(doc, short_title):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(0.82)
    sec.left_margin = sec.right_margin = Inches(0.88)
    sec.header_distance = sec.footer_distance = Inches(0.42)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(10.8)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18
    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 8),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 11.5, NAVY, 10, 5),
    ):
        s = styles[name]
        s.font.name = "Calibri"
        s._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        s._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        s.font.size, s.font.bold, s.font.color.rgb = Pt(size), True, color
        s.paragraph_format.space_before, s.paragraph_format.space_after = Pt(before), Pt(after)
        s.paragraph_format.keep_with_next = True
    for name in ("List Bullet", "List Number"):
        s = styles[name]
        s.font.name = "Calibri"
        s.font.size = Pt(10.8)
        s.paragraph_format.left_indent = Inches(0.38)
        s.paragraph_format.first_line_indent = Inches(-0.19)
        s.paragraph_format.space_after = Pt(4)
        s.paragraph_format.line_spacing = 1.18
    hp = sec.header.paragraphs[0]
    hp.text = f"GREEN DIAMOND  |  {short_title}"
    hp.style = styles["Normal"]
    hp.runs[0].font.size = Pt(8.5)
    hp.runs[0].font.color.rgb = MUTED
    add_page_number(sec.footer.paragraphs[0])
    sec.footer.paragraphs[0].runs[0].font.size = Pt(8.5)
    return doc


def cover(doc, kicker, title, subtitle, audience):
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(kicker.upper())
    r.bold = True; r.font.size = Pt(10); r.font.color.rgb = GREEN
    p.paragraph_format.space_after = Pt(18)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold = True; r.font.size = Pt(29); r.font.color.rgb = NAVY
    p.paragraph_format.space_after = Pt(8)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(subtitle)
    r.font.size = Pt(14); r.font.color.rgb = BLUE
    p.paragraph_format.space_after = Pt(24)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(audience + "\nVersión 1.0 · 1 de agosto de 2026")
    r.font.size = Pt(10); r.font.color.rgb = MUTED
    doc.add_page_break()


def add_callout(doc, label, text, color="E8EEF5"):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.08)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(8)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), color); p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left"); left.set(qn("w:val"), "single"); left.set(qn("w:sz"), "18"); left.set(qn("w:color"), "178F52"); left.set(qn("w:space"), "8"); borders.append(left); p_pr.append(borders)
    r = p.add_run(label + ": ")
    r.bold = True; r.font.color.rgb = NAVY
    p.add_run(text)


def _fresh_numbering(doc):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType"); multi.set(qn("w:val"), "singleLevel"); abstract.append(multi)
    lvl = OxmlElement("w:lvl"); lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start"); start.set(qn("w:val"), "1"); lvl.append(start)
    num_fmt = OxmlElement("w:numFmt"); num_fmt.set(qn("w:val"), "decimal"); lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText"); lvl_text.set(qn("w:val"), "%1."); lvl.append(lvl_text)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs"); tab = OxmlElement("w:tab"); tab.set(qn("w:val"), "num"); tab.set(qn("w:pos"), "540"); tabs.append(tab); p_pr.append(tabs)
    ind = OxmlElement("w:ind"); ind.set(qn("w:left"), "540"); ind.set(qn("w:hanging"), "270"); p_pr.append(ind)
    lvl.append(p_pr); abstract.append(lvl); numbering.append(abstract)
    num = OxmlElement("w:num"); num.set(qn("w:numId"), str(num_id))
    ref = OxmlElement("w:abstractNumId"); ref.set(qn("w:val"), str(abstract_id)); num.append(ref); numbering.append(num)
    return num_id


def add_steps(doc, steps):
    num_id = _fresh_numbering(doc)
    for step in steps:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.188)
        p.paragraph_format.space_after = Pt(4)
        num_pr = p._p.get_or_add_pPr().get_or_add_numPr()
        ilvl = OxmlElement("w:ilvl"); ilvl.set(qn("w:val"), "0"); num_pr.append(ilvl)
        num = OxmlElement("w:numId"); num.set(qn("w:val"), str(num_id)); num_pr.append(num)
        p.add_run(step)


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (header, width) in enumerate(zip(headers, widths)):
        cell = table.rows[0].cells[i]
        cell.width = Inches(width)
        cell.text = header
        set_cell_shading(cell, "E8EEF5")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.paragraphs[0].runs[0].bold = True
        cell.paragraphs[0].runs[0].font.color.rgb = NAVY
    set_repeat_table_header(table.rows[0])
    for row in rows:
        cells = table.add_row().cells
        for i, (value, width) in enumerate(zip(row, widths)):
            cells[i].width = Inches(width)
            cells[i].text = str(value)
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def build_executive():
    doc = configure(Document(), "WireGuard para ejecutivos")
    cover(doc, "Acceso privado", "WireGuard para ejecutivos", "Cómo entrar al CRM desde PC o teléfono", "Ejecutivos comerciales")
    doc.add_heading("1. Qué hace la VPN", 1)
    doc.add_paragraph("WireGuard crea un túnel cifrado entre tu dispositivo y el servidor de Green Diamond. Debe estar activo antes de abrir el CRM cuando trabajas fuera de la oficina. La VPN no reemplaza tu usuario del CRM.")
    add_callout(doc, "Regla principal", "Cada dispositivo tiene su propio perfil. Nunca compartas el archivo .conf, el QR, la contraseña del equipo ni tu usuario del CRM.", "E7F5EC")
    doc.add_heading("2. Qué recibirás de Oscar", 1)
    add_bullets(doc, ["Un perfil con nombre de marca y dispositivo, por ejemplo GOURMET-PC.", "La URL privada oficial: http://10.77.0.1/crm/.", "Tu usuario y contraseña personal del CRM por un canal separado."])
    doc.add_heading("3. Instalación en Windows", 1)
    add_steps(doc, ["Descarga WireGuard desde https://www.wireguard.com/install/ y completa la instalación.", "Abre WireGuard y selecciona Importar túnel desde archivo.", "Elige solamente el archivo .conf entregado para ese computador.", "Comprueba que el nombre coincida con tu marca y termine en -PC.", "Pulsa Activar y abre http://10.77.0.1/crm/ en Chrome o Edge.", "Inicia sesión con tus credenciales personales del CRM."])
    doc.add_heading("4. Instalación en macOS", 1)
    add_steps(doc, ["Instala WireGuard desde App Store.", "Importa el archivo .conf entregado para ese Mac.", "Autoriza a macOS a agregar la configuración VPN.", "Activa el túnel y abre la URL privada del CRM."])
    doc.add_heading("5. Instalación en iPhone o Android", 1)
    add_steps(doc, ["Instala WireGuard desde App Store o Google Play.", "Oscar mostrará un QR del perfil de ese teléfono; no le tomes una foto ni lo reenvíes.", "En WireGuard pulsa +, escanea el QR y autoriza la VPN.", "Activa el túnel y abre la URL privada en el navegador."])
    add_callout(doc, "Importante", "El perfil del teléfono no sirve para el PC y viceversa. No se reutilizan perfiles.", "FFF1D6")
    doc.add_heading("6. Prueba obligatoria", 1)
    add_steps(doc, ["Desconéctate del Wi-Fi de la oficina y usa datos móviles u otra red.", "Activa WireGuard.", "Abre http://10.77.0.1/crm/.", "Inicia sesión, consulta un lead autorizado y cierra sesión.", "Desactiva WireGuard cuando ya no necesites el CRM si así lo indica la política interna."])
    doc.add_heading("7. Problemas frecuentes", 1)
    add_table(doc, ["Síntoma", "Qué hacer"], [
        ("No abre la URL", "Comprueba Internet, activa el túnel correcto y reintenta con datos móviles."),
        ("Handshake vacío", "No modifiques el perfil. Envía a Oscar una captura de la pantalla de WireGuard."),
        ("El CRM pide login", "Es normal: VPN y CRM son controles separados."),
        ("403 o sin marca", "Solicita revisión de permisos del CRM; la VPN sí está funcionando."),
        ("Equipo perdido", "Avisa inmediatamente para revocar ese perfil."),
    ], [1.8, 4.7])
    doc.add_heading("8. Cómo reportar un error", 1)
    add_bullets(doc, ["Tu nombre, marca y dispositivo.", "Fecha y hora exactas.", "Red utilizada: oficina, casa o datos móviles.", "Captura de estado de WireGuard sin mostrar claves.", "Mensaje del navegador y acción que estabas realizando."])
    add_callout(doc, "Nunca", "No edites PrivateKey, PublicKey, Endpoint, Address ni AllowedIPs. No envíes el perfil por WhatsApp o correo sin autorización.", "FCE8E8")
    path = OUT / "Manual_WireGuard_Ejecutivos_Green_Diamond.docx"
    doc.save(path)
    return path


def build_admin():
    doc = configure(Document(), "Administración WireGuard")
    cover(doc, "Infraestructura", "Supermanual WireGuard", "Red, Entel, perfiles, seguridad, operación y recuperación", "Administrador: Oscar Mendoza")
    doc.add_heading("1. Arquitectura aprobada", 1)
    doc.add_paragraph("El servidor Ubuntu funciona como concentrador WireGuard. Los clientes reciben una IP /32 en 10.77.0.0/24 y solo enrutan hacia 10.77.0.1. El CRM permanece privado; Meta continúa usando el webhook público de InMotion durante el piloto.")
    add_table(doc, ["Elemento", "Valor previsto"], [("Interfaz VPN", "wg0"), ("Servidor VPN", "10.77.0.1/24"), ("Puerto", "UDP 51820"), ("CRM", "http://10.77.0.1/crm/"), ("Servidor LAN actual", "192.168.10.51 — no cambiar hasta diagnosticar"), ("Gestor", "/usr/local/sbin/crm-vpn-peer")], [2.0, 4.5])
    doc.add_heading("2. Diagnóstico de red antes de cambiar nada", 1)
    add_steps(doc, ["Conecta Mac, Proxmox y Ubuntu a la red que realmente se usará en la oficina.", "En Mac ejecuta ipconfig getifaddr en0 y route -n get default | grep gateway.", "En Ubuntu ejecuta sudo bash /opt/greendiamond/crm/deploy/ubuntu/inspect-office-network.sh.", "Entra al router Entel y anota LAN, máscara, DHCP, WAN/Internet IPv4 y estado CGNAT.", "Compara la WAN del router con PUBLIC_IPV4 del script.", "Documenta el resultado antes de tocar Proxmox o Netplan."])
    add_callout(doc, "Decisión", "Si Ubuntu y los usuarios se alcanzan en 192.168.10.0/24 y el gateway es estable, conservar 192.168.10.51. Que el Wi-Fi se llame 2.4G o 5G no obliga a usar 192.168.100.0/24.", "E7F5EC")
    doc.add_heading("3. Cómo reconocer CGNAT o doble NAT", 1)
    add_bullets(doc, ["WAN del router diferente de la IP pública observada.", "WAN dentro de 10.0.0.0/8, 100.64.0.0/10, 172.16.0.0/12 o 192.168.0.0/16.", "Port forwarding correcto pero ningún paquete llega a wg0.", "Entel confirma que la línea comparte IPv4."])
    doc.add_heading("4. Solicitud exacta a Entel", 1)
    add_callout(doc, "Texto", "Necesitamos retirar CGNAT y disponer de IPv4 pública para recibir UDP 51820. El tráfico debe redireccionarse al servidor Ubuntu. Confirmen si la IPv4 es fija o dinámica, el costo y si el equipo permite reserva DHCP y port forwarding.", "E8EEF5")
    doc.add_heading("5. Si la LAN realmente debe cambiar", 1)
    add_steps(doc, ["Crea snapshot de la VM en Proxmox y backup de /etc/netplan.", "Usa la consola de Proxmox, no una sesión SSH, para evitar quedar fuera.", "Reserva primero la nueva IP en el router.", "Edita Netplan con nueva dirección, máscara, gateway y DNS; ejecuta netplan try, no netplan apply a ciegas.", "Comprueba gateway, DNS, Internet, SSH, Nginx y crm-gd.service.", "Actualiza documentación, rsync y monitoreo. WireGuard 10.77.0.0/24 no cambia por modificar la LAN."])
    add_callout(doc, "Retroceso", "Si falla netplan try, no confirmes: el sistema revierte. Desde consola restaura el archivo respaldado.", "FFF1D6")
    doc.add_heading("6. Configurar el endpoint", 1)
    doc.add_paragraph("No existe IP pública predeterminada en el gestor. Después de confirmar IPv4 o DNS dinámico:")
    doc.add_paragraph("echo 'IP_PUBLICA_O_DNS:51820' | sudo tee /etc/wireguard/crm-endpoint\nsudo chmod 600 /etc/wireguard/crm-endpoint", style=None)
    doc.add_paragraph("Si la IP es dinámica, utilizar un nombre DDNS controlado y comprobar su actualización. Cuando cambia el endpoint, los perfiles ya entregados deben actualizarse o regenerarse.")
    doc.add_heading("7. Router y firewall", 1)
    add_bullets(doc, ["Reserva DHCP o IP fija del servidor.", "Redirecciona únicamente UDP 51820 al servidor.", "No publiques 22, 80, 443, 5432 ni Proxmox 8006.", "UFW permite UDP 51820 y tráfico del túnel al CRM; deniega acceso lateral innecesario.", "Prueba desde datos móviles, nunca solamente desde la LAN."])
    doc.add_page_break()
    doc.add_heading("8. Gestión de perfiles", 1)
    add_table(doc, ["Acción", "Comando"], [("Crear", "sudo crm-vpn-peer add MARCA-DISPOSITIVO"), ("Listar", "sudo crm-vpn-peer list"), ("Entregado", "sudo crm-vpn-peer delivered MARCA-DISPOSITIVO"), ("Revocar", "sudo crm-vpn-peer revoke MARCA-DISPOSITIVO")], [1.5, 5.0])
    add_callout(doc, "Secuencia segura", "Crear → transferir personalmente → importar → probar fuera de la oficina → ejecutar delivered. Ese último comando elimina la copia exportable de la clave privada.", "E7F5EC")
    doc.add_heading("9. Inventario inicial", 1)
    add_table(doc, ["Marca/persona", "PC", "Teléfono"], [("Gourmet", "GOURMET-PC", "GOURMET-TELEFONO"), ("Camaleón", "CAMALEON-PC", "CAMALEON-TELEFONO"), ("Express", "EXPRESS-PC", "EXPRESS-TELEFONO"), ("Del Sabor", "DEL-SABOR-PC", "DEL-SABOR-TELEFONO"), ("Oscar", "OSCAR-PC", "OSCAR-TELEFONO")], [2.0, 2.25, 2.25])
    doc.add_heading("10. Distribución segura", 1)
    add_bullets(doc, ["PC/Mac: copiar el .conf directamente al equipo y eliminar la copia después de importar.", "Teléfono: mostrar QR en persona; no generar una imagen persistente.", "No usar WhatsApp, correo personal, Drive compartido ni enlaces públicos.", "Registrar responsable, dispositivo, fecha de entrega, prueba y revocación."])
    doc.add_heading("11. Pruebas de aceptación", 1)
    add_steps(doc, ["sudo wg show confirma escucha en 51820.", "Desde datos móviles activar el cliente.", "Confirmar latest handshake reciente y transferencia en ambos sentidos.", "Abrir http://10.77.0.1/crm/.", "Confirmar que el ejecutivo no accede a SSH, Proxmox, base de datos ni otras LAN.", "Reiniciar Ubuntu y verificar que wg-quick@wg0 y crm-gd.service vuelven automáticamente."])
    doc.add_heading("12. Monitoreo y mantenimiento", 1)
    add_bullets(doc, ["Revisar semanalmente peers sin uso o con nombres ambiguos.", "Aplicar actualizaciones Ubuntu en ventana controlada.", "Respaldar wg0.conf y crm-peers.tsv cifrados y con permisos 0600; contienen material sensible.", "No respaldar perfiles entregados: la clave privada debe vivir solo en su dispositivo.", "Probar restauración y revocación trimestralmente."])
    doc.add_heading("13. Incidentes", 1)
    add_table(doc, ["Situación", "Respuesta"], [("Equipo perdido", "Revocar peer inmediatamente y desactivar usuario CRM."), ("Perfil enviado por error", "Revocar y crear otro; no basta con borrar el mensaje."), ("IP pública cambió", "Actualizar DDNS/endpoint y perfiles afectados; probar externamente."), ("Sin handshake", "Revisar endpoint, CGNAT, redirección UDP, firewall y hora del sistema."), ("Servidor comprometido", "Aislar, preservar evidencia, rotar todos los peers y restaurar desde base confiable.")], [2.0, 4.5])
    doc.add_page_break()
    doc.add_heading("14. Lista de salida a producción", 1)
    add_bullets(doc, ["Red LAN y gateway documentados.", "WAN comparada con IPv4 pública.", "CGNAT resuelto por Entel.", "Snapshot y backup restaurable.", "UDP 51820 comprobado desde fuera.", "Perfil Oscar probado en Mac y teléfono.", "Perfil piloto Gourmet probado y revocable.", "Permisos CRM comprobados.", "Inventario de perfiles firmado.", "InMotion mantiene webhook Meta durante el piloto."])
    doc.add_heading("15. Regla de oro", 1)
    add_callout(doc, "Administración", "Nunca cambies simultáneamente LAN, firewall, endpoint y perfiles. Cambia una capa, prueba, registra y recién entonces continúa.", "FCE8E8")
    path = OUT / "Supermanual_Administrador_WireGuard_Green_Diamond.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    for built in (build_executive(), build_admin()):
        print(built)
