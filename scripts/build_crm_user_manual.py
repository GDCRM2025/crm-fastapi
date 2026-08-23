from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "Manual_Usuario_CRM_Green_Diamond.docx"
IMG = ROOT / "docs" / "crm_manual_images"
GREEN = "128A45"; NAVY = "102A43"; PALE = "EAF6EE"; GOLD = "D9A441"; RED = "B42318"; GREY = "5B6870"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); tcPr.append(shd)

def margins(section):
    section.top_margin=Inches(.55); section.bottom_margin=Inches(.55); section.left_margin=Inches(.62); section.right_margin=Inches(.62)

def title(doc, text, level=1):
    p=doc.add_paragraph(); p.style=f"Heading {level}"; p.add_run(text); return p

def bullets(doc, items):
    for x in items: doc.add_paragraph(x, style="List Bullet")

def steps(doc, items):
    for x in items: doc.add_paragraph(x, style="List Number")

def callout(doc, heading, body, color=PALE):
    t=doc.add_table(rows=1, cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=True
    c=t.cell(0,0); shade(c,color); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p=c.paragraphs[0]; r=p.add_run(heading+"  "); r.bold=True; r.font.color.rgb=RGBColor.from_string(NAVY); p.add_run(body)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

def picture(doc, name, caption):
    path=IMG/name
    if path.exists():
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(path), width=Inches(6.45))
        c=doc.add_paragraph(caption); c.alignment=WD_ALIGN_PARAGRAPH.CENTER; c.style="Caption"

def page(doc): doc.add_page_break()

doc=Document(); sec=doc.sections[0]; margins(sec)
styles=doc.styles
styles["Normal"].font.name="Aptos"; styles["Normal"].font.size=Pt(9.2)
styles["Normal"].paragraph_format.space_after=Pt(4)
for i,size in [(1,20),(2,14),(3,11)]:
    st=styles[f"Heading {i}"]; st.font.name="Aptos Display"; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(GREEN)
    st.paragraph_format.space_before=Pt(8); st.paragraph_format.space_after=Pt(4)
styles["Caption"].font.size=Pt(8); styles["Caption"].font.italic=True; styles["Caption"].font.color.rgb=RGBColor.from_string(GREY)

# Cover
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(95)
r=p.add_run("GREEN DIAMOND"); r.bold=True; r.font.size=Pt(16); r.font.color.rgb=RGBColor.from_string(GREEN)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run("Manual de usuario del CRM"); r.bold=True; r.font.size=Pt(28); r.font.color.rgb=RGBColor.from_string(NAVY)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run("Guía operacional · Comercial · Reportes · Postevento · RRHH"); r.font.size=Pt(13); r.font.color.rgb=RGBColor.from_string(GREY)
doc.add_paragraph("\n")
callout(doc,"Objetivo", "Trabajar con un criterio común, mantener datos trazables y resolver errores sin improvisar.")
p=doc.add_paragraph("Versión 1.0 · 1 de agosto de 2026"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p=doc.add_paragraph("Documento interno. Contiene procesos operacionales; no incluye contraseñas, tokens ni claves."); p.alignment=WD_ALIGN_PARAGRAPH.CENTER

page(doc); title(doc,"Cómo usar este manual")
doc.add_paragraph("En Word o PDF usa Ctrl+F (Windows) o ⌘F (Mac) para buscar una palabra. Dentro del CRM, el botón ? abre la Ayuda buscable y permite saltar al módulo correspondiente.")
title(doc,"Ruta rápida",2)
tbl=doc.add_table(rows=1,cols=2); tbl.alignment=WD_TABLE_ALIGNMENT.CENTER; tbl.style="Light Shading Accent 1"
tbl.rows[0].cells[0].text="Necesito…"; tbl.rows[0].cells[1].text="Ir a…"
for a,b in [("Crear o actualizar oportunidad","Leads → Ver Leads"),("Organizar pendientes","Leads → Comercial 360"),("Preparar propuesta","Cotizador"),("Reabrir PDF","Cotizador → Historial"),("Medir gestión","Reportes"),("Enviar encuesta manual","Checklist → Encuestas post-evento"),("Conversar por WhatsApp","Tools → WhatsApp Greenie"),("Revisar trazabilidad","Settings → Auditoría"),("Marcar entrada/salida","RRHH → Marcación")]:
    cells=tbl.add_row().cells; cells[0].text=a; cells[1].text=b
title(doc,"Reglas de oro",2)
bullets(doc,["Un cliente y evento deben vivir en un solo lead; busca antes de crear.","Registra cada contacto con un resultado útil y una próxima acción.","No envíes cotizaciones, mensajes o encuestas sin revisar cliente, marca y teléfono.","No compartas usuarios, perfiles VPN, tokens ni capturas con datos personales.","Si algo falla, conserva hora, módulo, acción e ID del lead; evita reintentos repetidos."])

page(doc); title(doc,"1. Inicio, sesión y navegación")
picture(doc,"01_inicio_crm.png","Pantalla de inicio. La marcación puede aparecer primero según la política RRHH.")
steps(doc,["Abre la URL autorizada del CRM.","Inicia sesión con tu usuario personal.","Si trabajas fuera de la oficina, activa primero el perfil WireGuard asignado.","Usa el menú lateral para cambiar de módulo y el botón ? para ayuda."])
callout(doc,"Si aparece «Iniciar sesión» dentro de un módulo", "Vuelve al inicio, recarga una vez e inicia sesión nuevamente. Si persiste, registra la hora y avisa al administrador.","FFF3CD")

page(doc); title(doc,"2. Leads: buscar, mes y vistas")
picture(doc,"02_leads_mes.png","Ver Leads permite elegir cualquier mes; al abrir, propone el mes actual.")
bullets(doc,["Kanban: trabajo visual por estado.","Tabla: revisión rápida y comparación de campos.","Buscador: cliente, correo, comuna o marca.","Mes: muestra eventos del mes seleccionado con su estado real.","Todos: úsalo solo cuando necesites salir del alcance mensual."])
callout(doc,"No encuentro un lead", "Revisa mes, vista, buscador, estado y permisos. Después prueba Filtrar Leads. No lo recrees hasta confirmar que no existe.")

page(doc); title(doc,"3. Crear un lead")
picture(doc,"03_crear_lead.png","Formulario de creación. Completa la mayor cantidad de datos correctos desde el primer contacto.")
steps(doc,["Pulsa + Crear lead.","Busca primero al cliente para evitar duplicados.","Completa nombre, marca, comuna, plataforma, teléfono y fecha del evento.","Selecciona la plataforma real. Carta web y botones WhatsApp son opciones válidas y administrables.","Agrega dirección y notas útiles; no copies datos sensibles innecesarios.","Pulsa Crear una sola vez y confirma que la tarjeta aparezca."])
callout(doc,"Calidad mínima", "Nombre identificable + marca + fecha de evento + teléfono/comuna cuando estén disponibles. Las notas deben indicar necesidad, contexto y próxima acción.")

page(doc); title(doc,"4. Editar, estados, seguimiento y eliminación")
title(doc,"Editar",2); steps(doc,["Abre el lead y pulsa Editar.","Corrige el dato necesario; no borres el historial.","Guarda y vuelve a abrir para verificar."])
title(doc,"Estados",2)
doc.add_paragraph("Los estados representan avance comercial. Registrar WhatsApp, llamada o seguimiento no siempre debe cambiar automáticamente la etapa: el resultado manda.")
bullets(doc,["Nuevo: sin gestión efectiva.","Contactado: hubo contacto real o respuesta útil.","Cotizado: propuesta generada/enviada.","Confirmado: cliente aceptó bajo la regla comercial vigente.","Declinado: oportunidad cerrada sin venta; registra motivo."])
title(doc,"Seguimiento",2); bullets(doc,["Anota canal, resultado, compromiso y fecha próxima.","«No contesta» crea historial y seguimiento, pero no equivale a contacto real.","Comercial 360 mostrará vencidos y prioridades."])
title(doc,"Eliminar",2); callout(doc,"Acción restringida", "Elimina solo duplicados o registros de prueba autorizados. Confirma identidad, captura ID y motivo. La auditoría debe conservar la acción.","FDECEC")

page(doc); title(doc,"5. Comercial 360")
picture(doc,"04_comercial_360.png","Centro diario para priorizar tareas, vencimientos y próximos contactos.")
steps(doc,["Abre Comercial 360 al iniciar la jornada.","Trabaja primero vencidos y compromisos del día.","Abre el lead desde la tarea, ejecuta la gestión y registra el resultado.","Cierra o reprograma la tarea con fecha y motivo concreto.","Revisa nuevamente antes de terminar la jornada."])
callout(doc,"Importante", "Tarea y estado comercial son controles distintos. Cerrar una tarea no significa confirmar el lead.")

page(doc); title(doc,"6. Cotizador: crear propuesta")
picture(doc,"05_cotizador.png","Cotizador. La disponibilidad de campos y productos depende de marca, evento y permisos.")
steps(doc,["Entra desde el lead correcto o abre Cotizador.","Confirma cliente, marca, fecha y lugar.","Agrega productos, cantidades y observaciones.","Revisa neto, descuento, IVA, traslado y total.","Previsualiza antes de guardar.","Guarda una sola vez y descarga el PDF."])
bullets(doc,["No cotices desde otro lead para ahorrar tiempo.","No fuerces descuentos fuera de tus permisos.","Si el PDF abre en una pestaña blob, usa Descargar; el nombre debe ser reconocible."])

page(doc); title(doc,"7. Historial, vista previa y PDF")
picture(doc,"06_historial_cotizaciones.png","Historial centraliza búsqueda, vista previa y descarga de cotizaciones.")
steps(doc,["Busca por cliente, marca o código.","Usa el ojo para previsualizar y la flecha para descargar.","Comprueba que número, cliente, marca y total coincidan.","Si vas a reenviar, descarga una copia nueva desde el CRM."])
callout(doc,"«Token requerido» o «token inválido»", "Recarga una vez desde el CRM, no desde un enlace guardado. Si continúa, informa número de cotización, usuario, hora y acción. Nunca pegues el token en un chat.","FDECEC")
callout(doc,"El navegador dice «No seguro»", "Una URL http o una descarga iniciada desde ella no tiene TLS. No significa que el PDF esté infectado; sí indica que falta HTTPS. En oficina/VPN usa solo la URL indicada por administración y no ignores avisos fuera de ese contexto.","FFF3CD")

page(doc); title(doc,"8. Reportes: leer antes de decidir")
picture(doc,"07_reportes.png","Reportes reúne KPIs y gráficos según rango y filtros.")
bullets(doc,["Embudo: distribución y pérdida entre etapas.","Cierre: confirmados respecto de oportunidades trabajadas.","Origen/plataforma: qué canales generan demanda.","Marca y ejecutivo: desempeño dentro del mismo rango.","Productos: oferta cotizada o confirmada.","Encuestas: satisfacción, recomendación y reseñas."])
steps(doc,["Elige rango y filtros.","Espera que terminen de cargar KPIs y gráficos.","Contrasta total con Leads usando el mismo mes y filtros.","Exporta solo cuando la pantalla tenga datos completos."])
callout(doc,"Comparación justa", "No compares un mes parcial con uno cerrado ni ejecutivos con filtros diferentes.")

page(doc); title(doc,"9. Encuestas de satisfacción")
picture(doc,"08_encuestas.png","Encuestas postevento: generación, envío manual, seguimiento y resultados.")
steps(doc,["Abre Checklist → Encuestas post-evento.","Busca un evento confirmado y ya realizado.","Abre el enlace y revisa logo, marca, cliente y preguntas.","Copia el mensaje y envíalo manualmente por el canal autorizado.","Marca el envío y revisa luego la respuesta.","Si corresponde, usa el enlace de reseña específico de la marca."])
bullets(doc,["Cada token pertenece a un evento; nunca lo reutilices.","No marques enviada antes de hacer el envío real.","La encuesta evalúa atención comercial, cumplimiento, presentación/servicio, recomendación y comentario abierto.","Los colores y logo dependen de la marca del evento."])

page(doc); title(doc,"10. WhatsApp Greenie")
picture(doc,"09_whatsapp_greenie.png","Greenie concentra conversaciones y contexto comercial sin sondeo cada dos segundos.")
steps(doc,["Selecciona la conversación correcta.","Confirma teléfono, nombre, marca y lead asociado.","Lee el historial antes de responder.","Escribe o elige la plantilla permitida.","Envía una sola vez y revisa estado.","Registra la gestión en el lead cuando corresponda."])
callout(doc,"Aceptado no es entregado", "Aceptado significa que Meta recibió la solicitud. Revisa enviado, entregado, leído o fallido. Fuera de la ventana de 24 horas se requiere una plantilla aprobada.","FFF3CD")
callout(doc,"No llegó", "No insistas muchas veces. Registra hora, teléfono, ID del lead y estado; luego se revisan webhook, ventana, plantilla, coexistencia y respuesta de Meta.","FDECEC")

page(doc); title(doc,"11. Auditoría y seguridad")
picture(doc,"10_auditoria.png","Auditoría permite revisar acciones relevantes sin modificar la operación.")
bullets(doc,["Filtra por día, usuario, módulo o acción antes de concluir que falta información.","Úsala para confirmar creación, edición, eliminación, cambios de estado y cotizaciones.","No publiques capturas con teléfonos, correos, tokens o claves.","Los permisos se asignan por rol y necesidad; no por conveniencia.","Bloquea la pantalla al alejarte y nunca compartas la sesión."])
callout(doc,"Alcance", "La auditoría ayuda a investigar quién hizo qué y cuándo. No reemplaza respaldos ni corrige datos por sí sola.")

page(doc); title(doc,"12. Marcación")
picture(doc,"11_marcacion.png","Marcación registra entrada/salida con reglas de sede, dispositivo y método.")
steps(doc,["Abre RRHH → Marcación.","Confirma sede/punto y tipo Entrada o Salida.","Permite ubicación/cámara solo al sitio autorizado cuando se solicite.","Marca una vez y espera confirmación.","Si falla, conserva mensaje, hora y ubicación de trabajo."])
bullets(doc,["No marques por otra persona.","No compartas dispositivo enrolado ni usuario.","Si la empresa deshabilita marcación por usar un reloj autorizado, no es un error.","Correcciones administrativas deben tener motivo y auditoría."])

page(doc); title(doc,"13. Solución rápida de problemas")
tbl=doc.add_table(rows=1,cols=3); tbl.style="Light Shading Accent 1"; tbl.alignment=WD_TABLE_ALIGNMENT.CENTER
for i,h in enumerate(["Síntoma","Qué hacer","Qué reportar"]): tbl.rows[0].cells[i].text=h
rows=[
("No inicia sesión","Revisa URL, red/VPN y credenciales; recarga una vez.","Hora, URL, mensaje y usuario."),
("Lead no aparece","Revisa mes, Todos, filtros, estado y permisos.","ID/nombre, mes y filtros."),
("Parpadeo al abrir lead","Cierra modal/vista y entra una vez desde Ver Leads.","Navegador, hora e ID."),
("Token requerido/inválido","Reabre desde Historial; no uses favorito antiguo.","Nº cotización, acción y hora."),
("PDF bloqueado/no seguro","Usa URL oficial; conserva solo dentro de red/VPN.","URL sin token y navegador."),
("Reporte sin gráficos","Espera carga, reduce rango, recarga una vez.","Rango, filtros y captura."),
("Encuesta sin logo/preguntas","Reabre desde Encuestas y valida token/evento.","Marca, evento y hora."),
("WhatsApp aceptado pero no llega","No reintentes; revisa estado y ventana 24 h.","ID lead, hora y estado."),
("VPN no conecta","Prueba internet, activa solo un perfil y revisa túnel.","Dispositivo, red externa y hora."),
]
for row in rows:
    c=tbl.add_row().cells
    for i,v in enumerate(row): c[i].text=v

page(doc); title(doc,"14. Cierre diario del ejecutivo")
bullets(doc,["No quedan tareas vencidas sin comentario o reprogramación.","Cada contacto tiene resultado y próxima acción.","Cotizaciones del día abren y descargan correctamente.","Leads confirmados/declinados tienen estado y motivo coherente.","Mensajes pendientes tienen estado revisado.","Encuestas enviadas manualmente están marcadas solo después del envío.","La sesión queda cerrada en equipos compartidos."])
title(doc,"Cómo pedir soporte",2)
doc.add_paragraph("Envía un solo reporte por problema con: usuario, fecha/hora, módulo, acción exacta, ID de lead o cotización, mensaje completo y captura sin claves/tokens. Indica si estabas en oficina o por VPN.")
callout(doc,"Evita", "«No funciona» sin contexto, múltiples reintentos, crear duplicados, cambiar estados para probar o enviar datos sensibles.","FDECEC")
title(doc,"Glosario",2)
bullets(doc,["Lead: oportunidad comercial ligada a un cliente/evento.","Estado: etapa real del proceso comercial.","Seguimiento: compromiso futuro con fecha y contexto.","Token: identificador temporal o de acceso; no se comparte.","Webhook: aviso automático entre Meta y el servidor.","VPN: túnel cifrado para acceder al servidor local desde fuera.","Auditoría: registro trazable de acciones relevantes."])

# header/footer
for section in doc.sections:
    margins(section)
    hp=section.header.paragraphs[0]; hp.text="GREEN DIAMOND · Manual de usuario CRM"; hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    hp.runs[0].font.size=Pt(8); hp.runs[0].font.color.rgb=RGBColor.from_string(GREY)
    fp=section.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    fp.add_run("Uso interno · ")
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); fp._p.append(fld)
    for run in fp.runs: run.font.size=Pt(8); run.font.color.rgb=RGBColor.from_string(GREY)

OUT.parent.mkdir(parents=True,exist_ok=True); doc.save(OUT); print(OUT)
